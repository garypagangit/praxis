from pathlib import Path
from datetime import datetime, timezone
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parent
NAME = 'Choosing_the_Right_Evidence_for_Cybersecurity_AI'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = Inches(.48)
sec.bottom_margin = Inches(.48)
sec.left_margin = sec.right_margin = Inches(.62)
sec.header_distance = sec.footer_distance = Inches(.2)

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11.2)
normal.font.color.rgb = RGBColor.from_string('243342')
normal.paragraph_format.line_spacing = 1.04
normal.paragraph_format.space_after = Pt(5)
for name in ['Title', 'Heading 1']:
    doc.styles[name].font.name = 'Calibri'

def link(p, text, url, size=None, italic=False):
    rel = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rel)
    r = OxmlElement('w:r'); props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '006F7C'); props.append(color)
    if size:
        sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(round(size*2))); props.append(sz)
    if italic:
        props.append(OxmlElement('w:i'))
    r.append(props); t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r); p._p.append(h)

def paragraph(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    run = p.add_run(label + '  '); run.bold = True; run.font.color.rgb = RGBColor.from_string('006F7C')
    p.add_run(text)
    return p

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
r = p.add_run('Choosing the Right Evidence for Cybersecurity AI')
r.bold = True; r.font.size = Pt(21); r.font.color.rgb = RGBColor.from_string('172B4D')
p = doc.add_paragraph('RESEARCH BRIEF  |  SEPTEMBER 18, 2026')
p.paragraph_format.space_after = Pt(8)
for r in p.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('667584')

paragraph('Thesis', 'A checker that recognizes when retrieved security facts actually apply could improve cybersecurity AI while preserving the benefits of useful evidence.')
paragraph('Problem and impact', 'Security AI can be misled by apparently relevant evidence, potentially wasting analysts\' time and leading to poor security decisions (Hamzić et al., 2026; Liu et al., 2026).')
paragraph('Hypothesis', 'A checker trained to predict whether evidence will help will improve accuracy on unfamiliar security questions, limit harm from mismatched facts, and outperform simpler relevance checks.')
p = paragraph('Primary base paper and gap', '')
link(p, 'CoRM-RAG (Liu et al., 2026)', 'https://arxiv.org/html/2605.01302v1')
p.add_run(' already learns whether evidence is useful. Our target is a checker that recognizes when security facts apply to unfamiliar questions and beats simpler checks while using facts equally often. A new contribution has not yet been demonstrated. Hamzić et al. (2026) provide the broader CTI foundation.')

p = paragraph('Dataset and example', 'Development used 2,500 ')
link(p, 'CTIBench', 'https://arxiv.org/abs/2406.07599v3')
p.add_run(' questions; external testing used 1,247 qualifying SecEval questions from 10 security source families, with retrieved MITRE ATT&CK 19.1 facts (Li et al., 2023). SecEval answers were AI-generated/calibrated; independent human review is pending.')
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(.13)
p.paragraph_format.keep_together = True
r = p.add_run('Real example, shortened: '); r.bold = True
p.add_run('Which tool best tests a web application for SQL injection? A: sqlmap; B: wget; C: Find Security Bugs; D: OWASP O-Saft. Released answer: A.')
for r in p.runs: r.font.size = Pt(10.6)

paragraph('Experiment in simple terms', 'Freeze the checker before testing. Have Llama 3.1 8B and Qwen 2.5 7B answer every question with and without six retrieved facts, producing 4,988 fresh answers. Let the checker choose which answer to use; compare with always using facts, never using them, and four simpler checkers, including comparisons using facts equally often.')

p = paragraph('Results', 'The external test did not support the hypothesis.')
p.paragraph_format.keep_with_next = True
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [1.55, 1.8, 1.8, 2.08]
rows = [
    ['Model', 'No extra facts', 'Always add facts', 'Our checker'],
    ['Llama 3.1 8B', '86.53%', '84.20%', '86.21%'],
    ['Qwen 2.5 7B', '87.09%', '85.24%', '87.25%'],
]
for i, values in enumerate(rows):
    cells = table.rows[0].cells if i == 0 else table.add_row().cells
    for j, value in enumerate(values):
        cell = cells[j]; cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr = cell._tc.get_or_add_tcPr()
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '172B4D' if i == 0 else ('EDF4F5' if i == 2 else 'F6F8FA')); tcpr.append(shade)
        p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(3)
        if j: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value); r.font.size = Pt(10.3); r.bold = i == 0
        if i == 0: r.font.color.rgb = RGBColor(255,255,255)
    trpr = table.rows[i]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
p = doc.add_paragraph('The checker used facts on 6.82% of questions; its average change from no extra facts was -0.08 percentage points (95% interval: -0.52 to +0.36). The planned improvement and comparison tests failed. Earlier pilot gains of +12.24 and +9.64 points did not transfer to this test.')
p.paragraph_format.space_before = Pt(4)

paragraph('Conclusion and next step', 'A better checker must retain helpful facts while rejecting mismatches. Complete the prepared 50-question human review, investigate explicit checks for the relevant system, version, and conditions, then compare against strong published checkers on untouched questions. This is a direction to test, not a proven solution or established novelty.')

p = doc.add_paragraph('APA REFERENCES')
p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
for r in p.runs: r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('006F7C')
refs = [
    ('Hamzić, D., Skopik, F., Landauer, M., Wurzenberger, M., & Rauber, A. (2026). ', 'Beyond RAG for cyber threat intelligence: A systematic evaluation of graph-based and agentic retrieval', '. arXiv. ', 'https://doi.org/10.48550/arXiv.2604.11419'),
    ('Li, G., Li, Y., Wang, G., Yang, H., & Yu, Y. (2023). ', 'SecEval: A comprehensive benchmark for evaluating cybersecurity knowledge of foundation models', ' [Data set]. GitHub. ', 'https://github.com/XuanwuAI/SecEval'),
    ('Liu, P., Yan, Q., Cui, Z., Liang, D., Wang, X., & Ye, W. (2026). ', 'Beyond semantic relevance: Counterfactual risk minimization for robust retrieval-augmented generation', '. arXiv. ', 'https://doi.org/10.48550/arXiv.2605.01302'),
]
for authors,title,tail,url in refs:
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.14); p.paragraph_format.first_line_indent=Inches(-.14)
    p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_together=True
    p.add_run(authors); r=p.add_run(title); r.italic=True; p.add_run(tail)
    for r in p.runs: r.font.size=Pt(9)
    link(p,url,url,9)

doc.core_properties.title = 'Choosing the Right Evidence for Cybersecurity AI'
doc.core_properties.subject = 'One-page summary of evidence selection and external validation'
doc.core_properties.author = 'Gary Pagan'
doc.core_properties.keywords = 'cybersecurity; CTI; evidence checker; external validation'
doc.save(ROOT / (NAME+'.docx'))

source = Path('C:/Users/garyp/OneDrive/Documents/codex/reports/cti_external_validation_20260918')
results=json.loads((source/'RESULTS.json').read_text(encoding='utf-8'))
assert results['n']==1247 and results['fresh_outputs']==4988
assert results['status']=='EXTERNAL_TRANSPORT_CRITERIA_NOT_MET'
audit=json.loads((source/'INDEPENDENT_RESULTS_AUDIT.json').read_text(encoding='utf-8'))
assert audit['status']=='PASS'
metadata={'created_utc':datetime.now(timezone.utc).isoformat(), 'source_report':str(source/'REPORT.md'),
          'dataset_example_id':'seceval_2125','example_question_shortened':True,'example_choices_and_released_label_preserved':True,
          'source_commit':'719b8742cbecc96b882f81e9f118e748b1849228','human_review_status':'PENDING',
          'literature_verified_date':'2026-09-18','new_scientific_claims':False}
(ROOT/'CONTENT_PROVENANCE.json').write_text(json.dumps(metadata,indent=2)+'\n',encoding='utf-8')
print(ROOT/(NAME+'.docx'))
