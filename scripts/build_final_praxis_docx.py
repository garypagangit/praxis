"""Regenerable Markdown-to-DOCX Praxis reports, with PDF/page render artifacts.

Uses python-docx, LibreOffice and PyMuPDF. Scientific content comes only from each
paper/PRAXIS_REPORT.md; this script adds formatting and a documented author cover.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import re
import subprocess
import textwrap

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
EXPERIMENTS={"001":"001_outcome_state_verification","002":"002_cascade_containment",
             "003":"003_adaptive_investigation_stopping"}
DEFAULT_OUT=ROOT/"output/doc/final_praxis_20260908"
DASHES=str.maketrans({"\u2010":"-","\u2011":"-","\u2012":"-","\u2013":"-","\u2014":"-","\u2212":"-"})
ACCENT="173A4A"

def clean(text):
    return text.translate(DASHES).replace("<br>","\n").replace("<br/>","\n").replace("<br />","\n")

def shade(cell,fill):
    node=OxmlElement("w:shd");node.set(qn("w:fill"),fill);cell._tc.get_or_add_tcPr().append(node)

def add_link(paragraph,label,target):
    relationship=paragraph.part.relate_to(target,RT.HYPERLINK,is_external=True)
    link=OxmlElement("w:hyperlink");link.set(qn("r:id"),relationship)
    run=OxmlElement("w:r");props=OxmlElement("w:rPr")
    color=OxmlElement("w:color");color.set(qn("w:val"),ACCENT);props.append(color)
    underline=OxmlElement("w:u");underline.set(qn("w:val"),"single");props.append(underline)
    run.append(props);text=OxmlElement("w:t");text.text=clean(label);run.append(text);link.append(run)
    paragraph._p.append(link)

INLINE=re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\)|(?<!\*)\*[^*]+\*(?!\*)|https?://[^\s<>]+)")

def inline(paragraph,text):
    cursor=0
    for match in INLINE.finditer(text):
        if match.start()>cursor:paragraph.add_run(clean(text[cursor:match.start()]))
        token=match.group()
        if token.startswith("**"):
            paragraph.add_run(clean(token[2:-2])).bold=True
        elif token.startswith("`"):
            run=paragraph.add_run(clean(token[1:-1]));run.font.name="Consolas";run.font.size=Pt(9)
        elif token.startswith("["):
            label,target=re.match(r"\[([^\]]+)\]\(([^\)]+)\)",token).groups()
            add_link(paragraph,label,target)
        elif token.startswith("*"):
            paragraph.add_run(clean(token[1:-1])).italic=True
        elif token.startswith("http"):
            target=token.rstrip(".,;")
            add_link(paragraph,target,target)
            if len(target)!=len(token):paragraph.add_run(token[len(target):])
        cursor=match.end()
    if cursor<len(text):paragraph.add_run(clean(text[cursor:]))

def configure(document,study):
    for section in document.sections:
        section.page_width=Inches(8.5);section.page_height=Inches(11)
        section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
        section.header_distance=Inches(.45);section.footer_distance=Inches(.45)
    normal=document.styles["Normal"]
    normal.font.name="Times New Roman";normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=1.8
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for name,size in (("Title",22),("Heading 1",16),("Heading 2",14),("Heading 3",12)):
        style=document.styles[name];style.font.name="Times New Roman";style.font.size=Pt(size)
        style.font.color.rgb=RGBColor.from_string(ACCENT)
        style.font.bold=True;style.paragraph_format.keep_with_next=True
        style.paragraph_format.space_before=Pt(14);style.paragraph_format.space_after=Pt(8)
        style.paragraph_format.line_spacing=1.15
    for name in ("List Bullet","List Number"):
        document.styles[name].font.name="Times New Roman";document.styles[name].font.size=Pt(12)
        document.styles[name].paragraph_format.line_spacing=1.5
        document.styles[name].paragraph_format.space_after=Pt(4)
    caption=document.styles["Caption"];caption.font.name="Times New Roman";caption.font.size=Pt(10)
    caption.font.italic=True;caption.font.color.rgb=RGBColor.from_string("404040")
    caption.paragraph_format.line_spacing=1.1;caption.paragraph_format.space_after=Pt(8)
    props=document.core_properties
    props.author="Gary Pagan";props.title=f"Final Praxis {study} research report"
    props.subject="Verified experiment report with reproducibility and bounded claims"
    props.keywords=f"Final Praxis {study}; reproducibility; research"

def cover(document,title,study):
    p=document.add_paragraph();p.paragraph_format.space_after=Pt(62)
    p=document.add_paragraph(f"FINAL PRAXIS {study}");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:run.bold=True;run.font.size=Pt(14);run.font.color.rgb=RGBColor.from_string(ACCENT)
    p=document.add_paragraph(clean(title),"Title");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(20);p.paragraph_format.space_after=Pt(46)
    p=document.add_paragraph("Research Report");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=document.add_paragraph("Gary Pagan");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=document.add_paragraph("September 8, 2026");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=document.add_paragraph("Methods, frozen decisions, results and reproducibility");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:run.italic=True;run.font.size=Pt(11)
    section=document.add_section(WD_SECTION_START.NEW_PAGE)
    section.header.is_linked_to_previous=False;section.footer.is_linked_to_previous=False
    pg=OxmlElement("w:pgNumType");pg.set(qn("w:start"),"1");section._sectPr.append(pg)
    h=section.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=h.add_run(f"Final Praxis {study}");r.font.name="Times New Roman";r.font.size=Pt(9)
    r.font.color.rgb=RGBColor.from_string("666666")
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
    footer._p.append(field)

def table_cells(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]

def add_table(document,lines):
    values=[table_cells(line) for line in lines]
    values=[row for row in values if not all(re.fullmatch(r":?-+:?",cell.replace(" ","")) for cell in row)]
    columns=max(len(row) for row in values)
    if columns<1:return
    for row in values:row.extend([""]*(columns-len(row)))
    table=document.add_table(rows=1,cols=columns);table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    weights=[max(8,min(48,max(len(re.sub(r"\[[^\]]+\]\([^\)]+\)","source",row[c])) for row in values)))**.5 for c in range(columns)]
    widths=[6.5*weight/sum(weights) for weight in weights]
    font=8.5 if columns>=7 else 9.5 if columns>=5 else 10
    for c,width in enumerate(widths):table.columns[c].width=Inches(width)
    for ridx,rowvalues in enumerate(values):
        row=table.rows[0] if ridx==0 else table.add_row()
        prop=row._tr.get_or_add_trPr();no_split=OxmlElement("w:cantSplit");prop.append(no_split)
        if ridx==0:
            repeat=OxmlElement("w:tblHeader");prop.append(repeat)
        for c,value in enumerate(rowvalues):
            cell=row.cells[c];cell.width=Inches(widths[c]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade(cell,"E8EFF2" if ridx==0 else "F5F7F8" if ridx%2==0 else "FFFFFF")
            p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.05
            p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.keep_with_next=False
            inline(p,value)
            for run in p.runs:
                run.font.name="Times New Roman";run.font.size=Pt(font)
                if ridx==0:run.bold=True
            # Hyperlinks do not appear in paragraph.runs; set their size in raw XML too.
            for run in p._p.findall(".//"+qn("w:r")):
                rp=run.find(qn("w:rPr"))
                if rp is None:rp=OxmlElement("w:rPr");run.insert(0,rp)
                size=OxmlElement("w:sz");size.set(qn("w:val"),str(int(font*2)));rp.append(size)
    document.add_paragraph().paragraph_format.space_after=Pt(1)

def diagram_from_mermaid(source,path):
    """Render the actual node/edge graph; unsupported syntax stays as source text."""
    declarations=re.findall(r"([A-Za-z]\w*)\[([^\]]+)\]",source)
    nodes=dict(declarations)
    stripped=re.sub(r"([A-Za-z]\w*)\[[^\]]+\]",r"\1",source)
    edges=re.findall(r"([A-Za-z]\w*)\s*-->\s*([A-Za-z]\w*)",stripped)
    if not nodes or not edges or any(a not in nodes or b not in nodes for a,b in edges):return False
    ranks={name:0 for name in nodes}
    for _ in range(len(nodes)):
        changed=False
        for a,b in edges:
            if ranks[b]<=ranks[a]:ranks[b]=ranks[a]+1;changed=True
        if not changed:break
    if max(ranks.values())>=len(nodes):return False
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    depth=max(ranks.values())+1;height=min(7.2,max(2.6,depth*.69))
    fig,ax=plt.subplots(figsize=(6.5,height));ax.set_xlim(0,1);ax.set_ylim(-.2,depth-.1);ax.axis("off")
    positions={}
    for level in range(depth):
        row=[name for name in nodes if ranks[name]==level]
        for index,name in enumerate(row):positions[name]=((index+.5)/len(row),depth-1-level)
    for a,b in edges:
        x1,y1=positions[a];x2,y2=positions[b]
        ax.annotate("",xy=(x2,y2+.25),xytext=(x1,y1-.25),arrowprops={"arrowstyle":"->","color":"#47636E","lw":1.25},zorder=1)
    for name,(x,y) in positions.items():
        siblings=sum(ranks[n]==ranks[name] for n in nodes);width=.9/siblings
        box=FancyBboxPatch((x-width/2,y-.25),width,.5,boxstyle="round,pad=0.012",facecolor="#E8EFF2",edgecolor="#47636E",lw=1,zorder=2)
        ax.add_patch(box)
        text=clean(nodes[name]).strip('"').replace("<br/>","\n")
        ax.text(x,y,textwrap.fill(text,width=max(16,int(68/siblings))),ha="center",va="center",fontsize=10,fontfamily="serif",zorder=3)
    fig.tight_layout(pad=.3);fig.savefig(path,dpi=200,bbox_inches="tight",facecolor="white");plt.close(fig)
    return True

def add_image(document,path,caption):
    from PIL import Image
    with Image.open(path) as pic:width,height=pic.size
    available_width=6.3;available_height=6.9
    desired=min(available_width,available_height*width/height)
    p=document.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.line_spacing=1.0
    p.add_run().add_picture(str(path),width=Inches(desired))
    p=document.add_paragraph(clean(caption),"Caption");p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def markdown_to_document(source,study,out,assets):
    content=source.read_text(encoding="utf-8-sig");lines=content.splitlines()
    title=next((line[2:].strip() for line in lines if line.startswith("# ")),f"Final Praxis {study}")
    document=Document();configure(document,study);cover(document,title,study)
    i=0;title_seen=False;figures=[];tables=0;headings=[]
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        heading=re.match(r"^(#{1,6})\s+(.+)$",line)
        if heading:
            depth=len(heading[1]);text=heading[2].strip()
            if depth==1 and not title_seen:title_seen=True;i+=1;continue
            level=max(1,min(3,depth-1));p=document.add_paragraph(clean(text),f"Heading {level}")
            if re.match(r"(?:Chapter\s*\d|References\b|Appendix\b)",text,re.I):p.paragraph_format.page_break_before=True
            headings.append(text);i+=1;continue
        if line.startswith("```"):
            language=line[3:].strip();block=[];i+=1
            while i<len(lines) and not lines[i].strip().startswith("```"):block.append(lines[i]);i+=1
            i+=1
            picture=assets/f"{study}_diagram_{len(figures)+1}.png"
            if language=="mermaid" and diagram_from_mermaid("\n".join(block),picture):
                add_image(document,picture,"Graphical methodological representation (GMR).")
                figures.append({"path":str(picture),"kind":"mermaid_render","source_sha256":hashlib.sha256("\n".join(block).encode()).hexdigest()})
            else:
                for code in block:
                    p=document.add_paragraph();p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(0)
                    r=p.add_run(clean(code));r.font.name="Consolas";r.font.size=Pt(8.5)
            continue
        image_match=re.fullmatch(r"!\[([^\]]*)\]\(([^\)]+)\)",line)
        if image_match:
            caption,target=image_match.groups();path=(source.parent/target).resolve()
            if not path.is_file():raise FileNotFoundError(f"Referenced figure missing: {path}")
            if path.suffix.lower()==".svg":
                fallback=path.with_suffix(".png")
                if not fallback.exists():raise ValueError(f"A PNG rendering is required for {path}")
                path=fallback
            add_image(document,path,caption);figures.append({"path":str(path),"kind":"source_figure","sha256":hashlib.sha256(path.read_bytes()).hexdigest()});i+=1;continue
        if line.startswith("|") and i+1<len(lines) and re.match(r"^\s*\|?[\s:|-]+\|\s*$",lines[i+1]):
            table=[]
            while i<len(lines) and lines[i].strip().startswith("|"):table.append(lines[i]);i+=1
            add_table(document,table);tables+=1;continue
        if re.fullmatch(r"[-*_]{3,}",line):i+=1;continue
        bullet=re.match(r"^[-*+]\s+(.+)$",line)
        numbered=re.match(r"^(\d+)\.\s+(.+)$",line)
        if bullet:
            p=document.add_paragraph(style="List Bullet");inline(p,bullet[1]);i+=1;continue
        if numbered:
            p=document.add_paragraph();p.paragraph_format.left_indent=Inches(.25)
            p.paragraph_format.first_line_indent=Inches(-.25);p.paragraph_format.line_spacing=1.3
            inline(p,numbered[1]+". "+numbered[2]);i+=1;continue
        paragraph=[line];i+=1
        while i<len(lines) and lines[i].strip() and not re.match(r"^(#|```|\||!\[|[-*+]\s|\d+\.\s)",lines[i].strip()):
            paragraph.append(lines[i].strip());i+=1
        p=document.add_paragraph();inline(p," ".join(paragraph))
    document.save(out)
    return {"study":study,"source":str(source),"source_sha256":hashlib.sha256(source.read_bytes()).hexdigest(),
            "docx":str(out),"docx_sha256":hashlib.sha256(out.read_bytes()).hexdigest(),"tables":tables,
            "figures":figures,"headings":headings,"author_basis":"Gary Pagan named as owner in repository reports"}

def render(docx,outdir,work,soffice):
    profile=(work/(docx.stem+"_lo_profile")).resolve();profile.mkdir(parents=True,exist_ok=True)
    command=[str(soffice),f"-env:UserInstallation={profile.as_uri()}","--headless","--convert-to","pdf","--outdir",str(outdir),str(docx)]
    completed=subprocess.run(command,capture_output=True,text=True,timeout=180,
        creationflags=subprocess.CREATE_NO_WINDOW if os.name=="nt" else 0)
    pdf=outdir/(docx.stem+".pdf")
    if completed.returncode!=0 or not pdf.exists():raise RuntimeError(f"LibreOffice render failed: {completed.stdout}\n{completed.stderr}")
    if pdf.stat().st_mtime<docx.stat().st_mtime-1:
        raise RuntimeError("LibreOffice left a stale PDF; visual proof must correspond to the regenerated DOCX.")
    import pymupdf
    pages_dir=work/(docx.stem+"_pages");pages_dir.mkdir(parents=True,exist_ok=True)
    pdfdoc=pymupdf.open(pdf);page_paths=[];page_evidence=[]
    for index,page in enumerate(pdfdoc):
        path=pages_dir/f"page-{index+1:02d}.png";page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False).save(path)
        page_paths.append(str(path))
        blocks=page.get_text("blocks")
        outside=[list(block[:4]) for block in blocks if block[0]<30 or block[2]>page.rect.width-30 or block[1]<15 or block[3]>page.rect.height-15]
        page_evidence.append({"page":index+1,"text_characters":len(page.get_text()),"outside_page_safety_bounds":outside})
    pdfdoc.close()
    return {"pdf":str(pdf),"pages":page_paths,"page_count":len(page_paths),"render_stdout":completed.stdout.strip(),
            "page_geometry_checks":page_evidence,"visual_review":"PENDING - each page must be inspected before delivery"}

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--experiment",choices=list(EXPERIMENTS),action="append")
    parser.add_argument("--output-dir",type=Path,default=DEFAULT_OUT)
    parser.add_argument("--work-dir",type=Path,default=ROOT/"tmp/docs/final_praxis_20260908")
    parser.add_argument("--render",action="store_true")
    parser.add_argument("--soffice",type=Path,default=Path("C:/Program Files/LibreOffice/program/soffice.exe"))
    args=parser.parse_args();args.output_dir.mkdir(parents=True,exist_ok=True);args.work_dir.mkdir(parents=True,exist_ok=True)
    assets=args.work_dir/"assets";assets.mkdir(exist_ok=True)
    results=[];missing=[]
    for study in args.experiment or list(EXPERIMENTS):
        source=ROOT/"final_praxis"/EXPERIMENTS[study]/"paper/PRAXIS_REPORT.md"
        if not source.is_file():missing.append(str(source));continue
        out=args.output_dir/f"FINAL_PRAXIS_{study}_REPORT.docx"
        result=markdown_to_document(source,study,out,assets)
        if args.render:result.update(render(out,args.output_dir,args.work_dir,args.soffice))
        manifest=args.output_dir/f"FINAL_PRAXIS_{study}_BUILD_MANIFEST.json"
        manifest.write_text(json.dumps(result,indent=2)+"\n",encoding="utf-8");results.append(result)
    print(json.dumps({"created":results,"missing_reports":missing},indent=2))
    if args.experiment and missing:raise SystemExit(2)

if __name__=="__main__":main()
