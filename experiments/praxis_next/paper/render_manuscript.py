"""Render the reviewed Markdown with the repository's document formatter."""
from pathlib import Path
import importlib.util
import hashlib
import json
import subprocess
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
spec = importlib.util.spec_from_file_location('praxis_document', REPO/'experiments/apt_benchmark/lateral_protection_experiment/paper/build_document.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
url = 'https://github.com/garypagangit/praxis/blob/praxis-next-20260923'
renderer = module.Renderer(HERE/'manuscript.md', url, None)
document = renderer.render()
document.styles['Praxis Reference'].paragraph_format.space_after = Pt(3)
for paragraph in document.paragraphs:
    if paragraph.text.startswith(('Table 1.', 'Table 2.', 'Table 3.')):
        paragraph.paragraph_format.keep_with_next = True
    if paragraph.text.startswith('Design choices.'):
        paragraph.paragraph_format.keep_together = True
document.core_properties.subject = 'Historical evidence, attack-stage recognition, and temporal evaluation'
for section in document.sections:
    header = section.header.paragraphs[0]
    header.text = 'EMPIRICAL PRAXIS  |  Context and attack-stage recognition'
    for run in header.runs:
        run.font.name = 'Arial'; run.font.size = Pt(8)
    for footer in (section.footer, section.first_page_footer):
        p = footer.paragraphs[0]
        for run in p.runs:
            if 'False alarms' in run.text:
                run.text = 'Context and attack-stage recognition  |  '
docx = HERE/'historical_context_praxis.docx'
document.save(docx)
qa = Path('C:/w/praxis_next_document_qa')
qa.mkdir(parents=True, exist_ok=True)
command = [r'C:\Program Files\LibreOffice\program\soffice.exe',
           '-env:UserInstallation=file:///C:/w/praxis_next_lo_profile',
           '--headless','--convert-to','pdf','--outdir',str(HERE),str(docx)]
completed = subprocess.run(command, capture_output=True, text=True, timeout=120)
if completed.returncode:
    raise RuntimeError(completed.stdout + completed.stderr)
pdf = docx.with_suffix('.pdf')
if not pdf.is_file() or pdf.stat().st_mtime < docx.stat().st_mtime:
    raise RuntimeError('PDF missing or stale')
import fitz
from PIL import Image, ImageOps, ImageDraw
opened = fitz.open(pdf)
pages = []
for i, page in enumerate(opened):
    pix = page.get_pixmap(matrix=fitz.Matrix(1.45,1.45), alpha=False)
    path = qa/f'page_{i+1:02d}.png'; pix.save(path); pages.append(str(path))
    if not page.get_text().strip():
        raise ValueError('Unexpected empty page')
thumbs = []
for i, path in enumerate(pages):
    im = Image.open(path).convert('RGB'); im.thumbnail((360, 480))
    thumb = Image.new('RGB',(380,510),'#dddddd')
    thumb.paste(im,((380-im.width)//2,20))
    ImageDraw.Draw(thumb).text((12,490),f'Page {i+1}',fill='black'); thumbs.append(thumb)
for start in range(0,len(thumbs),8):
    chunk = thumbs[start:start+8]
    sheet = Image.new('RGB',(380*4,510*((len(chunk)+3)//4)), '#bbbbbb')
    for j, thumb in enumerate(chunk): sheet.paste(thumb,((j%4)*380,(j//4)*510))
    sheet.save(qa/f'contact_{start//8+1}.png')
receipt = {'source_sha256':hashlib.sha256((HERE/'manuscript.md').read_bytes()).hexdigest(),
           'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),
           'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),
           'pages':len(pages),'tables':renderer.tables,'images':renderer.images,
           'hyperlinks':renderer.links,'local_only_links':renderer.local_only_links,
           'visual_review':'PENDING', 'page_images':pages}
(HERE/'RENDER_RECEIPT.json').write_text(json.dumps(receipt,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt))
