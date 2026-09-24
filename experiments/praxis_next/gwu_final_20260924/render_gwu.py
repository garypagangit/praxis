"""Render the reviewed five-chapter praxis with GWU typography and live indexes.

This changes document presentation only. It performs no model fitting or analysis.
"""
from __future__ import annotations
import argparse
from datetime import datetime,timezone
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import re
import socket
import subprocess
import time

from docx import Document
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches,Pt,RGBColor
from PIL import Image,ImageDraw,ImageFont
import fitz
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
plt.rcParams['mathtext.fontset']='stix'

HERE=Path(__file__).resolve().parent
REPO=HERE.parents[2]
BASE_PATH=REPO/'experiments/apt_benchmark/lateral_protection_experiment/paper/build_document.py'
spec=importlib.util.spec_from_file_location('praxis_document_base',BASE_PATH)
base=importlib.util.module_from_spec(spec);spec.loader.exec_module(base)
base.TEXT_WIDTH=6.0
TITLE='When Better APT Scores Hide Missed Attack Warnings'
TABLE_TITLES={
    '2-1':'Closest literature and study scope',
    '3-2':'Dataset roles and qualification boundaries',
    '3-3':'UNRAVELED class support by partition',
    '3-4':'Prepared features and availability',
    '3-5':'Published aggregate example',
    '3-6':'Completed model-fitting inventory',
    '3-7':'Boosting iterations, tree size, and learning rate',
    '3-8':'Regularization and execution settings',
    '3-9':'Implemented and unexecuted approaches',
    '4-2':'Acquisition differences by condition and budget',
    '4-3':'Clean budget-three differences by fitting seed',
    '4-4':'Chronological history differences by fitting seed',
    '4-6':'Fixed-anchor temporal differences by fitting seed',
    '4-9':'Seed and capture omissions at clean budget three',
    '4-10':'Current-feature temporal gains after capture omission',
    'A-1':'Complete evaluation inventory',
    'A-2':'History-selection group means',
    'A-3':'Evidence-acquisition group means',
    'A-4':'Temporal-comparison group means',
    'A-5':'CasinoLimit policy group means',
    'A-6':'CAM-LDS policy group means',
    'B-1':'Experimental settings and fitting seeds',
    'B-2':'Source freezes and evidence bindings',
    'B-3':'Claim traceability',
}
SOFFICE=Path('C:/Program Files/LibreOffice/program/soffice.exe')
UNO_PYTHON=Path('C:/Program Files/LibreOffice/program/python.exe')


def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def field(paragraph,text):
    run=paragraph.add_run()
    begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin')
    instruction=OxmlElement('w:instrText');instruction.text=' '+text+' '
    separate=OxmlElement('w:fldChar');separate.set(qn('w:fldCharType'),'separate')
    value=OxmlElement('w:t');value.text='1'
    end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end')
    for node in (begin,instruction,separate,value,end):run._r.append(node)
    run.font.name='Times New Roman';run.font.size=Pt(12)


class GWURenderer(base.Renderer):
    def __init__(self,source,abstract,title,qa):
        self.abstract=abstract;self.qa=qa;self.chapter='0';self.figure_count=0;self.table_count=0
        self.equations=[];self.caption_inventory=[];self.table_splits=[]
        self.landscape_images=[];self.image_assets=[];self.fresh_section=False;self.in_landscape=False
        super().__init__(source,'',title)
        self.frontmatter()

    def configure(self):
        section=self.doc.sections[0]
        section.page_width,section.page_height=Inches(8.5),Inches(11)
        section.left_margin=section.right_margin=Inches(1.25)
        section.top_margin=section.bottom_margin=Inches(1)
        section.header_distance=section.footer_distance=Inches(.5)
        section.different_first_page_header_footer=True
        n=self.doc.styles['Normal'];n.font.name='Times New Roman';n.font.size=Pt(12)
        n.font.color.rgb=RGBColor(0,0,0)
        n.paragraph_format.line_spacing=2
        n.paragraph_format.space_after=Pt(0)
        n.paragraph_format.first_line_indent=Inches(.5)
        n.paragraph_format.widow_control=True
        for name in ('Heading 1','Heading 2','Heading 3'):
            s=self.doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12);s.font.bold=True
            s.font.color.rgb=RGBColor(0,0,0);f=s.paragraph_format
            f.first_line_indent=Inches(0);f.line_spacing=2;f.space_before=Pt(12);f.space_after=Pt(0)
            f.keep_with_next=True;f.keep_together=True
            if name=='Heading 1':f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.space_before=Pt(0);f.space_after=Pt(12)
        for name,size,line in [('GWU Front Heading',12,1),('GWU Front',12,1),('GWU Figure Caption',10,1),
            ('GWU Table Caption',10,1),('Praxis Table',10,1),('Praxis Code',10,1),('Praxis Reference',12,1),
            ('Praxis List',12,2),('GWU Equation',12,1)]:
            s=self.doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);s.base_style=n
            s.font.name='Times New Roman';s.font.size=Pt(size);f=s.paragraph_format
            f.first_line_indent=Inches(0);f.line_spacing=line;f.space_before=Pt(0);f.space_after=Pt(6)
            f.widow_control=True
        for name in ('GWU Figure Caption','GWU Table Caption'):
            s=self.doc.styles[name];s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
            s.paragraph_format.keep_together=True
        self.doc.styles['GWU Table Caption'].paragraph_format.keep_with_next=True
        self.doc.styles['GWU Front Heading'].font.bold=True
        self.doc.styles['GWU Front Heading'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
        self.doc.styles['GWU Front Heading'].paragraph_format.space_after=Pt(24)
        self.doc.styles['GWU Front Heading'].paragraph_format.keep_with_next=True
        s=self.doc.styles.add_style('GWU Index Heading',WD_STYLE_TYPE.PARAGRAPH)
        s.base_style=self.doc.styles['GWU Front Heading']
        s.paragraph_format.page_break_before=True
        r=self.doc.styles['Praxis Reference'].paragraph_format
        r.left_indent=Inches(.5);r.first_line_indent=Inches(-.5);r.space_after=Pt(12)
        r.keep_together=True
        self.doc.styles['Praxis Code'].font.name='Consolas'
        for name in ('TOC 1','TOC 2','TOC 3'):
            if name not in self.doc.styles:self.doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            s=self.doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12)
            s.paragraph_format.line_spacing=1;s.paragraph_format.space_after=Pt(6)
            s.paragraph_format.first_line_indent=Inches(0)
        numbering=OxmlElement('w:pgNumType');numbering.set(qn('w:fmt'),'lowerRoman');numbering.set(qn('w:start'),'1')
        section._sectPr.append(numbering)
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent=Inches(0);footer.paragraph_format.line_spacing=1
        field(footer,'PAGE')
        section.first_page_footer.paragraphs[0].text=''
        self.doc.core_properties.title=self.title;self.doc.core_properties.author='Gary Pagan'
        self.doc.core_properties.subject='Doctor of Engineering praxis; APT evaluation and attack warnings'
        settings=self.doc.settings.element
        update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)

    def front(self,text='',*,bold=False,before=0,after=0,center=True):
        p=self.doc.add_paragraph(style='GWU Front');p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before=Pt(before);p.paragraph_format.space_after=Pt(after)
        run=p.add_run(text);run.bold=bold
        return p

    def new_front(self,title):
        p=self.doc.add_paragraph(title,style='GWU Front Heading')
        p.paragraph_format.page_break_before=True
        return p

    def frontmatter(self):
        self.front(self.title,bold=True,before=65,after=36)
        self.front('by Gary Pagan',after=36)
        self.front('B.S. in Applied Math and Statistics, May 1999, Stony Brook University')
        self.front('M.S. in Engineering Management, May 2006, George Washington University',after=24)
        self.front('A Praxis submitted to',after=24)
        self.front('The Faculty of\nThe School of Engineering and Applied Science\nof The George Washington University\nin partial fulfillment of the requirements\nfor the degree of Doctor of Engineering',after=24)
        self.front('Manuscript prepared September 24, 2026',after=18)
        self.front('Degree-conferral and defense dates are not asserted in this manuscript.',after=18)
        self.front('Praxis direction and committee confirmation pending')
        self.new_front('Certification pending')
        p=self.doc.add_paragraph('This manuscript is prepared for academic review. It does not certify that a final examination has been passed or that the praxis has received institutional approval.')
        p.paragraph_format.first_line_indent=Inches(0)
        p=self.doc.add_paragraph('The confirmed director and committee, examination date, and approved certification wording must be supplied through the university process before formal submission.')
        p.paragraph_format.first_line_indent=Inches(0)
        p=self.front('\u00a9 Copyright 2026 by Gary Pagan\nAll rights reserved.',before=260)
        p.paragraph_format.page_break_before=True
        self.new_front('Abstract of Praxis');self.front(self.title,bold=True,after=24)
        for text in re.split(r'\n\s*\n',self.abstract.read_text(encoding='utf-8').strip()):
            if text.lstrip().startswith('#'):continue
            p=self.doc.add_paragraph();self.inline(p,' '.join(text.splitlines()))
        for marker in ('@@TOC@@','@@FIGURES@@','@@TABLES@@'):
            p=self.doc.add_paragraph(marker,style='GWU Front')
        self.new_front('List of Symbols')
        for text in ['x: observed feature vector','y: declared evaluation class','p: model class-score vector',
            'g: fitted history selector','w: declared stage-cost vector','B: simulated evidence budget',
            'D: simulated decision deadline','C(s,j): confusion count for true class s and predicted class j']:
            self.front(text,center=False,after=12)
        self.new_front('List of Acronyms')
        for text in ['APT: Advanced Persistent Threat','AUC: Area Under the Curve','EDA: Exploratory Data Analysis',
            'GML: Graph Machine Learning','GMR: Graphical Model of Research','IDS: Intrusion Detection System',
            'ML: Machine Learning','ROC: Receiver Operating Characteristic','RQ: Research Question']:
            self.front(text,center=False,after=12)
        section=self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.different_first_page_header_footer=False
        section.footer.is_linked_to_previous=False;section.header.is_linked_to_previous=False
        for p in section.header.paragraphs:p.text=''
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.first_line_indent=Inches(0);footer.paragraph_format.line_spacing=1
        field(footer,'PAGE')
        numbering=section._sectPr.find(qn('w:pgNumType'))
        if numbering is None:numbering=OxmlElement('w:pgNumType');section._sectPr.append(numbering)
        numbering.set(qn('w:fmt'),'decimal');numbering.set(qn('w:start'),'1')

    def inline(self,p,text,*,bold=False,italic=False):
        # Academic citations remain readable; implementation links do not add file-path clutter.
        text=re.sub(r'\b(Chapter|Figure|Table)\s+(?=[A-Z0-9])',lambda m:m[1]+'\u00a0',text)
        for part in base.TOKEN.split(text):
            if not part:continue
            match=base.LINK.fullmatch(part)
            if match:
                self.inline(p,match[1],bold=bold,italic=italic);continue
            if part.startswith('**'):
                self.inline(p,part[2:-2],bold=True,italic=italic);continue
            if part.startswith('*') and part.endswith('*'):
                self.inline(p,part[1:-1],bold=bold,italic=True);continue
            code=part.startswith('`') and part.endswith('`')
            run=p.add_run(base.safe_text(part[1:-1] if code else part));run.bold=bold;run.italic=italic
            run.font.name='Times New Roman'
            if code:run.font.name='Consolas';run.font.size=Pt(10)

    def heading(self,level,text):
        if level==1:
            match=re.match(r'Chapter\s+(\d+)',text,re.I) or re.match(r'Appendix\s+([A-Z])',text,re.I)
            if match:self.chapter=match[1];self.figure_count=self.table_count=0
        p=self.doc.add_paragraph(style='Heading '+str(min(level,3)))
        if level==1 and self.headings and not self.fresh_section:p.paragraph_format.page_break_before=True
        self.inline(p,text,bold=True)
        self.headings+=1;self.current_numbering=None
        self.in_references=text.lower().strip() in ('references','bibliography')

    def numbering(self,kind,start=1):
        number=super().numbering(kind,start)
        if kind=='bullet':
            abstract=self.doc.part.numbering_part.element.findall(qn('w:abstractNum'))[-1]
            abstract.find(qn('w:lvl')).find(qn('w:lvlText')).set(qn('w:val'),'\u2022')
        return number

    def caption(self,kind,text):
        attr='figure_count' if kind=='Figure' else 'table_count'
        setattr(self,attr,getattr(self,attr)+1)
        number=f'{self.chapter}-{getattr(self,attr)}'
        text=base.plain(text)
        text=re.sub(r'^(?:Figure|Table)\s+[A-Za-z0-9]+(?:[.-][A-Za-z0-9]+)*[.:]?\s*','',text,flags=re.I)
        text=re.sub(r'^(?:Figure|Table)\s*[:.]\s*','',text,flags=re.I)
        if kind=='Table':text=TABLE_TITLES.get(number,text)
        p=self.doc.add_paragraph(style='GWU '+kind+' Caption')
        self.inline(p,f'{kind} {number}. {text}')
        self.caption_inventory.append(dict(kind=kind,number=number,title=text))
        return p

    def table(self,lines):
        # Expand a wide table into print-sized panels, repeating its identity column.
        raw=[]
        for line in lines:
            cells=[x.strip() for x in line.strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-{3,}:?',c) for c in cells):raw.append(cells)
        if len(raw[0])>6:
            for start in range(1,len(raw[0]),4):
                indices=[0]+list(range(start,min(start+4,len(raw[0]))))
                rows=['| '+' | '.join(r[i] for i in indices)+' |' for r in raw]
                if start>1:self.caption('Table','Model settings, continued')
                super().table(rows)
                self.doc.tables[-1].alignment=WD_TABLE_ALIGNMENT.CENTER
            self.table_splits.append(dict(columns=len(raw[0]),panels=(len(raw[0])-2)//4+1))
        else:
            super().table(lines)
            self.doc.tables[-1].alignment=WD_TABLE_ALIGNMENT.CENTER

    def image(self,alt,target):
        path=(self.source.parent/target).resolve()
        with Image.open(path) as im:ratio=im.height/im.width
        landscape=path.name.startswith(('fig04_','fig05_','fig06_','fig07_','fig08_')) or path.name in {
            'gmr_original_reference.png','warning_destinations.png','temporal_effects.png','dapt_stage_support.png'}
        if landscape:
            self.figure_section(True)
            width=min(8.5,5.35/ratio)
            self.landscape_images.append(dict(source=str(path),width_in=width,height_in=width*ratio))
        else:
            if self.in_landscape:self.figure_section(False)
            width=min(6.0,7.45/ratio)
        self.image_assets.append(dict(source=str(path),sha256=sha(path),width_in=width,height_in=width*ratio,landscape=landscape))
        p=self.doc.add_paragraph(style='GWU Equation');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
        picture=p.add_run().add_picture(str(path),width=Inches(width))
        picture._inline.docPr.set('descr',base.plain(alt))
        self.caption('Figure',alt or path.stem);self.images+=1

    def figure_section(self,landscape):
        section=self.doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        section.page_width=Inches(11 if landscape else 8.5)
        section.page_height=Inches(8.5 if landscape else 11)
        section.left_margin=section.right_margin=Inches(1.25)
        section.top_margin=section.bottom_margin=Inches(1)
        numbering=section._sectPr.find(qn('w:pgNumType'))
        if numbering is not None:
            numbering.set(qn('w:fmt'),'decimal')
            numbering.attrib.pop(qn('w:start'),None)
        section.footer.is_linked_to_previous=True
        self.fresh_section=True;self.in_landscape=landscape

    def equation(self,text):
        if self.doc.paragraphs and self.doc.paragraphs[-1].text.rstrip().endswith(':'):
            self.doc.paragraphs[-1].paragraph_format.keep_with_next=True
        chunks=re.split(r'\\qquad\s*',text.strip())
        expressions=[]
        for chunk in chunks:
            if '\\begin{cases}' in chunk:
                prefix,cases=chunk.split('\\begin{cases}',1)
                cases=cases.split('\\end{cases}')[0]
                for row in cases.split('\\\\'):
                    if '&' in row:
                        value,condition=row.split('&',1)
                        expressions.append(prefix.strip()+value.strip().rstrip(',')+r'\quad\mathrm{if}\quad '+condition.strip().rstrip(',.'))
            elif chunk.strip():expressions.append(chunk.strip())
        for expression_index,expression in enumerate(expressions):
            formula=re.sub(r'\s+',' ',expression).strip()
            formula=formula.replace(r'\operatorname*',r'\operatorname').replace(r'\lVert',r'\Vert').replace(r'\rVert',r'\Vert')
            formula=formula.replace(r'\mathbf 1',r'\mathbf{1}').replace(r'\mathcal L',r'\mathcal{L}')
            formula=re.sub(r'\\(ge|le|ne)\b',lambda m:'\\'+{'ge':'geq','le':'leq','ne':'neq'}[m[1]],formula)
            formula=re.sub(r'\\mathcal\s+([A-Z])',r'\\mathcal{\1}',formula)
            formula=re.sub(r'\\text\{([^}]+)\}',lambda m:r'\mathrm{'+m[1].replace(' ',r'\ ')+r'}',formula)
            fig=plt.figure(figsize=(8,1),dpi=300)
            artist=fig.text(.02,.5,'$'+formula+'$',fontsize=14,fontfamily='serif')
            fig.canvas.draw();bounds=artist.get_window_extent(fig.canvas.get_renderer()).expanded(1.04,1.22)
            bbox=bounds.transformed(fig.dpi_scale_trans.inverted())
            name=self.qa/'equations'/f'equation_{len(self.equations)+1:03d}.png';name.parent.mkdir(parents=True,exist_ok=True)
            fig.savefig(name,dpi=300,bbox_inches=bbox,transparent=False,facecolor='white');plt.close(fig)
            natural_width=bbox.width;display_width=min(6,natural_width);effective=14*display_width/natural_width
            if effective<10:raise ValueError(f'Equation too wide for10pt minimum: {expression}')
            p=self.doc.add_paragraph(style='GWU Equation');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(8)
            p.paragraph_format.keep_with_next=expression_index<len(expressions)-1
            picture=p.add_run().add_picture(str(name),width=Inches(display_width));picture._inline.docPr.set('descr',expression)
            self.equations.append(dict(source=expression,mathtext=formula,image=str(name),effective_font_pt=effective,sha256=sha(name)))

    def render(self):
        i=0;pending_caption=None
        while i<len(self.lines):
            line=self.lines[i].strip()
            if not line:i+=1;continue
            h=base.HEADING.match(line);im=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
            if not im and self.in_landscape:self.figure_section(False)
            ordered=re.match(r'(\d+)\.\s+(.+)',line);bullet=re.match(r'[-+*]\s+(.+)',line)
            clean=base.plain(line)
            if line.startswith('$$'):
                content=line[2:];bits=[]
                if content.endswith('$$'):bits=[content[:-2]]
                else:
                    if content:bits.append(content)
                    i+=1
                    while i<len(self.lines) and self.lines[i].strip()!='$$':bits.append(self.lines[i]);i+=1
                self.equation('\n'.join(bits))
            elif line=='<!-- PAGE BREAK -->':self.doc.add_page_break()
            elif line.startswith('<!--'):
                while '-->' not in self.lines[i] and i+1<len(self.lines):i+=1
            elif line.startswith('```'):
                i+=1
                while i<len(self.lines) and not self.lines[i].strip().startswith('```'):
                    self.doc.add_paragraph(base.safe_text(self.lines[i]),style='Praxis Code');i+=1
            elif h:self.heading(len(h[1]),h[2])
            elif im:self.image(im[1],im[2])
            elif re.match(r'^Table\s+(?:[A-Z0-9]|:)',clean):
                title_note=re.fullmatch(r'\*\*(Table.+?)\*\*\s*(.*)',line)
                pending_caption=base.plain(title_note[1]) if title_note else clean
                self.caption('Table',pending_caption)
                if title_note and title_note[2]:
                    p=self.doc.add_paragraph()
                    p.paragraph_format.first_line_indent=Inches(0)
                    p.paragraph_format.keep_with_next=True
                    self.inline(p,title_note[2])
            elif line.startswith('|'):
                end=i
                while end<len(self.lines) and self.lines[end].strip().startswith('|'):end+=1
                if pending_caption is None:self.caption('Table','Summary of '+self.last_heading if hasattr(self,'last_heading') else 'Study information')
                self.table(self.lines[i:end]);pending_caption=None;i=end;continue
            elif ordered:self.list_item(ordered[2],int(ordered[1]))
            elif bullet:self.list_item(bullet[1],None)
            elif re.fullmatch(r'[-*_]{3,}',line):pass
            else:
                parts=[line]
                while i+1<len(self.lines):
                    n=self.lines[i+1].strip()
                    if not n or re.match(r'^(?:#{1,6}\s|\||```|\$\$|<!--|!\[|[-+*]\s|\d+\.\s)',n):break
                    parts.append(n);i+=1
                p=self.doc.add_paragraph(style='Praxis Reference' if self.in_references else 'Normal');self.inline(p,' '.join(parts))
            if h:self.last_heading=base.plain(h[2])
            if not im:self.fresh_section=False
            i+=1
        return self.doc


def server_ready():
    try:
        with socket.create_connection(('127.0.0.1',20824),timeout=.4):return True
    except OSError:return False


def start_server(profile):
    if server_ready():return None
    startup=subprocess.STARTUPINFO();startup.dwFlags|=subprocess.STARTF_USESHOWWINDOW;startup.wShowWindow=0
    proc=subprocess.Popen([str(SOFFICE),'-env:UserInstallation='+profile.resolve().as_uri(),'--headless',
        '--norestore','--nodefault','--nofirststartwizard','--accept=socket,host=127.0.0.1,port=20824;urp;StarOffice.ServiceManager'],
        startupinfo=startup,creationflags=subprocess.CREATE_NO_WINDOW,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    for _ in range(80):
        if server_ready():return proc
        time.sleep(.25)
    raise RuntimeError('Hidden LibreOffice server did not start')


def build(args):
    args.qa_dir.mkdir(parents=True,exist_ok=True)
    r=GWURenderer(args.source,args.abstract,args.title,args.qa_dir)
    doc=r.render();args.docx.parent.mkdir(parents=True,exist_ok=True);args.pdf.parent.mkdir(parents=True,exist_ok=True)
    doc.save(args.docx)
    service=start_server(args.qa_dir/'lo_profile')
    try:
        result=subprocess.run([str(UNO_PYTHON),str(HERE/'render_uno.py'),str(args.docx.resolve()),str(args.pdf.resolve())],
            capture_output=True,text=True,encoding='utf-8',errors='replace',timeout=180)
        if result.returncode:raise RuntimeError('UNO render failed: '+result.stderr+'\n'+result.stdout)
    finally:
        if service is not None:
            # The helper uses the dedicated port/profile. Terminate only that service through UNO.
            subprocess.run([str(UNO_PYTHON),str(HERE/'render_uno.py'),'--terminate'],capture_output=True,timeout=15)
    pdf=fitz.open(args.pdf);geometry=[];texts=[]
    for i,page in enumerate(pdf):
        path=args.qa_dir/f'page_{i+1:03d}.png';page.get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save(path)
        lines=[]
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines',[]):
                x0,y0,x1,y1=line['bbox']
                if x0<-.5 or y0<-.5 or x1>page.rect.width+.5 or y1>page.rect.height+.5:lines.append([x0,y0,x1,y1])
        text=page.get_text();texts.append(text)
        geometry.append(dict(page=i+1,width=page.rect.width,height=page.rect.height,outside_page=lines))
    for start in range(0,len(pdf),12):
        thumbs=[]
        for i in range(start,min(start+12,len(pdf))):
            im=Image.open(args.qa_dir/f'page_{i+1:03d}.png').convert('RGB');im.thumbnail((255,330))
            tile=Image.new('RGB',(275,360),'white');tile.paste(im,((275-im.width)//2,22));ImageDraw.Draw(tile).text((10,5),str(i+1),fill='black');thumbs.append(tile)
        sheet=Image.new('RGB',(275*4,360*3),'#dddddd')
        for j,im in enumerate(thumbs):sheet.paste(im,((j%4)*275,(j//4)*360))
        sheet.save(args.qa_dir/f'contact_{start+1:03d}_{min(start+12,len(pdf)):03d}.png')
    receipt=dict(status='RENDERED_PENDING_VISUAL_REVIEW',utc=datetime.now(timezone.utc).isoformat(),
        title=args.title,source=str(args.source),source_sha256=sha(args.source),abstract_sha256=sha(args.abstract),
        renderer_sha256=sha(__file__),uno_helper_sha256=sha(HERE/'render_uno.py'),base_renderer_sha256=sha(BASE_PATH),
        docx=str(args.docx),docx_sha256=sha(args.docx),pdf=str(args.pdf),pdf_sha256=sha(args.pdf),
        pages=len(pdf),headings=r.headings,tables=r.tables,figures=r.images,caption_inventory=r.caption_inventory,
        equations=r.equations,table_splits=r.table_splits,image_assets=r.image_assets,landscape_images=r.landscape_images,geometry=geometry,uno_stdout=result.stdout,
        native_index_fields=[x.text for x in Document(args.docx).element.iter(qn('w:instrText')) if x.text and 'TOC ' in x.text],
        marker_tokens_remaining=any('@@' in t for t in texts),visual_review='PENDING',
        notes='GWU typography and native indexes; certification is explicitly pending, no institutional approval asserted.')
    args.receipt.parent.mkdir(parents=True,exist_ok=True);args.receipt.write_text(json.dumps(receipt,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(dict(status=receipt['status'],pages=len(pdf),docx=str(args.docx),pdf=str(args.pdf),equations=len(r.equations))))


if __name__=='__main__':
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--source',type=Path,default=HERE/'manuscript.md')
    ap.add_argument('--abstract',type=Path,default=HERE/'abstract.md')
    ap.add_argument('--title',default=TITLE)
    ap.add_argument('--docx',type=Path,default=REPO/'output/doc/gwu_praxis_20260924/Gary_Pagan_GWU_APT_Evaluation_Praxis.docx')
    ap.add_argument('--pdf',type=Path,default=REPO/'output/pdf/gwu_praxis_20260924/Gary_Pagan_GWU_APT_Evaluation_Praxis.pdf')
    ap.add_argument('--qa-dir',type=Path,default=Path('C:/w/gwu_final_document_qa_20260924'))
    ap.add_argument('--receipt',type=Path,default=HERE/'RENDER_RECEIPT.json')
    build(ap.parse_args())
