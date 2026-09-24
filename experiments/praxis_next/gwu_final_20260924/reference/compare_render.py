"""Read-only mapping of the reviewed baseline to the polished document.

Caption-only substitutions and footer page numbers are excluded from the
content mapping. A match is not itself a visual review of the new page.
"""
from pathlib import Path
import hashlib
import json
import re
import fitz
from docx import Document

HERE=Path(__file__).resolve().parent.parent
BASE=Path('C:/w/gwu_final_document_qa_prepolish_20260924')
QA=Path('C:/w/gwu_final_document_qa_20260924')

def digest(data):return hashlib.sha256(data).hexdigest()
def norm(text):return re.sub(r'\s+',' ',text.replace('\u00a0',' ')).strip()

def records(path,receipt,qa):
    result=[]
    captions=sorted((('Table '+c['number']+'. '+c['title'],'TABLE_CAPTION_'+c['number'])
        for c in receipt['caption_inventory'] if c['kind']=='Table'),key=lambda p:-len(p[0]))
    pdf=fitz.open(path)
    for i,page in enumerate(pdf):
        blocks=page.get_text('dict')['blocks']
        lines=[line for b in blocks for line in b.get('lines',[])]
        text=norm(' '.join(' '.join(s['text'] for s in line['spans']) for line in lines
            if line['bbox'][3]<page.rect.height-50))
        exact_text=text
        for title,token in captions:text=text.replace(norm(title),token)
        # get_images() includes the shared document resource dictionary; image
        # blocks identify only images actually displayed on this page.
        images=[digest(fitz.Pixmap(b['image']).samples) for b in blocks if b.get('type')==1]
        pixel=qa/f'page_{i+1:03d}.png'
        result.append({'page':i+1,'text_hash':digest(exact_text.encode()),'normalized_text_hash':digest(text.encode()),
            'image_hashes':images,'pixel_sha256':digest(pixel.read_bytes()),
            'content_key':digest(json.dumps([text,images]).encode()),
            'body_in_footer':[line['bbox'] for line in lines if line['bbox'][3]>page.rect.height-55
                and len(' '.join(s['text'] for s in line['spans']))>5],
            'blank_except_footer':len(text)<4,'landscape':page.rect.width>page.rect.height})
    return result

def main():
    old=json.loads((BASE/'RENDER_RECEIPT.json').read_text())
    new=json.loads((HERE/'RENDER_RECEIPT.json').read_text())
    before=records(BASE/'before.pdf',old,BASE)
    after=records(new['pdf'],new,QA)
    by_content={r['content_key']:r for r in before}
    by_pixels={r['pixel_sha256']:r for r in before}
    mapped=[]
    for r in after:
        match=by_content.get(r['content_key'])
        pixels=by_pixels.get(r['pixel_sha256'])
        mapped.append({'final_page':r['page'],
            'baseline_exact_pixel_page':pixels['page'] if pixels else None,
            'baseline_same_content_page':match['page'] if match else None,
            'normalization':'Footer omitted; approved table-caption substitutions represented by table number; embedded image bytes included.'})
    output={'baseline_pdf_sha256':digest((BASE/'before.pdf').read_bytes()),
        'final_pdf_sha256':digest(Path(new['pdf']).read_bytes()),'baseline_pages':len(before),'final_pages':len(after),
        'mapping':mapped,'unmatched_content_pages':[r['final_page'] for r in mapped if r['baseline_same_content_page'] is None],
        'changed_pixel_pages':[r['final_page'] for r in mapped if r['baseline_exact_pixel_page'] is None],
        'blank_pages':[r['page'] for r in after if r['blank_except_footer']],
        'body_in_footer_pages':[r['page'] for r in after if r['body_in_footer']],
        'scope':'Mechanical mapping aids final review; it does not certify unseen pages.'}
    old_doc,new_doc=Document(BASE/'before.docx'),Document(new['docx'])
    cells=lambda doc:[[[c.text for c in row.cells] for row in table.rows] for table in doc.tables]
    output['table_cell_text_unchanged']=cells(old_doc)==cells(new_doc)
    output['manuscript_source_unchanged']=old['source_sha256']==new['source_sha256']
    output['abstract_source_unchanged']=old['abstract_sha256']==new['abstract_sha256']
    output['figure_input_bytes_unchanged']=[x['sha256'] for x in old['image_assets']]==[x['sha256'] for x in new['image_assets']]
    output['equation_sources_unchanged']=[x['source'] for x in old['equations']]==[x['source'] for x in new['equations']]
    (HERE/'reference/POLISH_PAGE_MAP.json').write_text(json.dumps(output,indent=2)+'\n')
    print(json.dumps({k:v for k,v in output.items() if k!='mapping'},indent=2))

if __name__=='__main__':main()
