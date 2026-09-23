"""Assemble a preserved-study review edition and defense brief; no model fitting."""
from pathlib import Path
import hashlib
import importlib.util
import json
import os
import re
import subprocess
from datetime import datetime, timezone
from urllib.parse import urlsplit

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
ORIGINAL = HERE.parent / 'measurement_praxis'
QA = Path('C:/w/praxis_submission_document_qa')
SHA = lambda path: hashlib.sha256(path.read_bytes()).hexdigest()


def write_text(path, text):
    path.write_text(text, encoding='utf-8')


def assemble():
    original = ORIGINAL / 'manuscript.md'
    expected = '7800012a34f9d716b1742c89e10058b62c40cc039fafa0c92e26d2a6cee087f2'
    assert SHA(original) == expected, 'Original manuscript changed'
    manifest = json.loads((ORIGINAL / 'PACKAGE_MANIFEST.json').read_text(encoding='utf-8'))
    for name, digest in manifest['files'].items():
        assert SHA(REPO / name) == digest, name
    audit = json.loads((HERE / 'sensitivity/AUDIT.json').read_text(encoding='utf-8'))
    assert audit['status'] == 'PASS' and audit['capture_omission_pairs'] == 180
    assert audit['seed_omission_means'] == 36 and audit['scalar_crosschecks'] == 3816
    repair = json.loads((HERE / 'clean_room/LINK_REPAIR_RECEIPT.json').read_text(encoding='utf-8'))
    assert repair['status'] == 'PASS' and not repair['original_extracted_members_changed']

    text = original.read_text(encoding='utf-8')
    def relocate(match):
        label, target = match.groups()
        if urlsplit(target).scheme or target.startswith('#'):
            return match.group(0)
        path, sep, fragment = target.partition('#')
        moved = os.path.relpath((ORIGINAL / path).resolve(), HERE).replace('\\', '/')
        return f'[{label}]({moved}{sep}{fragment})'
    text = re.sub(r'\[([^\]]+)\]\(([^\s)]+)\)', relocate, text)
    addition = json.loads((HERE / 'sensitivity/PAPER_DATA.json').read_text(encoding='utf-8'))['SENSITIVITY_ADDITION']
    addition = addition.replace('### Retrospective capture and seed sensitivity', '## D.1 Retrospective capture and seed sensitivity', 1)
    appendix = '''
# Appendix D. Reviewer Sensitivity and Public Reproduction

This review edition retains the completed study and adds the September 23, 2026 reviewer checks. The original manuscript and evidence package remain byte-identical. The additional calculations use existing public confusion counts and fixed predictions. No new models or independent attack executions were introduced.

'''
    appendix += addition
    appendix += '''
## D.2 Reproduction in a fresh local environment

The original evidence archive was extracted into a new directory and its public arithmetic verifier was run with Python 3.11.9 and NumPy 2.2.6 in a fresh virtual environment. All 36 paired comparisons and 720 reported metric intervals were reproduced from 66 public aggregate tables and the 2,000-draw resampling plan. All 206 original manifest hashes matched. Python file-open auditing recorded no access outside the allowed extracted files, new output/runtime directories and base Python runtime. This is a fresh local-runtime check, not an operating-system sandbox, another laboratory's replication or a refit from raw traces.

The initial package check exposed one missing relative link: the original README linked to the archive itself, which was not inside that archive. Adding the byte-identical downloaded archive at the linked location allowed the unchanged package verifier to pass all 74 local links. No original extracted member changed. Both the initial failure and repaired pass are retained in the [clean-extraction report](clean_room/REPORT.md). The complete review bundle includes the original archive alongside its extracted contents, so the original verification commands work after extraction.

These checks verify aggregate arithmetic and artifact bindings. They cannot independently reconstruct row identities, validate attack labels, establish successful exfiltration or reproduce private-source model fitting. The original access and inference limits remain applicable.

## D.3 Implications for the contribution and its next empirical boundary

The omission analysis supports a directional finding on the observed campaign: higher aggregate F1 can coexist with lower exfiltration warning recall, including after any one specified capture or fitting seed is removed from the means. Its magnitude is not stable across seeds. Removing seed 8101 from the clean budget-three comparison reduces the mean warning loss to 0.36 percentage points. This variation is part of the result, not evidence for a universal effect size. The analysis also records repeated aggregate signatures rather than counting them as independent evidence.

The completed praxis is a controlled measurement and source-qualification study. No new metric, general detector superiority, novel theorem or first-ever observation is claimed. Current literature already recognizes attack-to-normal errors and evaluation-protocol effects; the contribution remains the specific controlled evidence, complete error-destination accounting and reusable audit procedure described in Section 2.

A bounded intake check considered Sandworm and CAM-LDS as possible next sources. Sandworm was already evaluated in this project and lacks a native exfiltration flow class. CAM-LDS has distinct attack executions but lacks simulated normal-user activity, so it cannot by itself reproduce the benign-workload comparison. These are source-qualification findings, not additional classifier experiments. The [source decision and primary-source links](NEXT_REPLICATION_SOURCE.md) document a concrete log-first CAM-LDS intake and the execution, annotation and benign-background requirements for a broader replication. No additional corpus archive download or fit is counted as completed.

The [defense brief](praxis_defense_brief.md) and [prepared adviser handoff](ADVISER_HANDOFF.md) identify the remaining human review: contribution suitability, institutional requirements and whether independent-execution replication is required for the intended submission. They do not represent committee approval or external peer review.
'''
    write_text(HERE / 'praxis_review_edition.md', text.rstrip() + '\n\n' + appendix.lstrip())

    brief = (HERE / 'DEFENSE_BRIEF.template.md').read_text(encoding='utf-8')
    brief = brief.replace('{{SENSITIVITY_BRIEF}}', '''**The direction survived the planned omissions; its size varied.** The frozen audit computed 180 capture omissions and 36 seed omissions. All nine group means showing higher F1 with lower exfiltration warning recall retained that direction after every single-capture and single-seed omission. For the headline clean budget-three comparison, excluding seed 8101 shrank warning loss from **8.93 to 0.36 percentage points**. The 27 acquisition comparisons contain 24 distinct aggregate signatures; neither number represents independent experiments. The temporal current-feature F1 improvement remained positive in all 15 seed/capture omissions.

**The favorable history result also has a boundary.** Chronological history improved F1 in 14 of 15 individual seed/capture omissions; one decreased by 0.13 score points (0.0013 raw macro-F1). All omissions and their supports remain in the evidence. These are fixed-prediction sensitivity checks, not new model fits or confidence intervals.''')
    brief = brief.replace('{{REPRODUCTION_BRIEF}}', '''**The public arithmetic reproduced in a fresh local Python environment.** All 36 comparisons, 720 interval calculations and 206 original artifact hashes passed. One missing archive link was repaired by including the unchanged original ZIP beside its extracted contents; all 74 original local links then passed. This checks the shared calculations without private prediction files. It is not an external laboratory's replication or validation of the author labels.''')
    assert '{{' not in brief and '}}' not in brief
    write_text(HERE / 'praxis_defense_brief.md', brief)
    return {'original_manuscript_sha256': expected, 'original_manifest_files_preserved': len(manifest['files']),
            'sensitivity_audit_sha256': SHA(HERE / 'sensitivity/AUDIT.json'),
            'sensitivity_publication_sha256': SHA(HERE / 'sensitivity/PUBLICATION.json'),
            'repair_receipt_sha256': SHA(HERE / 'clean_room/LINK_REPAIR_RECEIPT.json')}


def render(source, target_stem, brief=False):
    from docx.shared import Pt
    import fitz
    from PIL import Image, ImageDraw
    spec = importlib.util.spec_from_file_location('praxis_renderer', REPO / 'experiments/apt_benchmark/lateral_protection_experiment/paper/build_document.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    renderer = module.Renderer(source, 'https://github.com/garypagangit/praxis/blob/praxis-next-20260923', None)
    doc = renderer.render()
    doc.core_properties.title = 'Praxis Defense Brief' if brief else 'When Better APT Scores Hide Missed Attack Warnings'
    doc.core_properties.subject = 'Measurement praxis review edition: temporal evaluation and attack warnings'
    doc.styles['Praxis Reference'].paragraph_format.space_after = Pt(5)
    for p in doc.paragraphs:
        if p.text.startswith(('Table ', 'Figure ')):
            p.paragraph_format.keep_with_next = True
        if p.style.name == 'Heading 1' and p.text not in ['Abstract', '1. Introduction', 'References']:
            p.paragraph_format.page_break_before = False
            p.paragraph_format.space_before = Pt(18)
        if p.text == 'Appendix D. Reviewer Sensitivity and Public Reproduction':
            p.paragraph_format.page_break_before = True
    if brief:
        doc.styles['Title'].font.size = Pt(22)
        doc.styles['Normal'].paragraph_format.space_after = Pt(5)
        for p in doc.paragraphs:
            if p.style.name.startswith('Heading'):
                p.paragraph_format.page_break_before = False
    for section in doc.sections:
        p = section.header.paragraphs[0]
        p.text = 'PRAXIS DEFENSE BRIEF  |  APT evaluation' if brief else 'EMPIRICAL PRAXIS  |  APT evaluation and attack warnings'
        for run in p.runs:
            run.font.name, run.font.size = 'Arial', Pt(8)
        for footer in [section.footer, section.first_page_footer]:
            for run in footer.paragraphs[0].runs:
                if 'False alarms' in run.text:
                    run.text = 'APT evaluation and attack warnings  |  '
    target = HERE / (target_stem + '.docx')
    doc.save(target)
    cmd = [r'C:\Program Files\LibreOffice\program\soffice.exe',
           '-env:UserInstallation=file:///C:/w/praxis_submission_lo_profile',
           '--headless', '--convert-to', 'pdf', '--outdir', str(HERE), str(target)]
    conversion = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    if conversion.returncode:
        raise RuntimeError(conversion.stdout + conversion.stderr)
    pdf = target.with_suffix('.pdf')
    assert pdf.exists() and pdf.stat().st_mtime >= target.stat().st_mtime
    qa = QA / target_stem
    qa.mkdir(parents=True, exist_ok=True)
    opened = fitz.open(pdf)
    images, geometry, identical = [], [], []
    for i, page in enumerate(opened):
        image = qa / f'page_{i+1:02}.png'
        pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
        pix.save(image)
        images.append(str(image))
        assert page.get_text().strip(), f'Blank page {i+1}'
        bad = []
        for block in page.get_text('dict')['blocks']:
            if block['type'] != 0:
                continue
            for line in block['lines']:
                x0, y0, x1, y1 = line['bbox']
                if x0 < 24 or x1 > page.rect.width - 24 or y0 < 12 or y1 > page.rect.height - 12:
                    bad.append(line['bbox'])
        geometry.append({'page': i+1, 'out_of_page_text_lines': len(bad), 'text_characters': len(page.get_text())})
        old = Path('C:/w/measurement_praxis_document_qa') / image.name
        if not brief and i < 24 and old.exists():
            previous = Image.open(old).convert('RGB')
            current = Image.open(image).convert('RGB')
            if previous.size == current.size and previous.tobytes() == current.tobytes():
                identical.append(i+1)
    thumbs = []
    for i, path in enumerate(images):
        im = Image.open(path).convert('RGB')
        im.thumbnail((340, 450))
        canvas = Image.new('RGB', (360, 480), '#e8e8e8')
        canvas.paste(im, ((360-im.width)//2, 10))
        ImageDraw.Draw(canvas).text((12,462), str(i+1), fill='black')
        thumbs.append(canvas)
    for start in range(0, len(thumbs), 8):
        sub = thumbs[start:start+8]
        canvas = Image.new('RGB', (1440, 480*((len(sub)+3)//4)), '#cccccc')
        for i, im in enumerate(sub):
            canvas.paste(im, ((i%4)*360, (i//4)*480))
        canvas.save(qa / f'contact_{start//8+1}.png')
    return {'manuscript_sha256': SHA(source), 'docx_sha256': SHA(target), 'pdf_sha256': SHA(pdf),
            'pages': len(images), 'tables': renderer.tables, 'images': renderer.images,
            'page_images': images, 'geometry_checks': geometry, 'visually_reviewed_original_pages_identical': identical,
            'visual_review': 'PENDING', 'new_or_changed_pages': [i+1 for i in range(len(images)) if i+1 not in identical]}


if __name__ == '__main__':
    inputs = assemble()
    receipt = {'utc': datetime.now(timezone.utc).isoformat(), 'source_sha256': SHA(Path(__file__)),
               'inputs': inputs, 'documents': {
                  'review_edition': render(HERE / 'praxis_review_edition.md', 'apt_praxis_review_edition'),
                  'defense_brief': render(HERE / 'praxis_defense_brief.md', 'praxis_defense_brief', brief=True)},
               'new_model_fits': 0, 'new_aws_spend_usd': 0, 'visual_review': 'PENDING'}
    write_text(HERE / 'DOCUMENT_RECEIPT.json', json.dumps(receipt, indent=2) + '\n')
    print(json.dumps({name: {key: record[key] for key in ('pages', 'tables', 'images', 'new_or_changed_pages')}
                      for name, record in receipt['documents'].items()}))
