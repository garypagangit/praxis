import argparse,importlib.util,json,re,shutil,zipfile
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree
ROOT=Path(__file__).resolve().parent
REPO=Path('C:/w/apt_benchmark_20260920')
SRC=REPO/'experiments/praxis_next/gwu_final_20260928'
OUT=ROOT.parents[1]/'output/praxis_committee_review_20260928'
QA=Path('C:/w/praxis_committee_review_20260928_qa')
OUT.mkdir(exist_ok=True);QA.mkdir(exist_ok=True)
md=(SRC/'manuscript.md').read_text(encoding='utf-8')
md=md.replace('a targeted primary-source review completed September 23, 2026, centred','a targeted primary-source review initially completed September 23, 2026 and supplemented on September 28, 2026 with the AIT dataset and its foundational sources, centred')
summary=json.loads((ROOT.parent/'campaign_validation_20260928/SUMMARY.json').read_text())
ds=[d for d in summary['deltas'] if d['execution']!='pooled']
labels=[d['execution'].title()+' / '+str(d['seed']) for d in ds]
fig,axes=plt.subplots(1,2,figsize=(10,4.7),sharey=True,layout='constrained')
colors=['#33556e' if d['execution']=='wilson' else '#a16c25' for d in ds]
for ax,key,title,unit in [(axes[0],'delta_macro_f1','Change in macro-F1','F1 units'),(axes[1],'delta_exfil_warning_pp','Change in exfiltration warning recall','Percentage points')]:
 values=[d[key] for d in ds];ax.barh(range(6),values,color=colors,height=.62)
 ax.axvline(0,color='#444444',lw=1);ax.set_title(title,fontsize=11);ax.set_xlabel(unit,fontsize=10)
 ax.spines[['top','right']].set_visible(False);ax.grid(axis='x',alpha=.18);ax.set_axisbelow(True)
 if key=='delta_macro_f1':ax.set_xlim(-.105,.021)
 else:ax.set_xlim(-.205,.075)
 for i,v in enumerate(values):ax.text(v+(-.002 if v<0 else .002),i,f'{v:+.4f}',ha='right' if v<0 else 'left',va='center',fontsize=9)
axes[0].set_yticks(range(6),labels,fontsize=10);axes[0].invert_yaxis()
fig.suptitle('AIT held-out comparisons: error-focused minus entropy',fontsize=13,fontweight='bold')
for ext in ['png','pdf','svg']:fig.savefig(SRC/'figures'/('ait_paired_deltas.'+ext),dpi=240)
plt.close(fig)
caption='Figure 4-9. AIT paired deltas for all six execution-by-seed comparisons. Both panels show error-focused minus entropy selection on matched rows. The frozen original direction requires positive macro-F1 change together with negative warning-recall change; none of the six comparisons meets both conditions. Wilson shows a small opposite-order tradeoff, while Harrison favors entropy on both outcomes. The panels use different horizontal units; fitting seeds are repeated fits, not independent campaigns.'
anchor='Giving each execution equal weight yields mean macro-F1'
if '![AIT paired deltas]' not in md:md=md.replace(anchor,'![AIT paired deltas](figures/ait_paired_deltas.png)\n\n'+caption+'\n\n'+anchor)
(SRC/'manuscript.md').write_text(md,encoding='utf-8')
spec=importlib.util.spec_from_file_location('gwu_renderer',REPO/'experiments/praxis_next/gwu_final_20260924/render_gwu.py')
mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
linear=[
'F_k(x) = F_0,k + eta * sum(m=1..M) h_mk(x).',
'p_k(x) = exp(F_k(x)) / sum(j=0..3) exp(F_j(x)).',
'L_CE = -sum(i=1..n) log(p_y_i(x_i)).',
'Predicted y_i = argmax(k in {0,1,2,3}) p_k(x_i).',
'loss_w(y,p) = w_y * indicator(argmax(k) p_k differs from y).',
'w = (1, 1, 4, 4).',
't_i = loss_w(y_i,p_i^H) - loss_w(y_i,p_i^C).',
'p_i^gate = p_i^H when g(z_i) < 0.',
'p_i^gate = p_i^C when g(z_i) >= 0.',
'u_i^harm(S,a) = loss_w(y_i,p_i^S) - loss_w(y_i,p_i^(S union {a})).',
'u_i^entropy(S,a) = H(p_i^S) - H(p_i^(S union {a})).',
'a* = argmax(a in E(S)) [q_a * predicted_u(S,a) / c_a].',
'Request only when q_a* * predicted_u(S,a*) / c_a* > 0.',
'a is not in A_attempted.',
'C_spent + c_a <= B.',
'T_elapsed + d_a^nominal <= D.',
'p(T1105 | x) = 1 / (1 + exp(-(b + beta^T * phi(x)))).',
'(b_hat,beta_hat) = argmin(b,beta) [sum(i=1..n) (t_i - b - z_i^T * beta)^2 + 10 * ||beta||_2^2].',
'P_s = TP_s / (TP_s + FP_s).',
'R_s = TP_s / (TP_s + FN_s).',
'F1_s = 2*TP_s / (2*TP_s + FP_s + FN_s).',
'F1_macro = sum(s=1..K) F1_s / K.',
'R_stage,s = C_s,s / N_s.',
'R_warning,s = 1 - C_s,b / N_s.',
'FPR_b = sum(j differs from b) C_b,j / N_b.'
]
symbols=['x; y: observed features; declared evaluation class','p; F_k: class-probability vector; additive class score','h_mk; M; eta: tree contribution; boosting iterations; learning rate',
'g; z; t: fitted selector; selector features; regression response','w; loss_w: declared stage costs; weighted hard-error loss','u; predicted_u: realized acquisition gain; fitted gain estimate','c_a; q_a: group-a request cost; supplied availability prior','H(p): predictive entropy; superscript H denotes the history expert','S; a; E(S): delivered evidence set; requested group; eligible groups','A_attempted: set of previously attempted evidence groups','B; D: simulated evidence budget; decision deadline','C_spent; T_elapsed; d_a: spent cost; elapsed time; request delay','b; beta; phi(x): intercept; coefficient vector; fixed feature representation','L_CE; n: multiclass cross-entropy data-fit loss; fitting-row count','C(s,j); N_s: confusion count; support for true class s','P_s; R_s; F1_s: class precision, recall and F1','TP; FP; FN: true-positive, false-positive and false-negative counts','K; F1_macro: declared class count; mean class F1','R_stage,s; R_warning,s: exact-stage recall; stage-conditioned warning recall','FPR_b: false-positive rate among benign records (class b)']
acronyms=['AIT: Austrian Institute of Technology','APT: Advanced Persistent Threat','AUC: Area Under the Curve','CRC: Cyclic Redundancy Check','DNS: Domain Name System','EDA: Exploratory Data Analysis','FPR: False-Positive Rate','GML: Graph Machine Learning','GMR: Graphical Model of Research','IDS: Intrusion Detection System','MD5: Message-Digest Algorithm 5','ML: Machine Learning','OOF: Out-of-Fold','ROC: Receiver Operating Characteristic','RQ: Research Question','SHA / SHA-256: Secure Hash Algorithm / 256-bit variant','SVD: Singular Value Decomposition','T1105: MITRE ATT&CK identifier for Ingress Tool Transfer','TCP: Transmission Control Protocol','UDP: User Datagram Protocol']
parent=mod.GWURenderer
class CorrectedRenderer(parent):
 def frontmatter(self):
  super().frontmatter()
  replacements={'Manuscript prepared September 24, 2026':'Manuscript prepared September 28, 2026','A Praxis submitted to':'A Praxis prepared for review by','Praxis direction and committee confirmation pending':'Committee review edition','Certification pending':'Review Status','The confirmed director and committee, examination date, and approved certification wording must be supplied through the university process before formal submission.':'Formal submission requires the university-confirmed director and committee, examination date and approved certification wording. This review copy does not supply or imply those approvals.'}
  for p in self.doc.paragraphs:
   for r in p.runs:
    for a,b in replacements.items():r.text=r.text.replace(a,b)
  ps=self.doc.paragraphs
  si=next(i for i,p in enumerate(ps) if p.text=='List of Symbols');ai=next(i for i,p in enumerate(ps) if p.text=='List of Acronyms')
  for heading,entries,remove in [(ps[si],symbols,ps[si+1:ai]),(ps[ai],acronyms,ps[ai+1:-1])]:
   for p in remove:p._element.getparent().remove(p._element)
   previous=heading._p
   for text in entries:
    p=self.doc.add_paragraph(text,style='GWU Front');p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1;p.paragraph_format.first_line_indent=Pt(0)
    previous.addnext(p._p);previous=p._p
 def equation(self,text):
  start=len(self.equations);before=len(self.doc.paragraphs)
  super().equation(text)
  equation_ps=self.doc.paragraphs[before:]
  for i,p in enumerate(equation_ps,start):
   p.paragraph_format.keep_with_next=True
   plain=self.doc.add_paragraph('Linear notation: '+linear[i]);plain.paragraph_format.first_line_indent=Pt(0);plain.paragraph_format.line_spacing=1;plain.paragraph_format.space_after=Pt(10)
   for r in plain.runs:r.font.size=Pt(10)
   p._p.addnext(plain._p)
mod.GWURenderer=CorrectedRenderer
mod.build(argparse.Namespace(source=SRC/'manuscript.md',abstract=SRC/'abstract.md',title=mod.TITLE,docx=OUT/'Gary_Pagan_Final_Praxis.docx',pdf=OUT/'Gary_Pagan_Final_Praxis.pdf',qa_dir=QA,receipt=ROOT/'COMMITTEE_RENDER_RECEIPT.json'))
# Cached PAGE results in footer parts were the exact reported trailing sequence.
# Keep live PAGE fields, clear their stale caches, and request field refresh on open.
docx=OUT/'Gary_Pagan_Final_Praxis.docx';ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(docx) as z:parts={n:z.read(n) for n in z.namelist()}
removed=[]
for name,data in list(parts.items()):
 if name.startswith('word/footer') and name.endswith('.xml'):
  root=etree.fromstring(data)
  if any('PAGE' in x for x in root.xpath('.//w:instrText/text()',namespaces=ns)):
   for t in root.findall('.//w:t',ns):
    if t.text and re.fullmatch(r'[ivxlcdm]+|\d+',t.text.strip()):removed.append(t.text);t.text=''
   for f in root.findall('.//w:fldChar',ns):
    if f.get(qn('w:fldCharType'))=='begin':f.set(qn('w:dirty'),'true')
   parts[name]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
settings=etree.fromstring(parts['word/settings.xml']);field=settings.find('w:updateFields',ns)
if field is None:field=etree.SubElement(settings,qn('w:updateFields'))
field.set(qn('w:val'),'true');parts['word/settings.xml']=etree.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(docx,'w',zipfile.ZIP_DEFLATED) as z:
 for n,data in parts.items():z.writestr(n,data)
(ROOT/'COMMITTEE_CORRECTIONS.json').write_text(json.dumps({'footer_page_caches_removed':removed,'live_page_fields_retained':True,'equations_with_linear_notation':len(linear),'ait_figure':'4-9','formal_submission':'Confirmed committee, director, examination date and approved certification wording still required; no approvals invented.'},indent=2))
shutil.copy2(SRC/'manuscript.md',OUT/'manuscript.md');shutil.copy2(SRC/'abstract.md',OUT/'abstract.md')
print('Corrected manuscript rendered; footer caches cleared.')
