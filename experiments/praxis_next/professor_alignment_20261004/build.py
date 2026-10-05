from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
import json

ROOT = Path(__file__).parent
OUT = Path('C:/Users/garyp/OneDrive/Documents/codex/output/doc/praxis_professor_alignment_20261004')
OUT.mkdir(parents=True, exist_ok=True)
d = Document()
s = d.sections[0]
s.top_margin = s.bottom_margin = Inches(.8)
s.left_margin = s.right_margin = Inches(.9)
s.page_width, s.page_height = Inches(8.5), Inches(11)
n = d.styles['Normal']; n.font.name = 'Calibri'; n.font.size = Pt(11)
n.paragraph_format.space_after = Pt(5)
records=[]
def h(t): d.add_heading(t, 1); records.append(t)
def p(t): d.add_paragraph(t); records.append(t)
def page(): d.add_page_break()

h('Praxis alignment for adviser review')
p('Gary Pagan | October 4, 2026 | Proposed replacement framing; completed results remain unchanged.')
h('Recommended title')
p('Auditing and Preserving Exfiltration Warnings in a LightGBM APT Detection Pipeline')
h('Problem statement: issue and consequence')
p('In the examined LightGBM APT stage-classification pipeline, a change in how the system selects available evidence increased macro-F1 while reducing the proportion of exfiltration-labeled flows that received any attack warning. Accepting this change on the overall score alone could therefore remove opportunities to investigate suspected data theft. The financial stakes are substantial: IBM (2026) reports a global average data-breach cost of US$4.99 million. Model review needs to measure retained attack warnings alongside classification scores and the additional review workload of proposed repairs.')
p('The first sentence is supported by our PX-081/PX-106 evidence, not attributed to Uddin et al. The IBM estimate supplies industry context; it is not an APT-specific loss estimate, a cost per missed flow, or a demonstrated saving from this study.')
h('Purpose and thesis')
p('Purpose: develop and evaluate a reproducible review procedure that identifies lost exfiltration warnings, distinguishes warnings discarded by the pipeline from attacks missed by all available models, and measures the benefits and costs of retaining available warnings.')
p('Thesis supported by completed evidence: paired warning accounting and recorded decision traces expose losses that an aggregate score alone does not describe. An OR gate retains every warning produced by its available members, but its detection gain and false-alert cost depend on the data and members. It cannot recover an attack that every member labels benign.')
h('Two wording corrections')
p('F1 does not charge all errors the same price. It combines precision and recall without explicitly representing the operational difference between a wrong attack-stage name and no warning. Also, a model warning is an opportunity for review; this study has not measured whether an analyst actually reviewed it.')

page(); h('Aligned questions and testable claims')
p('These organize completed studies. They are retrospective claims, not newly preregistered hypotheses. Failed tests remain visible; no threshold should be selected after looking at the test results.')
h('RQ1: When does a better score conceal fewer warnings?')
p('Under clean evidence and budget three, does error-focused evidence selection increase macro-F1 while decreasing exfiltration warning recall compared with entropy-based selection, using the same LightGBM experts and test records?')
p('H1: the three-seed mean change in macro-F1 is positive and the mean change in exfiltration warning recall is negative. Completed result: consistent with H1 on UNRAVELED; F1 rises from 0.7148 to 0.7379, while warning recall falls from 85.18% to 76.25%. This is descriptive evidence from one campaign, with substantial seed sensitivity. The original ordering did not replicate on AIT.')
h('RQ2: How much warning recovery does OR provide, and at what cost?')
p('Does the three-member OR gate recover more exfiltration warnings than probability averaging and the best individual member, on UNRAVELED and the separate AIT Wilson and Harrison executions? Count benign false alerts and grouped review cases alongside recall.')
p('H2: OR has strictly higher exfiltration warning recall than the best individual member in both AIT executions. Completed result: not supported. Wilson ties the best member. OR does improve on averaging, including 98.78% versus 68.04% on UNRAVELED, with 196 versus 182 benign false alerts. These are flow counts, not analyst hours or independent incidents.')
h('RQ3: Does added explanation improve diagnosis beyond a trace?')
p('On the predefined pipeline-fault scenarios, does replay-supported explanation diagnose the warning-loss cause more accurately than an ordinary decision trace?')
p('H3: replay-supported explanation has higher diagnosis accuracy than the trace baseline. Completed result: not supported. Both diagnose 170/170 scenarios and classify 510/510 repair-feasibility checks correctly in PX-105. The bot study verifies predefined evidence checks; it does not establish human analyst benefit.')
h('What this means for the proposal')
p('The completed evidence supports an applied audit and a conditional repair. It does not support a superior new explanation algorithm or a repair that consistently beats the strongest model. Hypotheses need not all succeed, but the claimed contribution must fit these outcomes.')

page(); h('Methodology in everyday language')
p('1. Start with recorded network flows and the dataset authors\' attack-stage labels. Use UNRAVELED for the examined campaign and adapted AIT Wilson/Harrison for external checks. Keep executions separate; many flows from one execution are not many independent attacks.')
p('2. Train LightGBM classifiers. LightGBM combines many small decision trees: each tree asks questions about traffic measurements, and their combined scores select benign, other attack, movement, or exfiltration. Evidence variants use current traffic measurements, host roles, and/or historical summaries.')
p('3. Prove the score-warning mismatch by replaying two evidence-selection policies on exactly the same test rows. Hold the fitted experts fixed. Count true exfiltration rows that change from any attack label to benign, including those previously given the wrong attack-stage name. Report gains separately from losses.')
p('4. Test the repair using three separately fitted current-traffic-plus-host-role LightGBM models, with seeds 8101, 8102 and 8103. These are different fits of the same architecture, not three independent kinds of evidence. Compare individual models, averaged probabilities and OR on the same rows.')
p('5. Check the recorded decisions. Identify whether a warning existed but was discarded, or whether every available model said benign. Replay only explicitly specified changes. This explains the pipeline decision; it does not prove what caused the attack.')
p('6. Report per-seed results, execution-level results and paired uncertainty using capture groups where supported. Fit thresholds and choose methods on training/calibration data. Already examined test data cannot become a fresh confirmation set. A new claim needs untouched executions and a frozen protocol.')
h('Simple repair diagram')
diagram='''Same network flow
        |
        +--> LightGBM fit A -- attack or benign --+
        +--> LightGBM fit B -- attack or benign --+--> OR gate
        +--> LightGBM fit C -- attack or benign --+       |
                                                       v
                              Any attack vote? Yes: keep warning
                                               No: no warning
                                                       |
                              Save member votes + final decision
                              Count recovered attacks + false alerts'''
q=d.add_paragraph(); q.paragraph_format.space_after=Pt(4)
r=q.add_run(diagram); r.font.name='Consolas'; r.font.size=Pt(8.5)
records.append(diagram)
p('This diagram is the repair comparison. The original score-loss experiment separately compares evidence-selection policies over fixed experts.')
h('Metrics that answer the questions')
p('Macro-F1: overall class performance. Exact-stage recall: exfiltration correctly named. Warning recall: exfiltration given any attack label. Warning losses/gains: paired changes in the same rows. Benign false alerts and grouped cases: workload proxies. Episode coverage: labeled episode proxies with at least one warning. Diagnosis accuracy: correct identification of the injected pipeline fault. No dollar savings, analyst time or prevented breaches were measured.')

page(); h('Novelty assessment and the next decision')
p('Novelty remains unresolved. OR combination, prediction regressions, false-negative diagnosis and SHAP-based intrusion explanations all have prior work. Merely adding explanations, an LLM narrator, or APT data does not establish a new method. The mathematical OR guarantees are established set-union properties; the empirical contribution concerns where the tradeoff is useful.')
p('The strongest current contribution candidate is the combined, reproducible evidence: wrong-stage warnings can disappear into benign predictions; available-warning losses and all-model misses require different repairs; and gains must be reported with false alerts and episode coverage. A targeted search cannot prove that this combination is first. Adviser acceptance of an applied contribution remains necessary.')
h('If a stronger XAI contribution is required')
p('Investigate whether a verifiable explanation helps choose a useful repair under a fixed review-case budget on unseen executions. For each proposed warning, show the supporting model, observed traffic evidence, recorded delivery/selection decision, and the exact change that would restore the warning. If all members are silent, explicitly report that the explanation cannot establish a recoverable warning.')
p('Compare the proposed method against ordinary trace diagnosis, confidence ranking, unconditional OR, a cost-sensitive decision rule and a binary attack detector. Give every method the same case budget and available information. The primary outcome should be additional labeled exfiltration episodes covered at that budget; secondary outcomes should include benign cases, explanation correctness and stability. Freeze the episode definition, budget and success criterion before evaluation.')
p('Proceed only if the proposed explanation changes a decision and beats these baselines on untouched executions. PX-103 already tied confidence ranking and PX-105 tied ordinary traces, so those implementations have not passed this test. Human usefulness would require an approved human study; a bot can test consistency but cannot stand in for SOC operators.')
h('Literature access and evidence limits')
p('GWU Libraries discovery was used on October 4: the CyberShapley title search returned one peer-reviewed Computers & Security record. The catalogue requested sign-in for complete results; authenticated subscription full-text access was not established. CyberShapley and AlertPro full-text review used the previously supplied local PDFs. Additional primary papers and IBM\'s report landing page were checked openly. This was a targeted overlap review, not a systematic literature review.')
p('No new fits or AWS jobs were needed for this alignment. Existing PX-081, PX-092, PX-103, PX-105 and PX-106 findings are summarized; results and protocols were not rewritten. This brief is an adviser-review supplement, not a replacement for the complete manuscript.')

page(); h('Selected references and their role')
refs=[
('IBM. (2026). Cost of a data breach report 2026. https://www.ibm.com/reports/data-breach','Financial context only: US$4.99 million global average.'),
('Uddin, M. A., Aryal, S., Bouadjenek, M. R., Al-Hawawreh, M., & Talukder, M. A. (2025). Hierarchical classification for intrusion detection system: Effective design and empirical analysis. Ad Hoc Networks, 178, 103982. https://doi.org/10.1016/j.adhoc.2025.103982','Related binary/stage classification design; not the source of our LightGBM score-loss result.'),
('Kittler, J., Hatef, M., Duin, R. P. W., & Matas, J. (1998). On combining classifiers. IEEE Transactions on Pattern Analysis and Machine Intelligence, 20(3), 226-239. https://doi.org/10.1109/34.667881','Classifier combination is established.'),
('Malach, A., Wudali, P. N., Momiyama, S., Furukawa, J., Araki, T., Elovici, Y., & Shabtai, A. (2025). CyberShapley: Explanation, prioritization, and triage of cybersecurity alerts using informative graph representation. Computers & Security, 150, 104270. https://doi.org/10.1016/j.cose.2024.104270','Prior explainable cybersecurity alert triage, including APT datasets.'),
('Wang, X., Yang, X., Liang, X., Zhang, X., Zhang, W., & Gong, X. (2024). Combating alert fatigue with AlertPro: Context-aware alert prioritization using reinforcement learning for multi-step attack detection. Computers & Security, 137, 103583. https://doi.org/10.1016/j.cose.2023.103583','Prior alert prioritization; its reported comparisons already include higher F1 with lower attack recall.'),
('Mia, M., Pritom, M. M. A., Islam, T., & Hasan, K. (2024). Visually analyze SHAP plots to diagnose misclassifications in ML-based intrusion detection [Preprint]. arXiv. https://arxiv.org/abs/2411.02670','Prior SHAP-based diagnosis of false positives and false negatives; inspect before claiming a new XAI contribution.'),
('Ficke, E., Schweitzer, K. M., Bateman, R. M., & Xu, S. (2019). Analyzing root causes of intrusion detection false-negatives: Methodology and case study [Preprint]. arXiv. https://arxiv.org/abs/1909.08725','Prior systematic false-negative root-cause analysis, demonstrated with Snort.'),
]
for ref,role in refs: p(ref); p('Role: '+role)
d.save(OUT/'Praxis_Professor_Alignment.docx')
(ROOT/'CONTENT.txt').write_text('\n\n'.join(records),encoding='utf-8')
(ROOT/'RESEARCH_ACCESS.json').write_text(json.dumps({'date':'2026-10-04','gwu_query':'CyberShapley','gwu_results':1,'peer_reviewed_record':True,'authenticated_full_text':False,'new_experiments':False,'sources':[x[0] for x in refs]},indent=2),encoding='utf-8')
print(OUT)
