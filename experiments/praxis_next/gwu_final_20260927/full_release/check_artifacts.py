from pathlib import Path
import json, zipfile, hashlib
import pymupdf
from PIL import Image, ImageChops
from pptx import Presentation
ROOT=Path(__file__).parent
OUT=ROOT.parents[1]/'output/praxis_final_20260927'
QA=Path('C:/w/gwu_final_document_qa_20260927')
changed=[]
for start in range(1,137,4):
 panel=Image.open(QA/'full_review'/f'pages_{start:03d}.png').convert('RGB')
 for k in range(4):
  n=start+k
  page=Image.open(QA/f'page_{n:03d}.png').convert('RGB')
  x=(k%2)*(panel.width//2);y=(k//2)*(panel.height//2)+24
  crop=panel.crop((x,y,x+page.width,y+page.height))
  if ImageChops.difference(crop,page).getbbox():changed.append(n)
paper=pymupdf.open(OUT/'Gary_Pagan_Final_Praxis.pdf')
deck=pymupdf.open(OUT/'Gary_Pagan_GWU_Praxis_Defense.pdf')
pres=Presentation(OUT/'Gary_Pagan_GWU_Praxis_Defense.pptx')
content=json.loads((ROOT/'DEFENSE_CONTENT.json').read_text(encoding='utf-8'))
if isinstance(content,dict):content=content['slides']
assert len(paper)==136 and len(deck)==len(pres.slides)==len(content)==32
titles=[]
for i,(slide,c) in enumerate(zip(pres.slides,content)):
 texts=[s.text for s in slide.shapes if s.has_text_frame]
 assert c['title'] in texts,(i,c['title'],texts)
 assert slide.has_notes_slide and len(slide.notes_slide.notes_text_frame.text)>100
 assert ''.join(c['title'].split()) in ''.join(deck[i].get_text().split()),i
 titles.append(c['title'])
for name in ['Gary_Pagan_Final_Praxis.docx','Gary_Pagan_GWU_Praxis_Defense.pptx']:
 with zipfile.ZipFile(OUT/name) as z:assert z.testzip() is None
text='\n'.join(p.get_text() for p in paper)
assert '300 boosting iterations' in text and '300 trees' not in text
assert 'This completion adds reanalysis' not in text
receipt=json.loads((ROOT/'RENDER_RECEIPT.json').read_text())
assert not receipt['marker_tokens_remaining']
outside=[]
for i,p in enumerate(deck):
 for b in p.get_text('blocks'):
  if b[0]<-1 or b[1]<-1 or b[2]>p.rect.width+1 or b[3]>p.rect.height+1:outside.append(i+1)
assert not outside,outside
for start in range(0,32,4):
 imgs=[]
 for j in range(start,start+4):
  pix=deck[j].get_pixmap(dpi=96);imgs.append(Image.frombytes('RGB',[pix.width,pix.height],pix.samples))
 w=max(x.width for x in imgs);h=max(x.height for x in imgs)
 panel=Image.new('RGB',(w*2,h*2),'white')
 for k,im in enumerate(imgs):panel.paste(im,((k%2)*w,(k//2)*h))
 panel.save(QA/f'deck_pdf_{start+1:02d}.png')
result={'status':'PASS','paper_pages':len(paper),'slides':len(deck),'speaker_notes':len(pres.slides),'changed_pages_since_visual_panels':changed,'slide_titles':titles,'outside_slide_text':outside}
(ROOT/'STRUCTURAL_QA.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result,indent=2))
