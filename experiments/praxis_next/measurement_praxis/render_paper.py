"""Build Word/PDF and local page images for visual review; no scientific fitting."""
from pathlib import Path
import importlib.util
import hashlib
import json
import subprocess
from docx.shared import Pt

HERE=Path(__file__).resolve().parent
REPO=HERE.parents[2]
QA=Path('C:/w/measurement_praxis_document_qa')
spec=importlib.util.spec_from_file_location('praxis_renderer',REPO/'experiments/apt_benchmark/lateral_protection_experiment/paper/build_document.py')
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
renderer=module.Renderer(HERE/'manuscript.md','https://github.com/garypagangit/praxis/blob/praxis-next-20260923',None)
doc=renderer.render()
doc.core_properties.title='When Better APT Scores Hide Missed Attack Warnings'
doc.core_properties.subject='Temporal evaluation, historical evidence, attack-warning retention and benchmark support'
doc.styles['Praxis Reference'].paragraph_format.space_after=Pt(5)
for p in doc.paragraphs:
    if p.text.startswith(('Table ', 'Figure ')):p.paragraph_format.keep_with_next=True
    if p.style.name=='Heading 1' and p.text not in ['Abstract','1. Introduction','References']:
        p.paragraph_format.page_break_before=False
        p.paragraph_format.space_before=Pt(18)
for section in doc.sections:
    p=section.header.paragraphs[0];p.text='EMPIRICAL PRAXIS  |  APT evaluation and attack warnings'
    for r in p.runs:r.font.name='Arial';r.font.size=Pt(8)
    for footer in [section.footer,section.first_page_footer]:
        for r in footer.paragraphs[0].runs:
            if 'False alarms' in r.text:r.text='APT evaluation and attack warnings  |  '
target=HERE/'apt_evaluation_praxis.docx';doc.save(target)
QA.mkdir(parents=True,exist_ok=True)
cmd=[r'C:\Program Files\LibreOffice\program\soffice.exe','-env:UserInstallation=file:///C:/w/measurement_praxis_lo_profile',
     '--headless','--convert-to','pdf','--outdir',str(HERE),str(target)]
p=subprocess.run(cmd,capture_output=True,text=True,timeout=120)
if p.returncode:raise RuntimeError(p.stdout+p.stderr)
pdf=target.with_suffix('.pdf')
if not pdf.exists() or pdf.stat().st_mtime<target.stat().st_mtime:raise RuntimeError('Missing or stale PDF')
import fitz
from PIL import Image,ImageDraw
opened=fitz.open(pdf);paths=[];geometry=[]
for i,page in enumerate(opened):
    target_page=QA/f'page_{i+1:02}.png';page.get_pixmap(matrix=fitz.Matrix(1.6,1.6),alpha=False).save(target_page);paths.append(str(target_page))
    if not page.get_text().strip():raise ValueError('Blank page '+str(i+1))
    bad=[]
    for block in page.get_text('dict')['blocks']:
        if block['type']!=0:continue
        for line in block['lines']:
            x0,y0,x1,y1=line['bbox']
            if x0<24 or x1>page.rect.width-24 or y0<12 or y1>page.rect.height-12:bad.append(line['bbox'])
    geometry.append({'page':i+1,'out_of_page_text_lines':len(bad),'text_characters':len(page.get_text())})
thumbs=[]
for i,path in enumerate(paths):
    im=Image.open(path).convert('RGB');im.thumbnail((340,450));canvas=Image.new('RGB',(360,480),'#e8e8e8')
    canvas.paste(im,((360-im.width)//2,10));ImageDraw.Draw(canvas).text((12,462),str(i+1),fill='black');thumbs.append(canvas)
for start in range(0,len(thumbs),8):
    sub=thumbs[start:start+8];canvas=Image.new('RGB',(1440,480*((len(sub)+3)//4)),'#cccccc')
    for i,im in enumerate(sub):canvas.paste(im,((i%4)*360,(i//4)*480))
    canvas.save(QA/f'contact_{start//8+1}.png')
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
receipt={'manuscript_sha256':sha(HERE/'manuscript.md'),'docx_sha256':sha(HERE/'apt_evaluation_praxis.docx'),
         'pdf_sha256':sha(pdf),'render_source_sha256':sha(Path(__file__)),'pages':len(paths),'tables':renderer.tables,
         'images':renderer.images,'hyperlinks':renderer.links,'local_only_links':renderer.local_only_links,
         'geometry_checks':geometry,'visual_review':'PENDING','page_images':paths}
(HERE/'RENDER_RECEIPT.json').write_text(json.dumps(receipt,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'pages':len(paths),'tables':renderer.tables,'images':renderer.images,'out_of_page_text_lines':sum(g['out_of_page_text_lines'] for g in geometry)}))
