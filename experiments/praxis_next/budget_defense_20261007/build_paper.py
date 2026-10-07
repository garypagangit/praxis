"""Generate the adviser-review manuscript from completed experiment results."""
import json,pathlib
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
HERE=pathlib.Path(__file__).resolve().parent; ROOT=HERE.parents[2]
OUT=ROOT/'output/doc/budget_defense_20261007'; OUT.mkdir(parents=True,exist_ok=True)
S=json.loads((HERE/'SUMMARY.json').read_text()); R=json.loads((HERE/'RESULTS.json').read_text())['rows']
TITLE='Preserving Attack Evidence Under Missing Telemetry and Fixed Investigation Budgets'
d=Document(); sec=d.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(1.25);sec.top_margin=sec.bottom_margin=Inches(1);sec.footer_distance=Inches(.5)
normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.paragraph_format.line_spacing=2;normal.paragraph_format.space_after=Pt(0)
for st in ['Heading 1','Heading 2','Heading 3']:
    d.styles[st].font.name='Times New Roman';d.styles[st].font.size=Pt(12);d.styles[st].font.color.rgb=None
for st in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
    style=d.styles[st];style.font.name='Times New Roman';style.font.size=Pt(12);style.font.color.rgb=RGBColor(0,0,0)
    rf=style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+key),None)
    rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')
for style in d.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
def numbering(s,fmt,start):
    s.footer.is_linked_to_previous=False
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
    n=OxmlElement('w:pgNumType');n.set(qn('w:fmt'),fmt);n.set(qn('w:start'),str(start));s._sectPr.append(n)
numbering(sec,'lowerRoman',1);sec.different_first_page_header_footer=True
def p(t): return d.add_paragraph(t)
def h(t): d.add_heading(t,level=2)
def chapter(t):
    pp=d.add_heading(t,level=1);pp.paragraph_format.page_break_before=True
def table(headers,rows):
    t=d.add_table(rows=1, cols=len(headers))
    for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
    for row in rows:
        for c,v in zip(t.add_row().cells,row):c.text=str(v)
    for row in t.rows:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_after=Pt(5)
                for run in pp.runs:run.font.size=Pt(10)
    return t
pp=p(TITLE);pp.paragraph_format.space_before=Inches(1)
for txt in ['by Gary Pagan','B.S. in Applied Math and Statistics, May 1999, Stony Brook University','M.S. in Engineering Management, May 2006, The George Washington University','A Praxis research draft submitted for adviser review','The Faculty of\nThe School of Engineering and Applied Science\nof The George Washington University','Doctor of Engineering','October 7, 2026','Research draft: degree-conferral date and committee approval are not asserted.']:
    pp=p(txt);pp.paragraph_format.space_before=Pt(14)
for pp in d.paragraphs:pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.line_spacing=1
d.add_page_break();d.add_heading('Manuscript Status',0)
p('This manuscript reports a completed exploratory experiment and is prepared for academic review. It does not certify a passed final examination, an approved committee, or institutional acceptance. Director, committee and conferral details must be confirmed through the university process before formal submission. Earlier degree details follow the existing project manuscript.')
p('The body follows GW university conventions for type, spacing, margins and pagination. The current GW doctoral policies page links a revised 2026 Praxis template, but that Box attachment was not retrievable during preparation. Therefore this is a GWU-style draft, not a claim of complete conformance to the revised template. The research conclusions and all negative results are retained for adviser review.')
d.add_page_break();d.add_heading('Abstract of Praxis',0);p(TITLE)
p('Missing security telemetry can increase threshold-based attack recall while reducing the useful evidence available within a fixed investigation budget. This study tests that problem and evaluates simple score-fusion mitigations using existing CADETS and THEIA provenance-graph score archives. Ten policies are compared across three fitting seeds, two observation conditions and five node-review budgets, yielding 600 comparisons without new data collection or neural-model fitting. At the primary 1% budget, mean GIN/KNN attack-node precision falls from 36.17% to 5.51% on CADETS and from 81.75% to 22.56% on THEIA after simulated 50% edge loss. Combining GIN and local-score percentiles raises THEIA loss-condition precision to 63.61%, but lowers clean precision to 75.63% and fails on CADETS. None of seven fusion policies meets the frozen requirement of improved loss-condition precision in both datasets with at most a five-percentage-point clean-condition penalty. Tie-aware accounting prevents arbitrary node ordering from being mistaken for detector utility. The contribution is an auditable engineering evaluation and a bounded negative finding about simple mitigation, rather than a new detector or demonstrated operational benefit. Previously inspected static graphs, one loss mask and benchmark labels limit generalization. Independent executions and audited entity mappings remain necessary before a defense deployment or originality claim.')
d.add_page_break();d.add_heading('Table of Contents',0)
entries=json.loads((HERE/'CONTENTS.json').read_text()) if (HERE/'CONTENTS.json').exists() else {'Chapter 1: Introduction':1,'Chapter 2: Related Work and Research Position':3,'Chapter 3: Methodology':5,'Chapter 4: Results':9,'Chapter 5: Discussion and Engineering Implications':12,'Chapter 6: Conclusions':15,'References':16,'Appendix A: Reproduction and Evidence':17,'Appendix B: Claim-to-Evidence Map':18}
for title,page in entries.items():p(title+' ... '+str(page))
d.add_page_break();d.add_heading('List of Tables and Figures',0)
for t in ['Table 1. Evaluation populations and primary budgets (Chapter 3).','Table 2. Primary-budget precision for all policies (Chapter 4).','Table 3. Primary-policy descriptive seed ranges (Chapter 4).','Figure 1. Budget sensitivity under 50% edge loss (Chapter 4).'] :p(t)
d.add_heading('List of Abbreviations',0)
for t in ['APT: advanced persistent threat','GIN: graph isomorphism network','KNN: nearest-neighbor anomaly scoring','MLP: multilayer perceptron','SOC: security operations center','TP: benchmark-labeled true-positive node']:p(t)
s=d.add_section(WD_SECTION_START.NEW_PAGE);s.different_first_page_header_footer=False;numbering(s,'decimal',1)
d.add_heading('Chapter 1: Introduction',1)
h('1.1 Engineering problem')
p('A defender has limited capacity to investigate suspicious activity. A detector may generate more alerts when logging deteriorates, yet place fewer attack-relevant entities near the top of the queue. The engineering problem is therefore not simply whether an anomaly score remains high. It is whether the available review capacity still reaches useful evidence when telemetry is incomplete. Accepting a detector update using threshold recall alone can hide this loss of practical ranking value.')
p('For example, an investigation queue with room for 100 entities cannot use all 10,000 alerts generated by a degraded detector. A higher count of labeled attack entities somewhere in those alerts does not establish that the first 100 are informative. This example is illustrative. In the experiment, the review unit is a graph node and the budget is a declared fraction of all nodes; it is not a measured number of analyst cases or hours.')
h('1.2 Purpose and scope')
p('The purpose is to build and evaluate a reproducible procedure for comparing detector evidence at an equal review budget under static telemetry loss, and to test whether simple label-free combinations of available scores mitigate the observed failure. Existing score archives make this possible without rerunning attack collection. The engineering artifact comprises a frozen protocol, scoring code, tie-aware metrics, full result inventory and independent arithmetic checks.')
p('This direction was selected after a broader project screen of model fingerprinting, incomplete logs, recovery behavior, unknown-model rejection, early decisions and human-versus-AI signals. Those studies exposed transfer and observation limitations. The present manuscript concentrates on direct defense. It does not infer whether an attacker is human or AI and does not establish an attribution signature.')
h('1.3 Research questions and success criterion')
p('RQ1: At a fixed node-review budget, how does simulated edge loss change attack-evidence yield for the existing GIN/KNN detector? RQ2: Can fixed combinations of GIN, MLP and local anomaly scores improve loss-condition precision in both datasets without an unacceptable clean-condition penalty? RQ3: How sensitive are those conclusions to review budget, fitting seed and score ties?')
p('For RQ2, the primary criterion requires higher mean precision than GIN under loss in both datasets, with no more than a five-percentage-point clean precision decrease in either dataset, at a 1% node budget. This is an engineering screening tolerance chosen for the exploratory follow-up; it is not a stakeholder-validated operational requirement. The protocol was committed before the new replay, after predecessor test outcomes had been inspected. It is not an independent preregistration.')
h('1.4 Thesis and contribution boundary')
p('The thesis is that missing-telemetry evaluation must constrain review capacity and expose ranking ambiguity to support defensible detector selection. The completed experiment supports the usefulness of that measurement procedure and rejects the declared simple-fusion mitigation criterion on these inputs. It does not prove that all fusion methods fail, that a production SOC benefits, or that the method is original. A negative result is informative when it prevents a weak mitigation from being represented as a general solution.')
chapter('Chapter 2: Related Work and Research Position')
h('2.1 Provenance-based detection')
p('Jia et al. (2024) introduce MAGIC, which uses masked graph representation learning for APT detection. That work provides an architectural and data-artifact reference for provenance learning. The present score archives use local, MLP and GIN representations with nearest-neighbor scoring. They are project baselines, not an assertion of exact MAGIC reproduction or a comparison against MAGIC headline performance. Dataset preparation and scoring differences prevent that interpretation.')
h('2.2 Evaluation and investigation utility')
p('Guerra et al. (2026) distinguish alerting from investigation metrics, emphasize calibration controls and show how benchmark properties shape architectural conclusions. Their audited E3 analysis uses entity mappings and artifact treatment that this prepared static archive does not provide. This is direct prior overlap: reporting investigation-oriented metrics instead of relying on a single detection score is not new. Our narrower question is how a fixed-budget replay and explicit score-tie accounting expose the failure of particular score combinations under a specified loss intervention.')
p('The comparison is methodological, not a numerical leaderboard. Our node populations and labels are not interchangeable with an audited process-only scoring target. The manuscript therefore reports its own denominators and does not compare its precision numerically against the prior paper. Recovering entity identifiers and applying audited exclusions are substantive future work, not cosmetic cleanup.')
h('2.3 Combining detector rankings')
p('Combining rankings is established research. Cormack, Clarke and Buettcher (2009) study reciprocal rank fusion in information retrieval. This experiment uses simpler midrank-percentile combinations, not their reciprocal-rank formula, and does not transfer their empirical claims to security telemetry. Mean, minimum, maximum and median combinations are deliberately transparent baselines. An observed gain from one combination would not, by itself, establish methodological novelty.')
h('2.4 Candidate contribution and remaining literature gap')
p('The candidate applied contribution is a reproducible qualification procedure that couples a loss intervention, an exact review budget, tie-aware evidence yield and a clean-performance guardrail. Its current empirical finding is that apparently promising mitigation on one graph does not qualify across the two tested datasets. The individual ingredients are established ideas. Originality would require a fuller systematic comparison with telemetry-loss robustness, alert prioritization, calibrated ensembles and selective prediction research, followed by a clearly differentiated method or validated engineering process.')
p('The present source review establishes overlap and bounds claims; it is not an exhaustive literature review. Consequently this draft should be discussed as a feasibility and evaluation Praxis direction. No claim of being the first fixed-budget telemetry-loss study is made. A defensible final contribution may be an independently validated qualification procedure even if the best technical conclusion is to reject an unsafe mitigation.')
chapter('Chapter 3: Methodology')
h('3.1 Data and provenance')
p('The experiment reuses prepared CADETS and THEIA graphs and previously fitted prediction archives. The graph manifest identifies the MAGIC repository source revision aa0b647eea74b6faa0e52eb444370c4411a32cbe. The original preparation assigns train0 through train2 to fitting, train3 to benign calibration and test0 to labeled evaluation. The replay verifies every prediction archive against its original RESULTS.json SHA-256 record before reading it. Twelve archives cover two datasets, three fitting seeds and two observation conditions.')
table(['Dataset','Nodes','Labeled positive','1% budget'],[['CADETS','357,173','12,846','3,572'],['THEIA','344,767','25,319','3,448']])
p('Table 1. Evaluation populations and primary budgets. Unannotated nodes are treated as benchmark negatives, not independently verified benign activity. A dataset contributes one labeled test graph; the three fitting seeds do not create three independent campaigns.')
p('Prepared graph edges encode source, destination and relation type. The native CADETS test graph contains 840,299 edges and THEIA contains 628,107. The intervention retains the existing nodes while removing edges according to the archived 50% random-loss condition, with mask seed 20260920. Original timestamps and entity UUID mappings are absent. Thus the intervention is static evidence removal, not a contiguous outage, event delay, or completed attack-chain reconstruction.')
h('3.2 Archived models and score reuse')
p('The archived scoring pipeline uses a local feature representation, an MLP representation and a GIN representation. Each is scored by nearest-neighbor distance against a benign reference bank. The inherited configuration uses bank size 8,192, ten neighbors, a standardization floor of 0.001, fitting seeds 101, 211 and 307, and a calibration false-positive target of 0.01. Exact-neighbor scoring retains reference-bank multiplicity. The original repository contains model code and configuration; this experiment does not retrain those models.')
p('Raw anomaly scores are used as the three individual ranking baselines. The replay consumes only the archived score vectors and uses labels solely for evaluation. Original threshold alerts, discussed as predecessor context, use nonnegative archived margins. The new primary comparison uses a fixed number of reviewed nodes and therefore does not tune thresholds to test labels.')
h('3.3 Fixed fusion policies')
p('For model m and node i, define q(m,i) as its ascending average rank divided by the number of nodes. Larger values indicate greater anomaly. Equal raw scores receive the same midrank. The seven fusion policies are the mean, minimum, maximum and median across the three q vectors, plus the equal-weight means of local/MLP, GIN/local and GIN/MLP. Weights and policy definitions are fixed; no label-based fitting or policy selection occurs in the ranking function.')
p('Percentile normalization uses the full unlabeled score distribution for each evaluation batch. This is a transductive retrospective ranking procedure. It is not suitable for claiming a live, prefix-only detector. Ranking an available batch under a fixed capacity is a legitimate engineering task, but it is different from choosing an online alert threshold before future activity is observed. A deployment study would have to define window boundaries and delayed-score availability explicitly.')
h('3.4 Budget and tie-aware outcomes')
p('For budget fraction f and N nodes, B equals ceil(fN). The primary f is 0.01; secondary fractions are 0.001, 0.005, 0.02 and 0.05. Let A be the nodes strictly above the Bth-score boundary and T the tied nodes on that boundary. Exactly k = B - |A| tied nodes are needed. If a positive nodes are in A and t positive nodes are in T, expected true positives under uniform random boundary selection equal a + kt/|T|.')
p('The minimum possible true positives are a + max(0, k - (|T| - t)); the maximum are a + min(k,t). Expected precision divides the expected count by B, and recall divides it by all labeled positive nodes. The expected value is analytic, not an average over extra simulated attack trials. Boundary tie bounds describe ambiguity in this queue, not statistical confidence intervals. This prevents arbitrary node IDs from silently deciding the reported outcome.')
p('Random-ranking precision equals the positive prevalence. At 1%, recall cannot exceed B divided by the positive count even with perfect ranking: approximately 27.81% on CADETS and 13.62% on THEIA. Reporting precision with recall is therefore essential. A small budget mathematically limits coverage; it does not alone establish poor ranking.')
h('3.5 Experimental inventory and verification')
p('The complete inventory is 2 datasets x 3 fitting seeds x 2 views x 10 policies x 5 budgets = 600 metric rows. Every row retains population size, positive count, budget, counts above and at the boundary, positive counts, expected yield and tie bounds. Means and seed ranges are descriptive. No p-values or independent-campaign confidence intervals are computed from correlated nodes.')
p('An audit checks budget arithmetic, expected precision and recall, tie bounds and code/protocol hashes. It also exhaustively enumerates all nonempty binary label configurations for five tied nodes and all possible selection sizes, comparing analytic outcomes with every allowed subset. Permutation checks verify that node order does not supply hidden ranking information. These checks total 3,765 assertions under the audit counting convention. They establish computational consistency, not validity of the source labels.')
chapter('Chapter 4: Results')
h('4.1 Primary-budget outcomes')
table(['Policy','CADETS clean','CADETS loss','THEIA clean','THEIA loss'],[[name]+[f"{next(x for x in S if x['dataset']==ds and x['policy']==name)[v]*100:.2f}%" for ds in ['cadets','theia'] for v in ['clean','loss']] for name in ['local','mlp','gin','mean','minimum','maximum','median','local_mlp','gin_local','gin_mlp']])
p('Table 2. Mean expected attack-node precision at the 1% budget across three archived fitting seeds. Loss denotes the archived 50% edge-removal condition. No fusion policy satisfies the primary criterion in both datasets.')
p('The GIN baseline loses 30.66 percentage points of mean precision on CADETS and 59.19 points on THEIA. On THEIA, GIN/local fusion improves loss precision by 41.05 points relative to loss-condition GIN, reaching 63.61%. However, its clean precision decreases by 6.12 points, exceeding the five-point tolerance. On CADETS the same policy yields only 0.19% loss precision. The result is a dataset-specific opportunity, not a qualified mitigation.')
p('The mean of all three percentiles also improves THEIA loss precision, to 50.15%, but has a large clean penalty and zero loss-condition precision on CADETS. Minimum fusion reaches 82.38% clean precision on THEIA, slightly above clean GIN, but falls to zero under loss. Selecting a method by its clean score alone would therefore favor a particularly fragile candidate in this comparison.')
h('4.2 Variability and reference performance')
table(['Dataset / policy','Clean seed range','Loss seed range'],[[x['dataset'].upper()+' / '+x['policy'],f"{100*x['clean_range'][0]:.2f}-{100*x['clean_range'][1]:.2f}%",f"{100*x['loss_range'][0]:.2f}-{100*x['loss_range'][1]:.2f}%"] for x in S if x['policy'] in ['gin','gin_local']])
p('Table 3. Descriptive fitting-seed ranges at the primary budget. Large CADETS variation cautions against reading its mean as stable behavior. These are initialization outcomes on shared test data, not intervals over independent attacks.')
p('The label prevalence is approximately 3.60% on CADETS and 7.34% on THEIA. Several policies produce less useful queues than a uniformly random node sample in the loss condition. The best individual mean precision is GIN in each dataset/view at the primary budget. This strongest-individual reference is an evaluation comparator; no test-label oracle chooses a production policy.')
h('4.3 Budget sensitivity')
fig,axs=plt.subplots(1,2,figsize=(10,3.7))
for ax,ds in zip(axs,['cadets','theia']):
    for name in ['gin','gin_local','mean','minimum']:
        xs=[.001,.005,.01,.02,.05]; ys=[np.mean([x['precision'] for x in R if x['dataset']==ds and x['policy']==name and x['view']!='clean' and x['budget_fraction']==f])*100 for f in xs]
        ax.plot(np.array(xs)*100,ys,marker='o',label=name)
    ax.set_title(ds.upper());ax.set_xlabel('Nodes reviewed (%)');ax.set_ylabel('Expected precision (%)');ax.grid(alpha=.2);ax.legend(fontsize=8)
fig.tight_layout();fig.savefig(OUT/'budget_curves.png',dpi=180);plt.close(fig)
d.add_picture(str(OUT/'budget_curves.png'),width=Inches(6))
p('Figure 1. Descriptive mean precision across fitting seeds under 50% edge loss. The full inventory retains all policies, budgets and seeds. The curves visualize sensitivity, not a post hoc replacement of the primary 1% criterion.')
h('4.4 Connection to threshold alerting')
p('The predecessor analysis of these same archives found that THEIA GIN threshold recall increased from approximately 90.33% to 93.00% under loss, while false-positive rate increased from 2.02% to 57.85%. Mean threshold alerts increased from 29,328 to approximately 208,342. Meanwhile primary-budget precision fell from 81.75% to 22.56%. This is the practical disagreement motivating RQ1: broader alerting can preserve more labeled nodes somewhere in the output while degrading the limited queue.')
p('CADETS was already poorly calibrated in the clean condition, with a predecessor false-positive rate of about 41.95%, rising to 53.17% under loss. It must not be presented as a deployable clean baseline that becomes unusable only after deletion. The threshold figures are reused predecessor results, not additional fitted experiments or independent replication.')
chapter('Chapter 5: Discussion and Engineering Implications')
h('5.1 What the experiment established')
p('RQ1 is supported for the tested GIN score archives: loss sharply reduces primary-budget precision. RQ2 fails the declared mitigation criterion: none of the seven fusion policies qualifies across both datasets. RQ3 shows substantial fitting-seed variation and budget sensitivity; explicit tie treatment is necessary to make the queue outcome well defined. These are empirical statements about the saved scores and intervention, not universal properties of graph learning.')
p('The positive THEIA result is useful as a mechanism-development lead. Local scores that are weak alone may change which GIN-ranked nodes enter the queue when combined. The experiment does not identify the causal reason for that improvement: percentile geometry, representation complementarity and benchmark artifacts remain competing explanations. Resolving them requires node-level inspection with original identifiers and an evaluation population that supports artifact exclusions.')
h('5.2 An implementable qualification procedure')
p('A defender can use this artifact as a research-stage acceptance test. First specify the review unit and capacity. Then preserve the clean operating point, impose a declared loss condition, rank without label access, and evaluate both useful evidence and workload. Finally reject any candidate whose benefit depends on one dataset, an arbitrary tie ordering, or an undisclosed clean-performance sacrifice. The present seven candidates would not be approved under the frozen gate.')
p('The procedure does not automatically suppress alerts or change a production detector. Its output is a documented qualification decision for an evaluated policy. Before deployment, the unit should be a validated investigation case or entity group rather than an arbitrary node. A fixed 1% budget still represents thousands of nodes here, and the experiment does not establish how many humans can review them.')
h('5.3 What would make this a stronger Praxis')
p('The most defensible next contribution would pair this qualification procedure with an independently validated mitigation or an operationally tested model-selection decision. Existing archived logs should first be recovered with timestamps and UUID mappings, enabling audited labels, known artifact exclusions and campaign-level separation. This may avoid new collection if the original source release remains available, but the currently prepared graphs cannot substitute for those fields.')
p('A follow-on protocol should reserve previously unexamined executions before choosing a mitigation, use several independently generated loss masks, compare random loss with source-specific and contiguous outages, and include simple audited baselines. It should test calibration-only normalization and causal windowed scoring against the current batch-percentile approach. Budget ranges should be tied to case grouping and observed analyst workflow. Human usefulness would require an appropriately approved study and measured outcomes; this manuscript supplies no analyst-time or cost savings estimate.')
h('5.4 Threats to validity')
p('Internal validity is limited by inherited model and graph preparation choices, including train/test vocabulary dimensions and upstream benign-filter assumptions. SHA-256 verification establishes that inputs match their recorded artifacts, not that those artifacts are correct. Label construction can dominate apparent performance, and unannotated nodes may include attack-relevant activity. The experiment does not resolve this source uncertainty.')
p('External validity is limited by two prepared graphs from the same benchmark family, one loss rate and one mask. Repeated fitting seeds describe initialization variation rather than independent environments or campaigns. The chosen fusion suite is finite. A failure of seven transparent methods is not evidence that calibrated, learned, graph-aware or reliability-conditioned alternatives cannot work.')
p('Construct validity is limited because node-level precision is a proxy for useful investigative evidence. The representation lacks chronology and cannot measure time to alert, causal attack-chain completeness, or intervention success. A budget selected after a complete batch is observed does not establish online feasibility. The five-point clean tolerance is a provisional design choice rather than a validated defense requirement.')
p('Conclusion validity is limited because predecessor test outcomes guided the direction of this follow-up. Although the policy list and criterion were frozen before the replay, these inputs are development data. The complete result inventory limits selective reporting but does not restore independent confirmation. Novelty and generalization must remain separate from reproducible arithmetic.')
chapter('Chapter 6: Conclusions')
p('This study built and ran a complete existing-data experiment on attack-evidence preservation under constrained review. It produced 600 comparisons and 3,765 audit checks. Missing edges reduced the useful top-budget evidence of the archived GIN detector, even where predecessor threshold recall increased. Simple fusion provided a substantial THEIA improvement under loss but did not satisfy the cross-dataset mitigation criterion. The correct engineering decision under the declared gate is to reject these candidates as generally qualified defenses.')
p('The strongest current Praxis direction is a defensible evaluation and qualification process for preserving investigative evidence under telemetry degradation. The present work supplies a reproducible feasibility study and an honest negative mitigation result. It does not yet establish an original, deployment-ready defense. Advancing that claim requires audited entity-level data, unseen executions, broader loss conditions and a differentiated method or independently demonstrated decision benefit.')
chapter('References')
for txt in [
'Cormack, G. V., Clarke, C. L. A., & Buettcher, S. (2009). Reciprocal rank fusion outperforms Condorcet and individual rank learning methods. Proceedings of SIGIR, 758-759. https://doi.org/10.1145/1571941.1572114',
'Guerra, L., Chapuis, T., Duc, G., Mozharovskyi, P., & Nguyen, V.-T. (2026). How benchmarks and evaluation protocols shape conclusions in provenance-based intrusion detection. arXiv:2608.01454. https://arxiv.org/abs/2608.01454',
'Jia, Z., Xiong, Y., Nan, Y., Zhang, Y., Zhao, J., & Wen, M. (2024). MAGIC: Detecting advanced persistent threats via masked graph representation learning. 33rd USENIX Security Symposium. https://www.usenix.org/conference/usenixsecurity24/presentation/jia-zian',
'The George Washington University. (2026). University formatting requirements. Accessed October 7, 2026. https://gradpostdoc.gwu.edu/gw-etd-formatting',
'The George Washington University. (2026). Policies and procedures: Doctoral programs. Accessed October 7, 2026. https://online.engineering.gwu.edu/policies-procedures-doctoral',
'Pagan, G. (2026a). Defense opportunities experiment archive, October 7. Project research artifact, commit d563364c7d51e50325c22070e6aaeef814fa3019. https://github.com/garypagangit/praxis/tree/research/defense-opportunities-20261007/experiments/praxis_next/defense_opportunities_20261007',
'Pagan, G. (2026b). Budget defense experiment archive, October 7. Project research artifact. https://github.com/garypagangit/praxis/tree/research/budget-defense-praxis-20261007/experiments/praxis_next/budget_defense_20261007'
]:p(txt)
chapter('Appendix A: Reproduction and Evidence')
p('The study directory is experiments/praxis_next/budget_defense_20261007. PROTOCOL.txt defines the fixed experiment; run.py performs hash-verified score replay; RESULTS.json contains all 600 rows and twelve source hashes; SUMMARY.json gives primary-budget summaries; audit.py and AUDIT.json record verification. build_paper.py regenerates this manuscript and its plot from the completed results.')
p('To audit the public numerical evidence, run python audit.py from the study directory. To regenerate results, run python run.py --scores PATH, where PATH contains the original cadets and theia output directories. The local source is the September 20 embedding-baseline collected output archive. NumPy and SciPy are required for replay; document generation additionally requires python-docx and Matplotlib. PDF conversion uses LibreOffice.')
p('The public aggregate file supports arithmetic verification. It does not contain all archived per-node scores or raw graph data and therefore is not, by itself, a full independent reproduction package. Input hashes allow holders of the original archives to verify exact sources. Neither checksums nor arithmetic checks independently validate benchmark ground truth. The repository retains the model implementation, configuration and predecessor provenance files for tracing the experiment.')
chapter('Appendix B: Claim-to-Evidence Map')
table(['Claim','Evidence','Limit'],[
['600 comparisons completed','RESULTS.json, 600 rows','Saved-score replay, not 600 fits'],
['No fusion passes primary gate','AUDIT.json; SUMMARY.json','Seven fixed candidates only'],
['THEIA GIN/local gain under loss','Table 2; all three seeds','Clean penalty and failed CADETS transfer'],
['Ties handled explicitly','run.py; exhaustive audit','Uniform boundary selection assumption'],
['Paper follows GWU style','Margins, font, pagination','2026 Box template unavailable; review draft'],
['New defensible Praxis direction','Chapters 2 and 5','Novelty and external benefit unresolved']])
p('No experiment in this manuscript establishes human-versus-AI attack identification, live APT prevention, analyst cost savings, or independent-campaign superiority. Those claims would require different evidence and are intentionally outside the completed study.')
d.save(OUT/'Gary_Pagan_Budget_Defense_Praxis.docx')
print(OUT)
