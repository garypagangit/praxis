"""Build a GWU-style working Praxis from saved evidence, not anticipated findings."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[2]
OUT=ROOT/'output/doc/aivh_praxis';OUT.mkdir(parents=True,exist_ok=True)
R=json.loads((HERE/'RESULTS.json').read_text())
Q=json.loads((HERE/'ROUXII_QUALIFICATION.json').read_text())
doc=Document()
title='Distinguishing Autonomous AI from Human Cyberattack Operators Using Explainable Command Behavior'
for sec in doc.sections:
 sec.page_width=Inches(8.5);sec.page_height=Inches(11)
 sec.left_margin=sec.right_margin=Inches(1.25);sec.top_margin=sec.bottom_margin=Inches(1)
 sec.footer_distance=Inches(.5)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=2;normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.first_line_indent=Inches(.5)
for name in ['Heading 1','Heading 2','Title']:
 s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(12);s.font.color.rgb=None
 s.paragraph_format.first_line_indent=Inches(0);s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6)
 s.font.bold=True
doc.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
for name in ['GWU Index Heading','GWU Figure Caption','GWU Table Caption']:
 s=doc.styles.add_style(name,1);s.base_style=normal;s.font.size=Pt(12)
 s.paragraph_format.first_line_indent=Inches(0);s.paragraph_format.line_spacing=1
 s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(8)
 s.paragraph_format.keep_with_next=True
doc.styles['GWU Index Heading'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER

def p(text):return doc.add_paragraph(text)
def h(text,level=2):return doc.add_heading(text,level)
def center(text):
 x=p(text);x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.first_line_indent=Inches(0);return x
def chapter(text):
 doc.add_page_break();h(text,1)
def table(caption,heads,rows):
 doc.add_paragraph(caption,'GWU Table Caption')
 t=doc.add_table(rows=1,cols=len(heads))
 for c,s in zip(t.rows[0].cells,heads):c.text=s
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in rows:
  rr=t.add_row();rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c,s in zip(rr.cells,row):c.text=str(s)
 for row_index,row in enumerate(t.rows):
  for c in row.cells:
   for pp in c.paragraphs:
    pp.paragraph_format.line_spacing=1;pp.paragraph_format.first_line_indent=Inches(0)
    pp.paragraph_format.space_after=Pt(6)
    pp.paragraph_format.keep_with_next=row_index<len(t.rows)-1
    for r in pp.runs:r.font.size=Pt(10)
 return t
def numbering(sec,fmt,start):
 el=OxmlElement('w:pgNumType');el.set(qn('w:fmt'),fmt);el.set(qn('w:start'),str(start));sec._sectPr.append(el)
 sec.footer.is_linked_to_previous=False
 x=sec.footer.paragraphs[0];x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.first_line_indent=Inches(0)
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');x._p.append(f)

numbering(doc.sections[0],'lowerRoman',1);doc.sections[0].different_first_page_header_footer=True
center(title);center('\nGary Pagan\n\nThe George Washington University\nSchool of Engineering and Applied Science\nDoctor of Engineering\n\nWorking Praxis Manuscript\n5 October 2026')
center('\nProposal and completed development evidence\nPrepared for author and committee review')
doc.add_page_break();h('Abstract of Praxis',1);center(title)
p('Security teams may want to know whether observed attack activity is being directed by a human operator or an autonomous AI agent. A classifier can appear accurate while relying on differences between the datasets used to represent those classes. Such a classifier may make unreliable attributions when the task or environment changes. This Praxis investigates whether observable command behavior can support attribution across those changes while limiting false AI labels on human activity.')
p('The research uses public human command records from GAMBiT and KYPO, autonomous-agent logs from Honey for the Agent, and a qualified external probe from Rouxii. Completed development work compares simple linear classifiers, removes argument information, and checks exact feature contributions. The follow-up adds a TRACE-style baseline, frequency and sequence representations, order shuffling, and four held-out AI-family probes. Thresholds are selected on calibration records. The primary operating targets are at least 80% AI recall and at most 5% mean human-group false-positive rate. Existing test records have been examined, so this work is developmental rather than confirmatory.')
p('A fresh matched collection remains necessary. The contribution under investigation is an explainable test of attribution across environments. General APT attribution and analyst benefit have not been established.')
doc.add_page_break();h('Manuscript Status and Research Transparency',1)
p('This is a new working Praxis built from the command-attribution experiments. It does not replace or rewrite the previously completed paper, When Better APT Scores Hide Missed Attack Warnings. Chapters 1 through 3 state the proposed research case. Chapter 4 reports completed development evidence. Chapter 5 states the conclusions currently justified and the work needed for a stronger defense.')
p('The manuscript follows the GW page layout and the chapter structure observed in the supplied Date and Okhankhuele Praxis examples. It uses 12-point Times New Roman, double-spaced body text, 1.25-inch side margins, one-inch top and bottom margins, Roman front matter and Arabic chapter pagination. Final degree, committee, approval and submission pages are deliberately absent from this review edition; no approval or degree award is represented.')
p('AI tools helped research sources, implement code, analyze saved results and draft this manuscript. The author must verify the evidence, revise the writing in their own scholarly voice, and apply the program\'s disclosure requirements. No outcome from an AI-writing detector is promised. Research quality is assessed through traceable sources, reproducible analysis and bounded claims.')
for marker in ['@@TOC@@','@@FIGURES@@','@@TABLES@@']:
 doc.add_page_break();p(marker)
doc.add_page_break();h('List of Acronyms and Terms',1)
table('Terms used in this manuscript',['Term','Meaning'],[['AI','Artificial intelligence'],['APT','Advanced persistent threat; not the validated scope of this study'],['AUROC','Ranking performance across all decision thresholds'],['FPR','False-positive rate: human observations labeled AI'],['SOC','Security operations center'],['TF-IDF','A weighting of command words by frequency and distinctiveness'],['XAI','Explainable artificial intelligence'],['Calibration','Data used to choose a decision threshold'],['Transfer','Performance when tasks, environments or operators differ'],['Window','Ten qualifying commands, with other commands omitted']])
sec=doc.add_section(WD_SECTION_START.NEW_PAGE);sec.different_first_page_header_footer=False;numbering(sec,'decimal',1)
h('Chapter 1: Introduction',1)
h('1.1 Background and motivation')
p('An autonomous agent can choose a command, observe its result and choose another command without a person deciding each step. A human operator can use many of the same commands. The presence of a familiar tool therefore does not identify who is directing an operation. This study asks whether combinations of actions contain useful evidence about the operator.')
p('Ediga and Chattopadhyay (2026) demonstrate passive AI model-family attribution from terminal commands. Maddaloni et al. (2026) report command-selection fingerprints in autonomous attacker experiments. These findings make behavioral attribution plausible. They do not establish that a classifier trained on separate human and AI datasets will retain its performance on a new environment.')
h('1.2 Problem statement')
p('Command-log classifiers trained on human and autonomous-AI records collected in different environments may learn the collection setting rather than the operator class. Security teams relying on those classifications could incorrectly attribute human activity to AI or miss autonomous activity when tasks and environments change. The engineering problem is to measure and reduce this dependence while keeping attribution decisions understandable and false AI labels within a stated limit.')
h('1.3 Purpose and proposed thesis')
p('The purpose is to develop and evaluate a reproducible command-attribution procedure with controlled source comparisons, explicit error limits and inspectable feature contributions. The proposed thesis is that command behavior stripped of unnecessary environment detail can retain useful autonomous-AI detection across held-out conditions. This is a proposition to test. The current experiments do not yet prove it.')
h('1.4 Practical impact')
p('An AI attribution label would be an additional investigation cue, not a replacement for attack detection or authorization rules. A wrong label can waste analyst attention or distort an incident account. For illustration, a 5% false-positive rate applied to 10,000 human windows creates 500 false AI labels. This arithmetic is a workload example, not an observed deployment result or an estimate of financial savings. Attribution should not control automatic blocking until its practical error cost is independently justified.')
h('1.5 Research questions and hypotheses')
table('Table 1. Research alignment',['Question','Hypothesis and decision rule'],[
 ['RQ1. Can attribution transfer at a useful operating point?','H1. The calibration-selected model achieves >=80% AI recall and <=5% mean human-group FPR on a fresh matched held-out environment. Both conditions must hold; point estimates and group uncertainty are reported.'],
 ['RQ2. Does command order add information beyond command frequency?','H2. Ordered verbs improve AI recall by >=5 percentage points over frequency-only verbs at separately calibrated <=5% human-group FPR, and the paired 95% interval for the improvement excludes zero on fresh data.'],
 ['RQ3. Can the score be explained without relying on irrelevant environment identifiers?','H3. Under a frozen identifier-renaming intervention, >=99% of decisions remain unchanged, mean top-five contribution Jaccard is >=0.90, and score reconstruction error is <1e-10. Validity of the intervention must be checked first.']])
p('These are proposed confirmatory hypotheses for the fresh matched study. PX-118 provides development diagnostics only. Its broad removal of all arguments is not the harmless renaming intervention required by H3. Failure of H2 would mean frequencies are sufficient in the tested setting; it would not invalidate H1. Exact score reconstruction is a correctness requirement and is not itself a novel research result.')
h('1.6 Scope and limitations')
p('Each prediction concerns ten discovery-like commands, not an entire APT campaign. Operator labels come from provenance, not attack-stage labels. Live response, incident cost and analyst outcomes are outside the measured scope.')
p('Human and AI labels align with GAMBiT and Honey respectively. Common preprocessing does not make their tasks equivalent. Related windows are not independent people. The small Rouxii subset cannot establish population-wide transfer.')

chapter('Chapter 2: Literature Review')
h('2.1 Closest technical precedent')
p('TRACE uses passive terminal features to identify the family of an AI attack model, followed by an active forensic step. Its use of command features means this Praxis cannot claim to introduce passive terminal fingerprinting. The present comparison adapts its TF-IDF and linear-support-vector approach to a binary operator problem. That is a strong baseline rather than a reproduction of the original seven-family task (Ediga & Chattopadhyay, 2026).')
h('2.2 Command behavior and environmental influence')
p('Honey for the Agent examines model behavior across SSH environments. Its institutional abstract reports stable command-selection patterns and changes caused by environmental cues. Both findings matter here: command choice may contain operator information, but observed choices also depend on what the operator encounters. The full paper was not accessible during this review; claims are limited to its verified abstract and released dataset (Maddaloni et al., 2026).')
h('2.3 Active detection and source qualification')
p('LLM Agent Honeypot uses planted instructions and timing as evidence of autonomous interaction. It demonstrates a different observation strategy from passive command attribution. A successful response to a trap is an operational detection signal; it is not a substitute for independently assigned operator labels in a classifier evaluation (Reworr & Volkov, 2025).')
p('Rouxii makes additional autonomous-agent transcripts publicly available. Its reports include framework operations, model statements and evaluation answers. Those fields cannot simply be concatenated into classifier input: they could reveal the source directly. This study therefore extracts only explicit command events for compatible frameworks and applies the existing minimum-length rule. The dataset is external, but its usable subset is small (Anonymous, 2026).')
h('2.4 Human command evidence')
p('GAMBiT documents human red-team exercises and multiple telemetry sources, including command histories. It offers a stronger basis for human operator labels than assuming all traffic in an attack-stage dataset is human-directed. KYPO supplies another public human command source from cybersecurity training. It is useful for an external false-label check, although training tasks differ from operational attacks (Beltz et al., 2026; Svabensky et al., 2021).')
h('2.5 Explainability and the proposed gap')
p('For a linear classifier, each weighted feature contributes exactly to the decision score. Displaying those contributions can show what the model used. It cannot establish why an actual operator chose a command. A classifier may faithfully explain a decision that depends on an irrelevant dataset artifact. This distinction motivates testing explanations under controlled changes to recorded identifiers.')
p('Rudin (2019) argues for models whose decision process can be inspected directly in consequential settings. This supports beginning with linear models here, while recognizing that a large feature vocabulary can still be difficult to interpret. Exact score contributions provide an inspectable record; they do not automatically make the whole model understandable to an analyst.')
p('Slack et al. (2020) show that adversarially constructed classifiers can produce misleading post-hoc explanations. Their setting differs from the present experiment, but it supports an important design principle: explanation appearance is insufficient evidence of reliability. This study therefore checks the implemented score directly and separates faithful reconstruction from stability under changes to the input.')
table('Table 2. Prior work and the candidate contribution',['Established work','Remaining question here'],[['Passive command fingerprinting: TRACE','Does a binary operator classifier transfer across source and task?'],['AI command behavior: Honey','Do differences persist with humans in the same environment?'],['Active honeypot signals','Can ordinary compatible command records support attribution?'],['Exact linear feature contributions','Are influential features stable under irrelevant identifier changes?']])
p('The candidate contribution is the combined evaluation of transfer, false-label cost and explanation dependence under matched conditions. No search can guarantee priority. The literature review must be updated before proposal approval and submission. A change of dataset alone does not establish novelty; the proposed study must resolve an evaluation question that prior work leaves open.')

chapter('Chapter 3: Research Methodology')
h('3.1 Study design in simple terms')
p('First, collect commands with independently known operator labels. Second, separate related observations before fitting. Third, compare models that see different levels of command detail. Fourth, choose each decision threshold using calibration data. Finally, test on held-out records and inspect what changes when environment clues are removed. A fresh matched collection is required to test whether the measured differences actually transfer.')
doc.add_paragraph('Figure 1. Command-attribution study workflow','GWU Figure Caption')
for text in ['1. Known human and autonomous-AI operators','2. Same tasks, starting state and recorder (fresh study)','3. Recorded commands -> common normalization','4. Separate training / calibration / untouched evaluation','5. Compare frequency, ordered verbs and TRACE-style models','6. Fixed threshold -> operator label + feature contributions','7. Measure AI recall, human errors and transfer limits']:
 pp=center(text);pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_after=Pt(7)
h('3.2 Data qualification and observation unit')
p('PX-117C qualified 17,092 GAMBiT commands after removing structural problems. Its common-command subset contained 244 windows across 42 conservative participant groups. Honey contributed 7,659 eligible sessions. A shared allowlist admitted simple discovery-like commands and rejected shell compounds and malformed records. Absolute paths, numeric tokens and IP addresses were normalized. Each observation used ten qualifying commands, omitting intervening commands outside the allowlist.')
p('The main saved test contains 1,300 AI windows and 61 human windows from nine conservative groups. Participant numeric suffixes are grouped across source experiment prefixes to reduce possible overlap. This cautious grouping is not proof of unique identity. The previously examined split is retained for paired development comparisons rather than described as new independent evidence.')
h('3.3 Models and controlled comparisons')
table('Table 3. PX-118 model comparisons',['Model','Information available'],[['Lexical logistic','Normalized words and adjacent word pairs'],['TRACE-style SVC','Word TF-IDF(1,2), 10,000 features, linear SVC, C=1'],['Frequency logistic','Executable-name unigrams only'],['Ordered logistic','Executable-name unigrams, pairs and triples'],['Shuffled logistic','Same verb representation after fixed within-window shuffling'],['Masked logistic','Verbs retained; every argument replaced by ARG']])
p('The logistic models use fixed regularization C=1. All models receive the same group/class fitting weights: the classes receive equal total mass and human groups receive equal mass within their class. This controlled weighting differs from TRACE\'s original class-balancing setup and is disclosed. No hyperparameter search is performed. Six models are tested in a main comparison and four family-exclusion conditions, for 30 fits.')
p('Shuffling tests order while retaining the multiset of command verbs. It is deterministic from the record ID and a fixed seed, independent of its label. Since the ordered model also contains unigrams, good ordered-model performance alone is not proof that ordering helps. Its direct comparison with the frequency model and shuffled model is necessary.')
h('3.4 Thresholds, metrics and uncertainty')
p('A threshold converts a model score into an AI label. Each model uses the lowest calibration threshold at which the average human-group false-positive rate is at most 5%. SVC outputs are decision scores, not probabilities. The model selected for the main comparison is the one with the highest calibration AUROC, with a fixed name tie-break. Test results do not select the threshold or model.')
p('AI recall is the number of AI windows correctly labeled AI divided by all eligible AI windows. Human-window FPR is the fraction of human windows incorrectly labeled AI. Mean human-group FPR first calculates each group\'s rate and then averages groups equally. The third quantity is primary because one prolific participant should not dominate the estimate. AUROC measures ranking over thresholds; it does not establish safety at the selected threshold.')
p('Bootstrap intervals resample whole human groups. They are conditional on the observed sample and do not prove risk outside it. With zero errors an empirical bootstrap can return an interval of zero width; that is not a population guarantee. Paired model comparisons should retain the same observations. Repeated family probes reuse human records and must not be counted as independent replications.')
h('3.5 Explanation and perturbation tests')
p('For a linear score, s(x) = b + sum(w_j x_j). The contribution of feature j is w_j x_j. The implementation checks that the sum reconstructs the model decision score with error below 1e-10. For logistic regression the probability is the logistic transformation of this score. Contributions explain the linear score, not an additive decomposition of probability.')
p('PX-118 masks all arguments at evaluation time and reports label changes, score changes and overlap between the five largest absolute contributions. This is a stress test, because arguments and flags can carry useful task information. It is not claimed to preserve command meaning. Verb models should remain unchanged by construction. The later H3 test instead requires validated renaming of irrelevant identifiers in both commands and their corresponding environment.')
h('3.6 Independent-source probe')
p('The Rouxii archive is verified against its published MD5. The parser admits HackingBuddy assistant exec_command events and PentestGPT explicit execute_command records. It excludes prose suggestions, reasoning, outputs and fixed Rouxii operator traces. Only the first ten qualifying commands per run are used; no run is padded or merged with another. Results are reported separately from the human-only KYPO check.')
h('3.7 Matched collection and feasibility')
p('The fresh collection must give human and autonomous-AI operators the same task text, initial access, available commands and machine image. Both use the same command recorder. The planned pilot uses objective-driven discovery tasks, dummy files, equal budgets and resets between runs. It is a controlled command study, not a full APT simulation. Known operator assignment supplies the label, independently of the classifier.')
p('The collection toolkit accepts one command per request and records the result, exit code and timing. It requires an isolated, non-root, resource-limited container. Human recruitment, institutional determination and task images remain outstanding. A researcher-written script cannot replace human observations. The accompanying task sheet specifies what must be frozen before collection.')
p('A 12-operator pilot can estimate collection yield and between-person variability; it cannot certify a low population error rate. For a different endpoint, any false alert per independent person, zero errors among 59 held-out people gives a one-sided 95% binomial upper bound of about 4.95%. Repeated sessions do not add independent people, and that bound does not certify mean window FPR. The pilot will inform the sample-size plan for the primary endpoint.')
h('3.8 Reproducibility and stopping rules')
p('PX-118 source, protocol and input hashes were committed before fitting. The AWS worker runs bounded jobs and is stopped after artifact publication. The analysis retains prediction IDs, scores, model artifacts and external eligibility counts. Commands in published logs are treated as inert strings and are never executed by the classifier experiment.')
p('The study stops short of a transfer claim if labels or command records cannot be independently qualified, if no adequate external sample exists, or if the fresh matched operating targets fail. A negative result is retained. Changing the command minimum or threshold after seeing an external result requires a new development protocol and new evaluation data.')

chapter('Chapter 4: Completed Development Results')
h('4.1 Preliminary PX-117C evidence')
table('Table 4. PX-117C main operating points',['Model','AI recall','Human false labels','Mean group FPR'],[['Command length','15.23%','5/61','9.88%'],['Command words','97.62%','4/61','8.95%'],['Executable sequences','71.54%','0/61','0%']])
p('Neither substantive command model met both original main targets. The word model identified more AI records but exceeded the human error limit. The sequence model made no observed human errors while missing more AI records. Its zero-error sample does not establish zero future risk. These results motivated the controlled PX-118 comparisons.')
h('4.2 PX-118 paired model comparison')
base=[r for r in R['results'] if r['condition']=='main']
names={'lexical':'Words','trace_svc':'TRACE-style','frequency':'Verb frequency','ordered':'Ordered verbs','shuffled':'Shuffled verbs','masked':'Masked arguments'}
table('Table 5. PX-118 main development results',['Model','AI recall','Human errors','Group FPR','Target'],[[names[r['model']],f"{r['agent_recall']:.2%}",f"{r['fp']}/{r['fp']+r['tn']}",f"{r['human_group_mean_fpr']:.2%}",'Pass' if r['meets_target'] else 'Fail'] for r in base])
p('Calibration selected '+names[R['selected_by_calibration']]+'. The comparisons use previously examined data, so numerical target attainment is a development finding. It is not a successful confirmatory test of H1. The complete results file includes all family probes rather than retaining only favorable conditions.')
p('The masked-argument model is the promising comparator: it identifies 1,207 of 1,300 AI windows (92.85%) with zero false AI labels among 61 human windows. It meets the point targets in all four family probes, including a small 16-record DeepSeek probe. However, calibration selected the TRACE-style model, which identifies 98% of AI windows and falsely labels seven human windows. Selecting the masked model now is a development decision that must be validated on fresh data, not a retroactive change to the selection rule.')
p('Ordered verbs improve recall over shuffled verbs by only 0.31 percentage points. The conditional paired-session bootstrap interval spans -1.77 to 2.54 points. Ordered verbs have 9.38 points lower recall than frequency-only verbs at their respective calibration thresholds, with fewer human false labels. This comparison does not support the claim that order supplies a large recall advantage. The masked model improves recall over ordered verbs by 21.31 points in this sample, with a conditional interval of 18.77 to 23.69 points. These intervals hold fitted models and source conditions fixed; they do not resolve source confounding.')
table('Table 6. AI recall under held-out families',['Model','Alibaba','DeepSeek','Google','Meta'],[[names[k]]+[f"{next(r for r in R['results'] if r['model']==k and r['condition']=='family_'+fam)['agent_recall']:.1%}" for fam in ['alibaba','deepseek','google','meta']] for k in names])
p('Family recall must be read together with its human-group FPR in the full results table. Family sample sizes differ substantially; the DeepSeek probe has only 16 eligible records. High recall in that small probe is not a precise estimate of family-wide performance.')
h('4.3 Argument removal and score explanations')
table('Table 7. Frozen-model argument-removal stress test',['Model','Decision flips','Top-five overlap','Max score error'],[[names[r['model']],f"{r['argument_masking']['flip_rate']:.2%}",f"{r['argument_masking']['top5_mean_jaccard']:.3f}",f"{r['explanation_reconstruction_max_error']:.1e}"] for r in base])
p('Top-five overlap is mean Jaccard similarity of features with the largest absolute score contributions. It measures model explanation stability under this particular transform, not agreement with expert truth. Exact reconstruction checks faithful implementation. It cannot establish that the contributing features causally identify AI. Models that ignore arguments are invariant to argument removal by design.')
p('The word models retain many predictions after argument removal, while their top-contribution feature overlap falls to about 0.22. Predictions can therefore appear fairly stable even when the features explaining them change substantially. Because the intervention deletes potentially useful information, this is a diagnostic of representation dependence rather than evidence that H3 passes or fails.')
h('4.4 External-source eligibility and outcomes')
table('Table 8. Rouxii qualification',['Framework','Runs','Usable ten-command runs'],[[k,str(v['runs']),str(v.get('eligible_runs',0))] for k,v in Q['frameworks'].items()])
p('Only two unique windows qualified. The small yield is a substantive finding about dataset compatibility. The study did not shorten its window rule to manufacture a larger external evaluation. Consequently, the following counts are a two-record probe and cannot establish general transfer.')
p('All six models label both eligible Rouxii windows as AI. The masked model labels three of 180 eligible KYPO windows as AI. These two one-class probes answer separate questions and must not be pooled into an apparently balanced external benchmark. The collection settings still differ between classes.')
table('Table 9. Independent-source probes',['Model','KYPO false AI labels','Rouxii AI detected'],[[names[r['model']],str(r['external_kypo'].get('fp','NA'))+'/'+str(r['external_kypo'].get('fp',0)+r['external_kypo'].get('tn',0)),str(r['external_rouxii'].get('tp','NA'))+'/'+str(r['external_rouxii'].get('tp',0)+r['external_rouxii'].get('fn',0))] for r in base])
h('4.5 What remains untested')
p('H1, H2 and H3 remain open: a fresh environment containing both classes under the same recorder has not been tested. Numerical development success does not establish transfer or analyst benefit.')

chapter('Chapter 5: Discussion and Conclusions')
h('5.1 Interpretation')
p('The completed work shows that command representation changes the attribution tradeoff. It also demonstrates why a high ranking score is insufficient: the actual human-error rate at the chosen operating point must be examined. The result should guide the next collection rather than be used to declare a universal AI signature.')
p('The largest threat to interpretation is the link between dataset source and class. A classifier can be highly accurate by learning differences in tasks, command vocabulary or starting access. Removing arguments is informative but cannot erase all environmental influence. A matched collection is the most direct way to test the proposed thesis.')
h('5.2 Engineering product and proposed contribution')
p('The engineering product is a reproducible attribution evaluation package: source qualification, common preprocessing, grouped splits, strong baselines, threshold selection, saved predictions and exact feature contributions. Its potential research contribution is a measured account of which behavioral evidence survives changes in source, task and environment at a stated human-error cost.')
p('The work does not claim that a linear classifier, TF-IDF, command ordering, additive explanations or a threshold is new. The defense must show that the experiment answers a specific unresolved transfer question. If it cannot, the honest product is a bounded dataset and evaluation study rather than a general attribution system.')
h('5.3 Decisions for the next phase')
p('First, retain every PX-118 result and inspect baseline and perturbation differences. Second, complete a matched-data pilot with both classes and independently assigned labels. Third, use the pilot to freeze model selection, sample size, environment holdouts and success rules for a fresh evaluation. Fourth, test the complete protocol once on that evaluation. Repeatedly trying new thresholds on the same records cannot replace the final step.')
p('Broader APT attribution should be considered only after multi-stage campaigns are observed for both classes under comparable conditions. Until then, the title and conclusions remain about cyberattack-operator command behavior. This narrower claim is easier to measure and more defensible.')
h('5.4 Conclusion')
p('This Praxis has a plausible research direction and completed development evidence, but transfer is not yet established. The present package makes that uncertainty measurable and specifies the data needed to resolve it. Success will mean detecting autonomous AI on fresh matched conditions while limiting human false labels and showing which command evidence supports each decision.')

chapter('References')
refs=[
'Anonymous. (2026). Rouxii: Exploiting honeypots with deception-aware AI pentesters [Data set]. Zenodo. https://doi.org/10.5281/zenodo.21986588',
'Beltz, B., Doty, J., Fonken, Y., Gurney, N., Israelsen, B., Lau, N., Marsella, S., Thomas, R., Trent, S., Wu, P., Yang, Y.-T., & Zhu, Q. (2026). Guarding against malicious biased threats (GAMBiT) datasets: Revealing cognitive bias in human-subjects red-team cyber range operations. Data in Brief, 65, 112476. https://doi.org/10.1016/j.dib.2026.112476',
'Ediga, M., & Chattopadhyay, S. (2026). Trace: Unmasking AI attack agents through terminal behavior fingerprinting [Preprint]. arXiv. https://doi.org/10.48550/arXiv.2605.01186',
'Maddaloni, D., Safargalieva, A., & Vasilomanolakis, E. (2026). Honey for the agent: Cyber deception and behavioral fingerprinting of LLM-based attackers. In Proceedings of the 5th Workshop on Active Defense and Deception (AD&D) [Accepted/in press]. Springer. https://orbit.dtu.dk/en/publications/honey-for-the-agent-cyber-deception-and-behavioral-fingerprinting/',
'Reworr, & Volkov, D. (2025). LLM agent honeypot: Monitoring AI hacking agents in the wild (Version 2) [Preprint]. arXiv. https://doi.org/10.48550/arXiv.2410.13919',
'Rudin, C. (2019). Stop explaining black box machine learning models for high stakes decisions and use interpretable models instead. Nature Machine Intelligence, 1, 206-215. https://arxiv.org/abs/1811.10154',
'Slack, D., Hilgard, S., Jia, E., Singh, S., & Lakkaraju, H. (2020). Fooling LIME and SHAP: Adversarial attacks on post hoc explanation methods (Version 2) [Preprint]. arXiv. https://doi.org/10.48550/arXiv.1911.02508',
'Svabensky, V., Vykopal, J., Seda, P., & Celeda, P. (2021). Dataset of shell commands used by participants of hands-on cybersecurity training. Data in Brief, 38, 107398. https://doi.org/10.1016/j.dib.2021.107398',
'The George Washington University. (n.d.). University formatting requirements. Retrieved October 5, 2026, from https://gradpostdoc.gwu.edu/gw-etd-formatting'
]
for ref in refs:
 pp=p(ref);pp.paragraph_format.first_line_indent=Inches(-.5);pp.paragraph_format.left_indent=Inches(.5);pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_after=Pt(12)
chapter('Appendix A: Evidence and Reproduction')
p('Repository: https://github.com/garypagangit/praxis. Development experiment directory: experiments/praxis_next/aivh_transfer_20261005. Frozen PX-118 protocol and code: commit a8a75b5. Earlier common-command experiment: experiments/praxis_next/aivh_gambit_20261005. Literature assessment: LITERATURE_REVIEW.txt in the earlier directory.')
p('The PX-118 protocol, cloud freeze, qualification record, RESULTS.json and saved prediction files specify what was run. The cloud receipt records worker closure. The paper builder reads result artifacts directly to populate the new comparison tables. The matched collection plan and recorder are preparation tools; they do not constitute collected human data.')
p('Reproduction requires the qualified private records and the pinned Python dependencies listed in the saved worker requirements. Public redistribution of source logs must follow each source\'s terms. Published aggregates do not require exposing participant command histories. Independent verification should reconstruct confusion counts from predictions and check split identities, weights, thresholds and source exclusions.')
chapter('Appendix B: Committee Review Checklist')
for text in ['Confirm that the narrower operator-attribution problem is an acceptable Praxis topic.','Approve or revise the three aligned RQ/H pairs before matched collection.','Verify the novelty boundary against TRACE and the full Honey paper when available.','Determine institutional requirements for collecting new human sessions.','Confirm the feasible number of independent operators and the untouched evaluation allocation.','Review the exact distinction between a faithful score explanation and causal operator evidence.','Review AI-use disclosure and final GW submission requirements.']:
 pp=p(text);pp.paragraph_format.first_line_indent=Inches(0)
doc.core_properties.title=title;doc.core_properties.author='Gary Pagan';doc.core_properties.subject='Working Praxis; development evidence and matched transfer protocol'
dest=OUT/'Gary_Pagan_AI_Operator_Praxis.docx';doc.save(dest)
(OUT/'Manuscript_Text.txt').write_text('\n\n'.join(pp.text for pp in doc.paragraphs),encoding='utf-8')
print(dest)
