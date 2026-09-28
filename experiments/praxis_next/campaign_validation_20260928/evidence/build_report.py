"""Generate the audited validation addendum and concise machine-readable summary."""
import json, shutil, statistics
from pathlib import Path
from datetime import datetime,timezone
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parents[1]/'output'/'campaign_validation_20260928'
OUT.mkdir(parents=True,exist_ok=True)
rows=json.loads((ROOT/'RESULTS.json').read_text())
freeze=json.loads((ROOT/'FREEZE.json').read_text())
prep=json.loads((ROOT/'PREPARATION.json').read_text())
audit=json.loads((ROOT/'AUDIT.json').read_text())
assert audit['failed']==0
policies=['none','always_history','entropy','harm']
names={'none':'Current flow','always_history':'Always add history','entropy':'Entropy selection','harm':'Error-focused selection'}
tests=freeze['test_executions']
seeds=freeze['seeds']
def avg(execution,policy,key):
    return statistics.mean(r[key] for r in rows if r['execution']==execution and r['policy']==policy)
def get(execution,policy,seed): return next(r for r in rows if r['execution']==execution and r['policy']==policy and r['seed']==seed)
def fmt_pct(x): return f'{100*x:.3f}%'
def table(headers,data): return '\n'.join(['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(str,r))+' |' for r in data])
def iso(ms): return datetime.fromtimestamp(ms/1000,timezone.utc).strftime('%Y-%m-%d %H:%M UTC')
deltas=[]
for n in tests+['pooled']:
    for seed in seeds:
        a,b=get(n,'entropy',seed),get(n,'harm',seed)
        d={'execution':n,'seed':seed,'delta_macro_f1':b['macro_f1']-a['macro_f1'],
           'delta_exfil_warning_pp':100*(b['exfil_warning_recall']-a['exfil_warning_recall']),
           'delta_exfil_warning_missed':b['exfil_warning_missed']-a['exfil_warning_missed'],
           'delta_benign_false_alerts':b['benign_false_alerts']-a['benign_false_alerts'],
           'delta_acquired_fraction':b['acquired_fraction']-a['acquired_fraction']}
        d['direction_matches_original']=d['delta_macro_f1']>0 and d['delta_exfil_warning_pp']<0
        deltas.append(d)
hits=sum(d['direction_matches_original'] for d in deltas if d['execution']!='pooled')
pooled_hits=sum(d['direction_matches_original'] for d in deltas if d['execution']=='pooled')
total=sum(s['raw_rows'] for s in prep)
eligible=sum(s['eligible_rows'] for s in prep)
training=sum(s['eligible_rows'] for s in prep if s['execution'] in freeze['train_executions'])
testing=sum(s['eligible_rows'] for s in prep if s['execution'] in tests)
scope='Completed an adapted validation on a different public source, using six development executions and two later held-out executions. This is not an exact replication of the original four-class, two-evidence-group experiment.'
verdict=(f'The original direction (higher macro-F1 together with lower exfiltration warning recall under error-focused selection) appeared in {hits} of the 6 execution-by-seed comparisons and {pooled_hits} of the 3 pooled seed comparisons. '
         + ('It was directionally consistent across the two held-out executions and all three seeds in this adapted experiment.' if hits==6 else
            'It did not replicate consistently across the held-out executions in this adapted experiment.' if hits else
            'The adapted experiment did not reproduce that direction on either held-out execution.'))
summary={'status':'completed_and_audited','scope':scope,'verdict':verdict,'raw_rows':total,'eligible_rows':eligible,
         'training_rows':training,'test_rows':testing,'train_executions':freeze['train_executions'],'test_executions':tests,
         'matching_execution_seed_comparisons':hits,'matching_pooled_seed_comparisons':pooled_hits,
         'deltas':deltas,'audit_passed':audit['passed'],'audit_failed':audit['failed']}
secondary=('Wilson shows a small score/warning tradeoff in the opposite policy order: entropy has higher F1 but misses 1, 3 and 8 more exfiltration warnings across seeds (4 of 21,219 on average, or 0.018851 percentage points). On Harrison, entropy improves both F1 and warning recall. This secondary descriptive finding is small and execution-specific; the frozen primary direction remains unreplicated.')
summary['secondary_finding']=secondary
(ROOT/'SUMMARY.json').write_text(json.dumps(summary,indent=2))
pooled_table=[[names[p],f"{avg('pooled',p,'macro_f1'):.6f}",fmt_pct(avg('pooled',p,'exfil_warning_recall')),
               f"{avg('pooled',p,'exfil_warning_missed'):.1f}",f"{avg('pooled',p,'benign_false_alerts'):.1f}",
               fmt_pct(avg('pooled',p,'acquired_fraction'))] for p in policies]
execution_table=[[n,names[p],f"{avg(n,p,'macro_f1'):.6f}",fmt_pct(avg(n,p,'exfil_warning_recall')),
                  f"{avg(n,p,'benign_false_alerts'):.1f}"] for n in tests for p in policies]
delta_table=[[d['execution'],d['seed'],f"{d['delta_macro_f1']:+.6f}",f"{d['delta_exfil_warning_pp']:+.4f}",
              f"{d['delta_exfil_warning_missed']:+d}",f"{d['delta_benign_false_alerts']:+d}"] for d in deltas if d['execution']!='pooled']
coverage=[[s['execution'],'Test' if s['execution'] in tests else 'Train',f"{s['raw_rows']:,}",
           f"{s['eligible_rows']:,}",f"{s['class_counts']['other_attack']:,}",f"{s['class_counts']['exfiltration']:,}"]
          for s in sorted(prep,key=lambda s:(s['raw_min_start_ms'],s['execution']))]
macro_table=[[names[p],f"{statistics.mean(avg(n,p,'macro_f1') for n in tests):.6f}",
              fmt_pct(statistics.mean(avg(n,p,'exfil_warning_recall') for n in tests)),
              fmt_pct(statistics.mean(avg(n,p,'benign_false_alert_rate') for n in tests))] for p in policies]

md=f'''# Campaign validation addendum - 28 September 2026

Gary Pagan | Companion to *When Better APT Scores Hide Missed Attack Warnings*

## Result

{scope}

{verdict}

**Secondary descriptive finding:** {secondary}

All {total:,} rows in the eight AIT netflow archives were inspected. After explicit exclusions, {training:,} rows train the final models and {testing:,} rows evaluate them. All eligible rows are used; no class caps or row sampling. The independent computational audit passed {audit['passed']:,} checks with zero failures.

### Pooled held-out outcomes

Means across seeds 8101, 8102, 8103. Fractional alert counts are means of integer counts. Each seed evaluates the same {testing:,} rows; this is not three times as many observations.

{table(['Policy','Macro-F1','Exfil warnings','Missed exfil warnings','Benign alerts','History acquired'],pooled_table)}

Warning recall means any non-benign prediction on a publisher-labeled exfiltration flow. Exact-stage recall, confusion matrices, per-class F1, error, weighted error, spending and all seed results are in RESULTS.json/RESULTS.csv. Equal available per-row budgets do not imply equal actual spending.

## Complete source coverage

{table(['Execution','Use','Raw rows','Eligible rows','Other attack','Exfiltration'],coverage)}

Excluded {sum(s['unknown_label_rows'] for s in prep):,} rows with unknown/empty labels, {sum(s['invalid_numeric_rows'] for s in prep):,} rows with invalid numeric values, {sum(s['duplicate_rows_removed'] for s in prep):,} exact duplicate rows and {sum(s['conflicting_identity_rows'] for s in prep):,} conflicting-identity rows. The final identity fingerprints have no train/test overlap. Repeated observations across probes can remain because sensor identity is absent from the released CSVs.

The latest training completion is {iso(max(s['max_end_ms'] for s in prep if s['execution'] in freeze['train_executions']))}; the earliest evaluation start is {iso(min(s['min_start_ms'] for s in prep if s['execution'] in tests))}. The final split is both execution-disjoint and calendar-separated. Development forward folds are ordered by execution start; overlapping development execution intervals mean those folds are not strict calendar-forward stream simulations.

## Held-out runs and sensitivity

{table(['Execution','Policy','Macro-F1','Exfil warnings','Benign alerts'],execution_table)}

Error-focused minus entropy-focused acquisition, separately for every execution and fitting seed:

{table(['Execution','Seed','Delta F1','Delta warning pp','Delta missed exfil','Delta benign alerts'],delta_table)}

Execution-macro averages give each of the two held-out runs equal weight:

{table(['Policy','Macro-F1','Exfil warning recall','Benign false-alert rate'],macro_table)}

Seed variation measures fitting sensitivity, not campaign-population uncertainty. With two related held-out laboratory executions, no population confidence interval or real-world prevalence estimate is claimed.

## Fixed experiment

Three native-label classes: benign, other attack, exfiltration. Current inputs use 10 direction/volume/duration statistics. Optional history uses eight summaries of same-initiator flows completed strictly before the current start, within one hour. No labels enter history; absolute time, IP strings, ports, FQDN, role, network and execution identifiers are not model inputs. Endpoint identities only group historical observations.

Two LightGBM classifiers per fit use 150 boosting iterations, 15 leaves, learning rate .05, minimum child count 10 and L2=1. No class weighting, tuning or early stopping. Five development forward folds generate out-of-fold predictions for two selectors (100 boosting iterations, nine leaves). Selectors observe only current features and current probabilities. Targets are entropy reduction or reduction of weighted hard error, with illustrative weights [1,1,4]. Final models fit all six development runs.

The four policies are current flow only, always-history, predicted-entropy-reduction >0, and predicted-weighted-error-reduction >0. All have an available per-row budget of 2; history costs 2 and arrives after .75 time units before a deadline of 1. This is a clean-delivery offline replay. It measures neither real retrieval latency/cost nor warning before a flow completes. There are 36 classifier fits and six selector fits across three seeds.

## What the evidence can justify

1. The pipeline now has external-source evaluation with two later, whole executions held out. This addresses the lack of any disjoint-execution test for the adapted configuration.
2. The original warning-loss result retains its original population, seed sensitivity and class definitions. This new test must not be presented as a rerun of every original intervention on independent real APT campaigns.
3. The AIT execution generator and scenario are shared. Simulated legitimate traffic is present, but source diversity is not equivalent to enterprise field validation.
4. Author flow labels are based on topology, service ports and attack windows. The released UDP notebook marks remaining unlabeled traffic as exfiltration. A descriptive audit confirms that all 183,121 native exfiltration rows use DNS port 53 and include a publisher-identified attacker endpoint. This supports consistency with the stated scenario, not independent payload verification. These topology/port fields are excluded from predictors.
5. There is no native movement flow class here. The comparison uses one history group, rather than the original role and history groups. Three-class macro-F1 is not directly comparable with four-class macro-F1.
6. Classification uses completed flows. The results are not claims about forecasting an attack, stopping exfiltration, successful lateral movement, analyst workload, or real collection savings.

### Suggested defense wording

"I extended the work to the complete AIT netflow release and froze a chronological whole-execution evaluation before fitting. Six executions trained the models and two later executions supplied 1,067,211 held-out rows. The adapted experiment uses the native three-class task and one optional history group. {verdict} The new evidence bounds the claim rather than making a universal statement about all APT campaigns."

### Remaining scope for a stronger campaign claim

An exact PX-081 campaign replication still needs a separate execution with defensible benign, other-stage, movement and exfiltration flow labels; support in the required chronological partitions; role/history groups established independently of outcome labels; and the same frozen two-group policy/latency comparisons. Real operational benefit additionally requires measured collection and analyst outcomes. These are limits to the claims, not unfinished computation on the AIT data.

## Reproduction and evidence

- `ACQUISITION.json`: all eight publisher MD5s, local SHA256s, download URLs and CRC checks.
- `QUALIFICATION.json`: every native CSV's class labels and time coverage.
- `NATIVE_LABEL_DIAGNOSTIC.json`: DNS/attacker-endpoint consistency of all 183,121 native exfiltration labels; descriptive only, with no model changes.
- `PROTOCOL.md`, `FREEZE.json`: design, software versions, features, split and pre-fit code/data hashes.
- `PREPARATION.json`: row conservation, exclusions, support and strict history checks.
- `FIT_LOG.json`: every classifier training population and elapsed time.
- `RESULTS.json`, `RESULTS.csv`: every execution, policy and seed, including pooled results.
- `AUDIT.json`: independent count-based recomputation and artifact checks.
- `ARTIFACTS.json`: model/prediction hashes and paths.
- `C:/w/campaign_validation_20260928`: complete source archives, prepared arrays, models and per-row predictions.

Run from the report directory with the recorded environment: `python acquire_qualify.py`, `python validate.py prepare`, `python validate.py run`, `python audit.py`, then `python build_report.py`. Scripts currently use the documented Windows data root; edit that path when reproducing elsewhere and record the amendment. Source downloads and original publisher notebooks are retained; the notebooks were inspected, not executed.

## Primary sources

Soro, F., Landauer, M., Skopik, F., Hotwagner, W., and Wurzenberger, M. (2024). *AIT Netflow Data Set*, version 2. Zenodo. https://doi.org/10.5281/zenodo.13168643

Landauer, M., Skopik, F., Frank, M., Hotwagner, W., Wurzenberger, M., and Rauber, A. (2023). Maintainable Log Datasets for Evaluation of Intrusion Detection Systems. *IEEE Transactions on Dependable and Secure Computing*, 20(4), 3466-3482. Author manuscript: https://arxiv.org/abs/2203.08580

Landauer, M., Skopik, F., Wurzenberger, M., Hotwagner, W., and Rauber, A. (2021). Have it Your Way: Generating Customized Log Datasets With a Model-Driven Simulation Testbed. *IEEE Transactions on Reliability*, 70(1), 402-415. https://doi.org/10.1109/TR.2020.3031317

The publisher's `label_info.txt` and `2_label_logs.ipynb` document the exact native labels and labeling rules used in this evaluation.
'''
(OUT/'CAMPAIGN_VALIDATION_REPORT.md').write_text(md,encoding='utf-8')

# A compact, explicitly paginated addendum accompanies the complete Markdown report.
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=Inches(.7);sec.right_margin=Inches(.7)
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for name in ['Normal','Body Text']:
    style=doc.styles[name];style.font.name='Calibri';style.font.size=Pt(10.5)
    style.paragraph_format.space_after=Pt(6)
for name,size in [('Title',24),('Heading 1',17),('Heading 2',12)]:
    st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('003B5C')
sec.header.paragraphs[0].text='GARY PAGAN  |  PRAXIS EVIDENCE ADDENDUM  |  28 SEPTEMBER 2026'
sec.header.paragraphs[0].style='Caption'
foot=sec.footer.paragraphs[0];foot.alignment=2
foot.add_run('Campaign validation  |  ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
def para(t): doc.add_paragraph(t)
def head(t): doc.add_heading(t,1)
def tbl(headers,data):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
    for c,x in zip(t.rows[0].cells,headers): c.text=str(x)
    for row in data:
        for c,x in zip(t.add_row().cells,row): c.text=str(x)
    for row in t.rows:
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for run in p.runs:run.font.size=Pt(8.5)
    doc.add_paragraph()
def page():doc.add_page_break()
doc.add_heading('Campaign validation',0)
para('Companion to When Better APT Scores Hide Missed Attack Warnings')
head('1. Outcome and scope')
para(scope);para(verdict);para(secondary)
para(f'Complete source coverage: {total:,} raw rows across eight testbeds. Final training: {training:,} rows. Held-out evaluation: {testing:,} rows. Audit: {audit["passed"]:,} checks passed, zero failures.')
doc.add_heading('Pooled held-out results',2)
tbl(['Policy','Macro-F1','Exfil warnings','Missed warnings','Benign alerts','Acquire'],pooled_table)
para('Values are means over three fitting seeds. Fractional alert counts are averages of integer counts. The same evaluation population is reused across seeds; it is not three independent samples.')
para('Warning recall counts any attack prediction on a true exfiltration-labeled flow; exact-stage recall is reported separately in the complete results. All policies have the same available per-row budget, while realized acquisition fractions differ.')
page();head('2. Full release and temporal separation')
tbl(['Execution','Use','Raw rows','Eligible','Other attack','Exfil'],coverage)
para(f'Excluded: {sum(s["unknown_label_rows"] for s in prep):,} unknown/empty labels; {sum(s["invalid_numeric_rows"] for s in prep):,} invalid numeric records; {sum(s["duplicate_rows_removed"] for s in prep):,} exact duplicate rows; {sum(s["conflicting_identity_rows"] for s in prep):,} conflicting-identity rows. All eligible rows are used without class sampling or fitting caps.')
para(f'Latest training completion: {iso(max(s["max_end_ms"] for s in prep if s["execution"] in freeze["train_executions"]))}. Earliest evaluation start: {iso(min(s["min_start_ms"] for s in prep if s["execution"] in tests))}. There is no identity-hash intersection across the final split.')
para('The test partition is both execution-disjoint and calendar-separated. Development folds advance through execution start order, but overlapping development intervals mean those folds are not a strict online calendar simulation.')
doc.add_heading('Fixed inputs and models',2)
para('Ten current-flow statistics describe protocol, duration, packets and bytes. Eight history statistics describe same-initiator flows completed strictly earlier, within one hour. Labels never enter history. Raw identities, ports, absolute timestamps, publisher roles, networks and execution identifiers are excluded from predictors.')
para('LightGBM: 150 boosting iterations, 15 leaves, learning rate .05, minimum child support 10, L2=1. Selectors: 100 boosting iterations and nine leaves. Five forward development folds; final classifiers fit all six development runs. Three seeds; no tuning, early stopping or class weighting. Total: 36 classifier and six selector fits.')
para('Native classes are benign, other attack and exfiltration. Acquisition selectors use only current features and current probabilities. History is acquired when predicted entropy reduction or weighted-error reduction is positive. Error weights are [1,1,4]. Clean-delivery budget 2, history cost 2, latency .75 and deadline 1 are illustrative simulated units.')
page();head('3. Each held-out execution')
tbl(['Execution','Policy','Macro-F1','Exfil warnings','Benign alerts'],execution_table)
doc.add_heading('Error-focused minus entropy-focused selection',2)
tbl(['Execution','Seed','Delta F1','Warning pp','Missed exfil','Benign alerts'],delta_table)
para('Positive delta missed-exfil means more exfiltration flows predicted benign. Positive delta benign-alerts means more legitimate flows flagged as attack. Seed spread measures fitting sensitivity; it does not estimate campaign-population uncertainty.')
doc.add_heading('Equal weight for each held-out execution',2)
tbl(['Policy','Macro-F1','Exfil warnings','Benign false-alert rate'],macro_table)
page();head('4. Defensible interpretation')
for t in [
    'External-source evidence is now available for the adapted configuration, with whole executions held out. The original single-campaign result remains bounded by its original class definitions and seed sensitivity.',
    'AIT uses a shared laboratory generator and attack scenario with simulated legitimate background traffic. Two held-out runs do not establish performance across unrelated APT families or enterprise environments.',
    'Publisher labels use topology, ports and attack windows. The UDP notebook assigns residual unlabeled flows to exfiltration. All 183,121 such rows use DNS port 53 and include a publisher-identified attacker endpoint. This supports scenario consistency, not independent payload verification; those identity/port fields are excluded from models.',
    'There is no native movement flow class. This test uses three classes and one optional history group, instead of the original four classes and two optional groups. Macro-F1 values cannot be directly compared across those tasks.',
    'Completed-flow classification does not measure attack forecasting, interruption of exfiltration, confirmed lateral movement, real acquisition costs or analyst workload.',
    'Sensor identity was discarded in source aggregation. Exact duplicate handling cannot guarantee that all near-duplicate observations of the same network activity were removed. Rows are not independent attack executions.']:
    para(t)
doc.add_heading('Suggested defense wording',2)
para('I evaluated the complete AIT netflow release using a frozen chronological split of six training executions and two held-out executions. The adapted experiment has 1,067,211 held-out rows and uses the source-native three-class task with one optional history group. '+verdict)
doc.add_heading('Requirement for an exact campaign replication',2)
para('A separate execution still needs defensible benign, other-stage, movement and exfiltration labels, support at the required chronological cutoffs, role/history groups defined independently of outcome labels, and the original two-group policy and latency comparisons. Operational benefit also requires measured collection and analyst outcomes. These are scope limits, not unfinished AIT computation.')
page();head('5. Reproduction and evidence')
para(f'The independent audit passed {audit["passed"]:,} checks. It verifies frozen code/data hashes, feature widths, strict history timing, disjoint execution membership, OOF targets, policy decisions, probability selection and normalization, budgets, every confusion matrix and every reported metric. Source acquisition separately verifies all eight publisher MD5s and ZIP CRCs.')
for title,body in [
    ('Source and preparation','ACQUISITION.json; QUALIFICATION.json; PREPARATION.json; publisher metadata, README, label dictionary and labeling notebooks.'),
    ('Frozen design','PROTOCOL.md and FREEZE.json include the partition, feature lists, seeds, software versions and pre-fit code/data hashes.'),
    ('Complete outcomes','RESULTS.json and RESULTS.csv retain every seed, execution and policy. SUMMARY.json records all paired deltas. FIT_LOG.json records classifier fitting populations.'),
    ('Audit and artifacts','AUDIT.json contains individual checks; ARTIFACTS.json binds all models and per-row predictions by SHA256.'),
    ('Local data root','C:/w/campaign_validation_20260928 contains the eight source archives, prepared arrays, fitted models and per-row predictions.')]:
    doc.add_heading(title,2);para(body)
para('Reproduce with the recorded Python environment: acquire_qualify.py; validate.py prepare; validate.py run; audit.py; build_report.py. Scripts currently use the documented Windows data root; a reproduction on another machine must record any path amendment. No publisher notebook was executed.')
doc.add_heading('Primary references',2)
para('Soro, F., et al. (2024). AIT Netflow Data Set, version 2. Zenodo. https://doi.org/10.5281/zenodo.13168643')
para('Landauer, M., et al. (2023). Maintainable Log Datasets for Evaluation of Intrusion Detection Systems. IEEE Transactions on Dependable and Secure Computing, 20(4), 3466-3482. Author manuscript: https://arxiv.org/abs/2203.08580')
para('Landauer, M., et al. (2021). Have it Your Way: Generating Customized Log Datasets With a Model-Driven Simulation Testbed. IEEE Transactions on Reliability, 70(1), 402-415. https://doi.org/10.1109/TR.2020.3031317')
para('Native label definitions and rules: publisher label_info.txt and 2_label_logs.ipynb, retained with the evidence.')
doc.save(OUT/'Gary_Pagan_Campaign_Validation_Addendum.docx')
for n in ['SUMMARY.json','RESULTS.csv','RESULTS.json','PROTOCOL.md','FREEZE.json','AUDIT.json']:
    shutil.copy2(ROOT/n,OUT/n)
print(json.dumps(summary,indent=2))
