"""Create the preliminary Praxis from actual saved results, with explicit limits."""
import json,os
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=Path(__file__).resolve().parent
OUT=Path(os.environ.get('PX117_OUTPUT','C:/w/px117_aws_20261005/collected/outputs'))
DEST=HERE/'delivery';DEST.mkdir(exist_ok=True)
results=json.loads((OUT/'RESULTS.json').read_text())
base=[r for r in results['results'] if r['name']=='record_holdout']
selected=results['selected_model']
families=[r for r in results['results'] if r['model']==selected and r['name'].startswith('family_')]
tasks=[r for r in results['results'] if r['model']==selected and r['name'].startswith('human_task_')]
doc=Document();section=doc.sections[0]
section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
section.page_width=Inches(8.5);section.page_height=Inches(11)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=2;normal.paragraph_format.space_after=Pt(0)
for name in ['Heading 1','Heading 2']:
    style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
    style.font.size=Pt(14 if name=='Heading 1' else 12)
    style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(6)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def p(text):return doc.add_paragraph(text)
def h(text):doc.add_heading(text,2)
def chapter(text):
    par=doc.add_heading(text,1);par.paragraph_format.page_break_before=True
def table(headers,rows):
    t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    for row in t.rows:
        props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.line_spacing=1.0;par.paragraph_format.space_after=Pt(4)
                for run in par.runs:run.font.size=Pt(10)
    props=t.rows[0]._tr.get_or_add_trPr();props.append(OxmlElement('w:tblHeader'))
    return t

for text,size in [('Distinguishing Autonomous AI Attacks from Human-Operated Attacks',20),
                  ('PX-117: A Preliminary Praxis Study',14),('Gary Pagan',12),
                  ('Doctor of Engineering Research',12),('5 October 2026',12)]:
    par=p(text);par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after=Pt(24)
    for run in par.runs:run.font.size=Pt(size);run.bold=(size>=14)
p('Status: development experiment and working classifier completed. Matched confirmation and defense readiness remain open. This document does not claim successful real-world APT attribution.')

chapter('Abstract')
p('This study asks whether recorded command sequences can distinguish fully autonomous AI attacks from human-operated attacks. Fully autonomous means that an agent chooses and executes the steps after an initial objective, without human guidance during the run. AI-assisted humans are outside the primary scope. The intended application is investigation of APT-style intrusions, but the present evidence comes from public laboratory and cybersecurity-training records.')
p('An AWS development experiment compared a session-length baseline, a command n-gram classifier and a small gradient-boosted tree model. A strict input audit retained 267 human record files and 13,397 agent prefixes. Each observation used up to ten commands, with timing, command outputs and source metadata excluded. Thresholds were selected on calibration data, and repeated normalized prefixes were separated across fitting, calibration and test partitions. The experiment comprised 36 descriptive comparisons.')
p('The selected n-gram model detected all 2,383 agent records in the main test set and mislabeled 4 of 54 human records, a 7.4% human false-positive rate. Its AUROC was 0.9995. A human-task holdout produced 39 false positives among 62 human records, or 62.9%. These results show strong separation within the available collections but substantial sensitivity to unfamiliar human tasks. Dataset source is perfectly associated with operator class, so the scores do not establish that autonomous AI itself caused the difference. The deliverables are a runnable classifier, a saved evidence package and a matched-evaluation plan. The proposed 5% false-positive criterion and broader APT claim are not established.')

chapter('Contents')
for text in ['Chapter 1. Introduction','Chapter 2. Literature Review','Chapter 3. Methodology','Chapter 4. Results','Chapter 5. Conclusions and Required Confirmation','References','Appendix A. Reproducibility and Classifier Use']:
    p(text)

chapter('Chapter 1. Introduction')
h('1.1 Problem and practical importance')
p('A security team may recognize an intrusion without knowing whether a person or an autonomous AI agent is directing it. A classifier could help analysts characterize an operation, but a wrong label could also distort an investigation. The practical question is therefore not simply whether a model obtains a high score. It is whether it recognizes unfamiliar autonomous agents while rarely assigning an AI label to human activity.')
p('Published work demonstrates that AI agents can perform multi-step cyber tasks and that terminal commands can reveal differences among agent families. Those findings motivate an empirical test of operator classification. They do not establish a reliable detector of AI-operated APT campaigns. The present study evaluates a small, explicit first step and measures its limitations.')
h('1.2 Problem statement')
p('Security teams lack validated evidence that a passive command-sequence classifier can distinguish fully autonomous AI attack activity from human-operated activity across unfamiliar tasks and models at a low human false-positive rate. Without that evidence, a classifier may label differences in datasets or tasks as differences in operator type, leading analysts to overstate AI involvement.')
h('1.3 Purpose and scope')
p('The purpose is to develop and test a binary classifier whose target is autonomous AI versus human operation. The development batch uses recorded shell behavior. It does not identify a threat group, establish malicious intent from ordinary commands, or infer the sponsor of an APT. It does not measure analyst time saved, attacks prevented or financial losses avoided. Those are separate possible applications requiring additional evaluation.')
p('An autonomous run begins with a task assignment and proceeds without human command selection, corrective hints or intervention. A human run is performed without generative AI assistance. Conventional scripts form a useful later control for automation; AI-assisted humans are an optional later boundary condition, not a required class in this study.')
h('1.4 Research questions and hypotheses')
table(['Question','Proposed confirmation hypothesis'],[
    ['RQ1: Can command behavior distinguish autonomous AI from human attacks?','H1: At a calibration-fixed threshold, autonomous-AI recall is at least 80% and the human false-positive rate is at most 5%.'],
    ['RQ2: Does detection survive an unfamiliar AI family?','H2: Each prespecified held-family condition retains at least 70% recall and at most 5% human false positives.'],
    ['RQ3: Does detection use more than speed?','H3: Removing timing reduces recall by no more than 10 percentage points at the same human false-alert ceiling.']])
p('These are proposed confirmation targets from the working design. This development batch does not provide a matched cohort or a powered confirmation test. It uses no timing features, so it also does not measure a paired timing-removal effect. A high development AUROC cannot be substituted for any hypothesis.')

chapter('Chapter 2. Literature Review')
h('2.1 Passive terminal fingerprinting')
p('Ediga and Chattopadhyay (2026) introduce Trace, whose passive stage represents terminal command sequences using TF-IDF and a linear support vector classifier. Its main attribution target is the AI model family, and it examines transfer across agent scaffolds. This establishes a close method precedent. Our study cannot claim to invent passive terminal fingerprinting. Its proposed distinction is a carefully controlled autonomous-AI versus human comparison, which remains to be confirmed.')
h('2.2 Earlier agent detection')
p('Reworr and Volkov (2024) investigate LLM agents using a honeypot that combines behavioral and timing observations with active interventions. This provides a foundation for asking whether passive observations are sufficient. It does not justify treating a reaction to a prompt-injection trap as independent ground truth for all operator classes.')
h('2.3 Human and agent behavior sources')
p('Svabensky et al. (2021) publish shell commands from participants in hands-on cybersecurity training. The later Zenodo release provides human-generated records suitable for source and feature qualification. Maddaloni et al. (2026) publish Honey for the Agent, containing automated interactions with SSH environments. These sources establish that relevant records are available. Their different tasks and collection systems prevent a direct pooled comparison from being a matched experiment.')
h('2.4 Multi-step cyber capability')
p('Measuring AI Agents\' Progress on Multi-Step Cyber Attack Scenarios (2026) examines agents performing extended cyber tasks. It supports the relevance of multi-stage evaluation. Its human completion-time figures are estimates, not timed matched human trials, so they cannot serve as the human ground truth for this classifier.')
h('2.5 Research gap and limits of the novelty claim')
p('The candidate contribution is an independently labeled comparison of human and autonomous-AI operations under equivalent tasks and observation methods, including unfamiliar model families and human task variants. The present literature review supports this as a question worth investigating; it does not establish that no prior work has answered every part. The current development result is not a new classification algorithm. The supplied package\'s broad first-passive-classifier claim was removed from the adopted research framing.')

chapter('Chapter 3. Methodology')
h('3.1 Design and evidence boundary')
p('The development protocol was frozen before model outcomes in the Git repository. The run used the two downloaded public archives and executed no logged attack commands. Both publisher checksums matched. The cloud batch was bounded by a stop watchdog and a $2 allocation within the existing research budget. No new human activity or autonomous attack activity was collected in this batch.')
p('The pipeline is: published logs -> strict prefix qualification -> identical text normalization -> record-group fitting/calibration/test separation -> classifier -> fixed-threshold prediction -> audit. Source, task and family fields are used for auditing and splits only. They are not model inputs.')
h('3.2 Data qualification')
table(['Source','Observed archive contents','Development eligibility'],[
    ['KYPO human training','267 useractions files; 21,411 JSON records','267 nonempty Bash-command record groups'],
    ['Honey for the Agent','16,199 session JSON files','13,397 prefixes with explicit parsed-command fields'],
    ['Agent exclusions','2,802 prefixes','Missing or ambiguous command fields in the observation prefix']])
p('The human README lists 275 trainees and 21,762 records, which do not match the archive counts observed here. A record file is not assumed to be a unique person. Agent logs include two-field turns that lack the explicit parsed-command pair. The supplied parser silently skipped these turns. The development builder instead rejects an entire selected prefix if it cannot establish its command strings from the explicit field. This preserves an audit of exclusion rather than pretending missing actions did not occur. Exclusion may itself introduce selection bias.')
h('3.3 Observation and preprocessing')
p('Each record contains its first ten eligible commands, or all commands if the session is shorter. Human records use Bash-command entries; agent records require explicit parsed-command strings. IP addresses and standalone numbers are normalized identically. Timing, terminal output, edit counters, input-file paths and operator metadata are excluded. No logged command is executed. A length-only model checks whether the number of available commands alone explains separation.')
h('3.4 Partitions and comparisons')
p('A deterministic seed of 11701 creates nominal 60/20/20 fitting, calibration and test partitions within each class. Repeated normalized prefixes are removed from later partitions if present in earlier ones. Within a partition, one example is retained per prefix-label pair; conflicting labels are not silently resolved in favor of one class. The main fitting sample has 160 human and 640 agent records after the prespecified agent cap. Calibration has 2,444 records and test has 2,437.')
p('The models are logistic regression on command count, TF-IDF word unigrams/bigrams with logistic regression, and a small histogram gradient-boosted classifier over 16 content/sequence features. The n-gram model is related to, but not an exact reproduction of, Trace\'s linear-SVM method. A direct author-matched Trace baseline remains required for the proposed confirmatory comparison.')
p('For each model, the threshold is selected using human calibration scores to obtain an empirical false-positive rate no greater than 5%. The selected export is the model with the highest calibration AUROC, with a prespecified alphabetical tie-break. Test results do not choose the winner or threshold. The same three methods are evaluated on four held-out agent-family conditions and seven held-out human-task conditions, giving 36 comparisons in total.')
h('3.5 Metrics and audit')
p('Agent recall is the fraction of agent records labeled AI. Human false-positive rate is the fraction of human records mislabeled AI. AUROC describes ranking across thresholds; it does not guarantee a useful operating point. Saved per-record predictions permit independent recomputation of confusion counts, AUROC and threshold selection. The audit also checks that held groups and repeated prefixes do not leak across the applicable partitions.')
p('The reused human groups and task-dependent records do not provide 36 independent replications. Participant identity across files is not established. Accordingly, the report does not offer a population risk certificate or use command counts as independent samples. The supplied conformal implementation is not used: prediction-set coverage would not by itself bound error among accepted decisions, and the draft had a small-calibration quantile defect.')

chapter('Chapter 4. Results')
h('4.1 Main held-out record comparison')
table(['Model','AUROC','Agent recall','Human false positives'],[
    [r['model'],f"{r['test_auroc']:.4f}",f"{100*r['agent_recall']:.1f}% ({r['tp']}/{r['tp']+r['fn']})",f"{100*r['human_fpr']:.1f}% ({r['fp']}/{r['fp']+r['tn']})"] for r in base])
p('Calibration selected the n-gram model. It labeled every held-out agent record AI, but four human records also received that label. The test false-positive rate of 7.4% exceeds the proposed 5% ceiling even though calibration was below it. The length baseline was near chance in ranking, while the tree model obtained substantial separation with a larger false-positive rate. None of these numbers resolves the source confound.')
h('4.2 Held-out agent families')
table(['Held family','Agent recall','Human false positives'],[
    [r['name'].replace('family_',''),f"{100*r['agent_recall']:.2f}%",f"{100*r['human_fpr']:.2f}% ({r['fp']}/{r['fp']+r['tn']})"] for r in families])
p('The n-gram classifier retained high recall across these held-family probes. However, the agent cases still come from Honey and the human cases from KYPO. The probes therefore show transfer within this source arrangement, not independent evidence that operator type has been isolated. The repeated human test groups also make the fold results dependent.')
h('4.3 Held-out human tasks')
table(['Held human exercise','Agent recall','Human false positives'],[
    [r['name'].replace('human_task_','').replace('_',' '),f"{100*r['agent_recall']:.1f}%",f"{100*r['human_fpr']:.1f}% ({r['fp']}/{r['fp']+r['tn']})"] for r in tasks])
p('Task transfer is the clearest weakness. In Junior hacker adaptive, 39 of 62 human records were labeled AI. In SQL injection, five of eight were labeled AI; that small denominator must remain visible. Other tasks had lower rates, including two with zero observed errors on small groups. These results warn that familiar-task accuracy does not establish reliable operator classification.')
h('4.4 What the classifier learned')
p('Post-hoc coefficient inspection of the selected linear model shows associations with ordinary command tokens such as cat and etc. Such coefficients describe patterns in the fitting collection; they are not causal evidence of AI involvement. Differences in tasks and command habits are plausible contributors to the score. The experiment does not isolate which contributor caused each prediction.')
h('4.5 Verification and hypothesis status')
p('All 361 artifact-audit checks passed. The exported CLI reproduced saved-model scores on held-out examples of both labels. These checks establish arithmetic and implementation consistency; they do not validate the labels as real APT campaigns or establish independent participants. H1 and H2 remain unconfirmed because the matched design is absent; the main empirical human false-positive target was also exceeded. H3 is not tested as a paired comparison because the current batch omitted timing entirely.')

chapter('Chapter 5. Conclusions and Required Confirmation')
h('5.1 Main conclusion')
p('A working command-based classifier can strongly distinguish the two available public collections. This is a completed development result, not proof that it distinguishes autonomous-AI APT attacks from human-operated APT attacks. The human-task holdouts expose a substantial failure mode that the near-perfect AUROC would otherwise hide. The current exported model is a research prototype and should not be used to assert AI involvement in a live investigation.')
h('5.2 What must happen next')
p('The next valid advance is a matched task experiment. The public Junior hacker scenario and its historical repository revisions have been recovered. It includes discovery, remote access, collection and transfer. That provides a concrete route toward reusing documented human tasks for agent runs. However, the supplied human topology and current scenario differ, and the historical execution environment has not been reproduced. A current download is not sufficient evidence of equivalence.')
p('Before new agent collection, fix the exact task revision, environment, instructions, command recorder and assistance rules. Keep all human task material separate from the classifier inputs. Audit whether the existing human population and logs support the intended split. Collect a matched human cohort if adequate equivalence cannot be established from the published source. Any new participant work must follow the institution\'s applicable determination; no bot-generated activity can substitute for human ground truth.')
p('Then reserve people, task variants and AI families before fitting. Use a fixed false-alert target and report intervals that respect the actual independent sampling units. Add a conventional-script control to test whether the classifier merely detects automation. Evaluate the exact Trace-style baseline alongside the simpler methods. Broader APT claims require independently verified multi-stage or campaign data, not merely adding APT terminology to shell-session labels.')
h('5.3 Research product and contribution')
p('The completed product is an executable research classifier with reproducible data qualification, frozen development rules, saved predictions, threshold selection and negative transfer evidence. A potential doctoral contribution would come from the matched confirmation and its operationally meaningful false-alert limits. The current study does not establish a novel classifier algorithm, universal detection guarantee, SOC benefit or defense-ready Praxis.')
h('5.4 Disclosure')
p('AI assistance was used to develop code, review literature, audit artifacts and draft this document. The candidate remains responsible for source verification, interpretation, institutional requirements and the final submitted text. The document makes no promise about an AI-writing detector score. References, calculations and substantive claims should be checked during faculty review.')

chapter('References')
refs=[
'Ediga, M., & Chattopadhyay, S. (2026). Trace: Unmasking AI attack agents through terminal behavior fingerprinting [Preprint]. arXiv. https://arxiv.org/abs/2605.01186',
'Reworr, & Volkov, D. (2024). LLM agent honeypot: Monitoring AI hacking agents in the wild [Preprint]. arXiv. https://arxiv.org/abs/2410.13919',
'Svabensky, V., Vykopal, J., Seda, P., & Celeda, P. (2021). Dataset of shell commands used by participants of hands-on cybersecurity training. Data in Brief, 38, 107398. https://doi.org/10.1016/j.dib.2021.107398',
'Svabensky, V., Vykopal, J., Seda, P., & Celeda, P. (2022). Dataset: Shell commands used by participants of hands-on cybersecurity training (Version 2) [Data set]. Zenodo. https://zenodo.org/records/6670113',
'Maddaloni, D., Safargalieva, A., & Vasilomanolakis, E. (2026). Honey for the Agent: Cyber deception and behavioral fingerprinting of LLM-based attackers [Data set]. Zenodo. https://doi.org/10.5281/zenodo.20818246',
'Folkerts, L., Payne, W., Inman, S., Giavridis, P., Skinner, J., Deverett, S., Aung, J., Zorer, E., Schmatz, M., Ghanem, M., Wilkinson, J., Steer, A., Hong, V., & Wang, J. (2026). Measuring AI agents\' progress on multi-step cyber attack scenarios [Preprint]. arXiv. https://arxiv.org/abs/2603.11214',
'Angelopoulos, A. N., & Bates, S. (2021). A gentle introduction to conformal prediction and distribution-free uncertainty quantification [Preprint]. arXiv. https://arxiv.org/abs/2107.07511',
'MUNI-KYPO-TRAININGS. (n.d.). Junior hacker training [Source code and training design]. GitLab. https://gitlab.ics.muni.cz/muni-kypo-trainings/games/junior-hacker']
for ref in refs:
    par=p(ref);par.paragraph_format.first_line_indent=Inches(-.3);par.paragraph_format.left_indent=Inches(.3)

chapter('Appendix A. Reproducibility and Classifier Use')
p('The research evidence is in experiments/praxis_next/aivh_20261005. DEVELOPMENT_PROTOCOL.txt records the frozen rules; DEVELOPMENT_DATA.json records exclusions and input hash; CLOUD_FREEZE.json identifies executed code. RESULTS_TABLE.csv gives all 36 comparisons. AUDIT_RESULTS.json records the independent recount. Model versions are pinned in requirements-lock.txt, and the worker receipt records compute and shutdown separately.')
p('The classifier reads a JSON array of recorded command strings. It never executes those strings. Install the pinned dependencies, then run predict.py with --model pointing to the trusted supplied ngrams.joblib and --commands pointing to the JSON file. The output gives a research label, AI score, fixed threshold and the model scope. The score is not a calibrated probability that a real attack is AI-operated.')
p('The delivery ZIP contains the selected classifier, inference code, its feature dependencies, an example input, the dependency lock and model card. The original supplied package remains preserved separately; its synthetic results and certification claims are not substituted for this AWS experiment. Reproduction should verify archive and artifact hashes and preserve all exclusions.')
doc.save(DEST/'PX117_Preliminary_Praxis.docx')
print(DEST/'PX117_Preliminary_Praxis.docx')
