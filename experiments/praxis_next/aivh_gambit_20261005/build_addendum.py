"""Plain-language Praxis results addendum, generated from saved evidence."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
HERE=Path(__file__).resolve().parent
R=json.loads((HERE/'RESULTS.json').read_text());doc=Document()
section=doc.sections[0]
section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Inches(1)
section.page_width=Inches(8.5);section.page_height=Inches(11)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7)
def p(s):doc.add_paragraph(s)
def h(s):doc.add_heading(s,1)
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for c,s in zip(t.rows[0].cells,headers):c.text=s
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 for row in rows:
  rr=t.add_row();rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c,s in zip(rr.cells,row):c.text=s
doc.add_heading('Autonomous AI versus Human-Operated Attack Commands',0)
p('PX-117C • Preliminary Praxis results addendum • Gary Pagan • 5 October 2026')
h('What we tested')
p('We tested whether command behavior can distinguish autonomous AI activity from human-operated activity. Both sides were restricted to the same set of discovery-like commands. Each observation contained ten commands, and related human records stayed in one partition. The models never executed any command.')
p('The human data came from GAMBiT; the AI data came from Honey for the Agent. After cleaning and command filtering, the development pool contained 244 human windows across 42 conservative participant groups and 7,659 AI sessions. The main test contained 61 human windows from nine groups and 1,300 AI windows.')
h('Main findings')
base=[r for r in R['results'] if r['condition']=='participant_holdout']
names={'length':'Command length','lexical':'Command words','verbs':'Executable sequences'}
table(['Model','AI identified','Human false AI calls','Mean group error'],[[names[r['model']],f"{r['agent_recall']:.1%}",f"{r['fp']}/{r['fp']+r['tn']} ({r['human_window_fpr']:.1%})",f"{r['human_group_mean_fpr']:.1%}"] for r in base])
p('The command-word model caught 97.6% of AI windows, but exceeded the human error target. The executable-sequence model caught 71.5%, with no observed false AI calls. Neither met both main targets: at least 80% AI recall and no more than 5% mean human-group error.')
p('Mean group error gives every human group equal weight. It differs from the percentage of all human windows mislabeled because some groups contribute more windows than others.')
h('A separate human dataset')
table(['Model','KYPO false AI calls'],[[names[r['model']],f"{r['external_kypo']['fp']}/180 ({r['external_kypo']['human_window_fpr']:.1%})"] for r in base])
p('KYPO was not used to fit these models or choose their thresholds. Its result supports a narrower claim about transfer to another human command source. A human-only test cannot establish AI detection on an independent AI dataset.')
h('Unfamiliar AI families')
families=[r for r in R['results'] if r['model']=='verbs' and r['condition'].startswith('family_')]
table(['Held-out family','AI identified','Mean human-group error'],[[r['condition'][7:],f"{r['tp']}/{r['tp']+r['fn']} ({r['agent_recall']:.1%})",f"{r['human_group_mean_fpr']:.1%}"] for r in families])
p('The sequence model met the separate 70% recall and 5% group-error point targets in two of four family probes. DeepSeek had only 16 eligible records. The same human test groups were reused, so these probes are dependent. All results are retained.')
h('How we checked it')
p('AWS ran 18 comparisons across three model types. The protocol was committed before scoring. Thresholds came from calibration data; test results did not change them. An independent audit passed 135 checks covering group separation, duplicate exclusion, withheld families, confusion counts, AUROC, threshold choice and model selection. Six CLI checks reproduced saved-model scores.')
p('Shuffling fitting labels reduced the word and sequence models to AUROC values of 0.558 and 0.572. That is a useful diagnostic, but one shuffle is not a statistical significance test. Calibration selected the word model; the sequence model remains a reported comparator, not a replacement selected after seeing test results.')
h('What this means for the Praxis')
p('This is a useful preliminary result: command sequences offer a more conservative operating point, while command wording identifies more AI activity. The experiment measures that tradeoff. It has not yet established the full success criterion or a novel, general APT attribution method.')
p('The most important limitation remains the source difference. Human and AI records came from different collections. A shared command filter does not make their tasks or environments identical. The discovery subset also does not cover an entire APT campaign.')
p('Only nine human groups were tested internally. Zero observed errors does not prove zero risk. The empirical bootstrap interval becomes [0,0] for an all-zero sample, which is not a population guarantee. The word model’s group-error interval was broad: 1.2% to 20.1%.')
h('Next experiment')
p('Keep both candidates and freeze a new test before further fitting. Measure sensitivity to participant splits and qualify an independent AI source. Choose the required number of independent human groups before making a low-false-alert claim. Do not adjust the current thresholds against these exposed test records.')
h('Sources and evidence')
p('GAMBiT dataset paper: https://doi.org/10.1016/j.dib.2026.112476')
p('Honey for the Agent: https://zenodo.org/records/20818246')
p('KYPO human command data: https://zenodo.org/records/6670113')
p('Protocol and executable freeze: Git commit f8ce737. Experiment directory: experiments/praxis_next/aivh_gambit_20261005. Full comparisons: RESULTS_TABLE.csv. Saved predictions and runnable classifier are included in delivery/.')
p('AI assistance was used to implement, audit and draft this research package. The candidate remains responsible for verifying the evidence and submitted manuscript. This addendum is preliminary and does not claim defense readiness.')
dest=HERE/'delivery';dest.mkdir(exist_ok=True);doc.save(dest/'PX117C_Praxis_Addendum.docx')
print(dest/'PX117C_Praxis_Addendum.docx')
