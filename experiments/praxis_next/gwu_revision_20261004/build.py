"""Preserve reviewed GWU format and equation accessibility; add executive summary."""
import argparse,ast,importlib.util,json,re
from pathlib import Path
from docx.shared import Pt
H=Path(__file__).resolve().parent;BASE=H.parent;OUT=H/'delivery';QA=Path('C:/w/praxis_revision_20261004_qa')
spec=importlib.util.spec_from_file_location('gwu',BASE/'gwu_final_20260924/render_gwu.py');mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
tree=ast.parse((BASE/'gwu_final_20260928/integration/correct_committee.py').read_text(encoding='utf-8'));data={}
for node in tree.body:
    if isinstance(node,ast.Assign):
        for target in node.targets:
            if isinstance(target,ast.Name) and target.id in ['linear','symbols','acronyms']:data[target.id]=ast.literal_eval(node.value)
class Revised(mod.GWURenderer):
    def caption(self,kind,text):
        p=super().caption(kind,text)
        if kind=='Table' and self.caption_inventory[-1]['number'] in ['3-1','5-1']:p.paragraph_format.page_break_before=True
        return p
    def frontmatter(self):
        super().frontmatter()
        replacements={'Manuscript prepared September 24, 2026':'Manuscript revised October 4, 2026','A Praxis submitted to':'A Praxis prepared for review by','Praxis direction and committee confirmation pending':'Committee review edition','Certification pending':'Review Status','The confirmed director and committee, examination date, and approved certification wording must be supplied through the university process before formal submission.':'Formal submission requires the university-confirmed director and committee, examination date and approved certification wording. This review copy does not supply or imply those approvals.'}
        for p in self.doc.paragraphs:
            for r in p.runs:
                for a,b in replacements.items():r.text=r.text.replace(a,b)
        ps=self.doc.paragraphs;si=next(i for i,p in enumerate(ps) if p.text=='List of Symbols');ai=next(i for i,p in enumerate(ps) if p.text=='List of Acronyms')
        for heading,entries,remove in [(ps[si],data['symbols'],ps[si+1:ai]),(ps[ai],data['acronyms']+['OR gate: Warn if any member warns','SOC: Security Operations Center','SSH: Secure Shell'],ps[ai+1:-1])]:
            for p in remove:p._element.getparent().remove(p._element)
            previous=heading._p
            for text in entries:
                p=self.doc.add_paragraph(text,style='GWU Front');p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1;p.paragraph_format.first_line_indent=Pt(0);previous.addnext(p._p);previous=p._p
        marker=next(p for p in self.doc.paragraphs if p.text=='@@TOC@@');marker.paragraph_format.page_break_before=True
        blocks=(H/'executive_summary.md').read_text(encoding='utf-8').strip().split('\n\n')
        for i,text in enumerate(blocks):
            p=marker.insert_paragraph_before('',style='GWU Front Heading' if i==0 else 'GWU Front')
            if i==0:p.text=text.removeprefix('# ');p.paragraph_format.page_break_before=True;p.paragraph_format.space_after=Pt(18)
            else:self.inline(p,text);p.paragraph_format.space_after=Pt(12);p.paragraph_format.line_spacing=1.15
    def equation(self,text):
        start=len(self.equations);before=len(self.doc.paragraphs);super().equation(text)
        for i,p in enumerate(self.doc.paragraphs[before:],start):
            p.paragraph_format.keep_with_next=True;plain=self.doc.add_paragraph('Linear notation: '+data['linear'][i]);plain.paragraph_format.first_line_indent=Pt(0);plain.paragraph_format.line_spacing=1;plain.paragraph_format.space_after=Pt(10)
            for r in plain.runs:r.font.size=Pt(10)
            p._p.addnext(plain._p)
mod.GWURenderer=Revised
mod.build(argparse.Namespace(source=H/'manuscript.md',abstract=H/'abstract.md',title=mod.TITLE,docx=OUT/'Gary_Pagan_Praxis_Review.docx',pdf=OUT/'Gary_Pagan_Praxis_Review.pdf',qa_dir=QA,receipt=H/'RENDER_RECEIPT.json'))
