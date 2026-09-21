"""Build the proposal DOCX from its reviewed Markdown; no research execution."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = "17365D"
GRAY = "596573"
TOKEN = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\)|\*[^*]+\*|https?://[^\s]+)")


def hyperlink(paragraph, text, url):
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), NAVY); props.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); props.append(underline)
    run.append(props)
    value = OxmlElement("w:t"); value.text = text; run.append(value)
    node.append(run); paragraph._p.append(node)


def inline(paragraph, text):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("["):
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
            hyperlink(paragraph, match[1], match[2])
        elif part.startswith("http"):
            label = "Project evidence (commit 9dd2b81)" if "/blob/9dd2b8" in part else part
            hyperlink(paragraph, label, part)
        else:
            paragraph.add_run(part)


def number_id(doc):
    numbering = doc.part.numbering_part.element
    abstracts = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(abstracts + [0]) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl"); level.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1."), ("lvlJc", "left")):
        node = OxmlElement("w:" + tag); node.set(qn("w:val"), value); level.append(node)
    props = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind"); indent.set(qn("w:left"), "300"); indent.set(qn("w:hanging"), "240"); props.append(indent)
    level.append(props); abstract.append(level); numbering.append(abstract)
    return numbering.add_num(abstract_id).numId


def set_table_geometry(table, widths):
    table.autofit = False
    props = table._tbl.tblPr
    existing = props.find(qn("w:tblW"))
    existing.set(qn("w:type"), "dxa"); existing.set(qn("w:w"), str(round(sum(widths) * 1440)))
    layout = props.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout"); props.append(layout)
    layout.set(qn("w:type"), "fixed")
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        cannot = OxmlElement("w:cantSplit"); trpr.append(cannot)
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, value in (("top", "60"), ("bottom", "60"), ("left", "80"), ("right", "80")):
                margin = OxmlElement("w:" + side); margin.set(qn("w:w"), value); margin.set(qn("w:type"), "dxa"); margins.append(margin)
            tcpr.append(margins)


def add_table(doc, lines):
    cells = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    cells = [r for r in cells if not all(re.fullmatch(r":?-+:?", c) for c in r)]
    count = len(cells[0])
    widths = [3.15, 3.75] if count == 2 else [2.30, 2.25, 2.35]
    table = doc.add_table(rows=len(cells), cols=count)
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, row in enumerate(cells):
        for j, text in enumerate(row):
            p = table.cell(i, j).paragraphs[0]
            p.style = doc.styles["Proposal Table"]
            inline(p, text)
            if i == 0:
                shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "E9EFF5"); table.cell(i, j)._tc.get_or_add_tcPr().append(shade)
                for run in p.runs:
                    run.bold = True; run.font.color.rgb = RGBColor.from_string(NAVY)
    repeat = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    after = doc.add_paragraph(); after.paragraph_format.space_after = Pt(1); after.paragraph_format.line_spacing = 0.2


def build(source: Path, output: Path):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(.8)
    section.top_margin = Inches(.72); section.bottom_margin = Inches(.7)
    section.header_distance = Inches(.28); section.footer_distance = Inches(.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("222222")
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in (("Title", 23, 0, 10), ("Heading 1", 15, 0, 9), ("Heading 2", 11.5, 8, 5)):
        style = doc.styles[name]
        style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.05; style.paragraph_format.keep_with_next = True
    table_style = doc.styles.add_style("Proposal Table", 1)
    table_style.base_style = normal; table_style.font.size = Pt(9.5)
    table_style.paragraph_format.line_spacing = 1.0
    table_style.paragraph_format.space_after = Pt(2); table_style.paragraph_format.space_before = Pt(2)
    list_style = doc.styles.add_style("Proposal Numbered", 1)
    list_style.base_style = normal
    list_style.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = "PRAXIS RESEARCH PROPOSAL"
    for run in header.runs:
        run.font.name = "Arial"; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Preliminary evidence | Proposed protection study | ")
    run.font.name = "Arial"; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(GRAY)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    doc.core_properties.title = "Reducing False Alarms While Preserving Lateral-Movement Detection"
    doc.core_properties.subject = "Final praxis proposal draft; audited preliminary evidence and prospective study"
    doc.core_properties.author = "Praxis research project"
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0; active_num = None
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line == "<!-- PAGE BREAK -->":
            doc.add_page_break(); active_num = None
        elif line.startswith("|"):
            end = i
            while end < len(lines) and lines[end].strip().startswith("|"):
                end += 1
            add_table(doc, lines[i:end]); i = end; continue
        elif line.startswith("# "):
            p = doc.add_paragraph(style="Title"); inline(p, line[2:])
        elif line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1"); inline(p, line[3:]); active_num = None
        elif line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2"); inline(p, line[4:]); active_num = None
        elif re.match(r"\d+\. ", line):
            match = re.match(r"(\d+)\. (.*)", line)
            if match[1] == "1" or active_num is None:
                active_num = number_id(doc)
            p = doc.add_paragraph(style="Proposal Numbered")
            numpr = p._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0; numpr.get_or_add_numId().val = active_num
            inline(p, match[2])
        else:
            p = doc.add_paragraph(); inline(p, line); active_num = None
        i += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return {"source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(), "docx_sha256": hashlib.sha256(output.read_bytes()).hexdigest(), "planned_pages": lines.count("<!-- PAGE BREAK -->") + 1}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=Path(__file__).with_name("PRAXIS_PROPOSAL.md"))
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(build(args.source, args.output)))
