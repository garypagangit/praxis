"""Format a reviewed empirical-praxis Markdown manuscript as a DOCX.

No research results are generated or changed. Relative repository links become
clickable public-source links; local-only links retain a readable file label.
Requires python-docx; Pillow is used only when local images are included.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import re
from urllib.parse import quote, unquote, urlsplit

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


NAVY, INK, GRAY, PALE = "17365D", "20252A", "5B6570", "EDF2F7"
DEFAULT_REPO_URL = "https://github.com/garypagangit/praxis/blob/apt-benchmark"
REPO_ROOT = Path(__file__).resolve().parents[4]
TEXT_WIDTH = 6.7
TOKEN = re.compile(r"(`[^`]+`|\*\*.+?\*\*|\[[^\]]+\]\([^\s)]+(?:\s+\"[^\"]*\")?\)|(?<!\*)\*[^*]+\*(?!\*)|https?://[^\s<>]+)")
LINK = re.compile(r'\[([^\]]+)\]\(([^\s)]+)(?:\s+"[^"]*")?\)')
HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def safe_text(text: str) -> str:
    """Use ordinary hyphens while retaining scientific symbols and accents."""
    return text.translate({ord(c): "-" for c in "\u2010\u2011\u2012\u2013\u2014\u2212"})


def plain(text: str) -> str:
    text = LINK.sub(lambda m: m[1], text)
    return re.sub(r"[`*]", "", safe_text(text))


def slug(text: str) -> str:
    text = re.sub(r"[^\w\s-]", "", plain(text).lower())
    return re.sub(r"\s+", "-", text.strip())


def add_field(paragraph, instruction: str):
    node = OxmlElement("w:fldSimple")
    node.set(qn("w:instr"), instruction)
    paragraph._p.append(node)


def link_run(paragraph, label: str, target: str | None = None, anchor: str | None = None, *, bold=False, italic=False):
    node = OxmlElement("w:hyperlink")
    if anchor:
        node.set(qn("w:anchor"), anchor)
    elif target:
        rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
        node.set(qn("r:id"), paragraph.part.relate_to(target, rel, is_external=True))
    node.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, value in (("color", NAVY), ("u", "single")):
        child = OxmlElement("w:" + tag); child.set(qn("w:val"), value); props.append(child)
    if bold:
        props.append(OxmlElement("w:b"))
    if italic:
        props.append(OxmlElement("w:i"))
    run.append(props)
    value = OxmlElement("w:t"); value.text = safe_text(label)
    value.set(qn("xml:space"), "preserve")
    run.append(value); node.append(run); paragraph._p.append(node)


class Renderer:
    def __init__(self, source: Path, repo_url: str, title: str | None):
        self.source = source.resolve()
        self.repo_url = repo_url.rstrip("/")
        self.lines = source.read_text(encoding="utf-8-sig").splitlines()
        self.doc = Document()
        first = next((HEADING.match(x.strip()) for x in self.lines if HEADING.match(x.strip())), None)
        self.title = title or (plain(first[2]) if first else source.stem.replace("_", " "))
        self.bookmarks: dict[str, str] = {}
        self.heading_occurrences: dict[str, int] = {}
        self.local_only_links: list[str] = []
        self.tables = self.images = self.headings = self.links = 0
        self.in_references = False
        self.current_numbering = None
        self.title_seen = False
        self.bookmark_counter = 1
        for line in self.lines:
            match = HEADING.match(line.strip())
            if match:
                key = slug(match[2])
                count = self.heading_occurrences.get(key, 0)
                self.heading_occurrences[key] = count + 1
                key = key + ("-" + str(count) if count else "")
                self.bookmarks[key] = "section_" + str(len(self.bookmarks) + 1)
        self.heading_occurrences.clear()
        self.configure()

    def configure(self):
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(.9)
        section.top_margin, section.bottom_margin = Inches(.78), Inches(.76)
        section.header_distance, section.footer_distance = Inches(.3), Inches(.32)
        section.different_first_page_header_footer = True
        normal = self.doc.styles["Normal"]
        normal.font.name, normal.font.size = "Times New Roman", Pt(11)
        normal.font.color.rgb = RGBColor.from_string(INK)
        normal.paragraph_format.line_spacing = 1.08
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.widow_control = True
        for name, size, before, after in (("Title", 24, 25, 15), ("Subtitle", 13, 0, 15),
                ("Heading 1", 16, 0, 12), ("Heading 2", 12.5, 13, 6), ("Heading 3", 11.5, 10, 5)):
            style = self.doc.styles[name]
            style.font.name, style.font.size = "Arial", Pt(size)
            style.font.bold = name != "Subtitle"
            style.font.color.rgb = RGBColor.from_string(NAVY if name != "Subtitle" else GRAY)
            style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
            style.paragraph_format.line_spacing = 1.07
            style.paragraph_format.keep_with_next = True
        for name, size in (("Praxis Table", 9.5), ("Praxis Code", 9), ("Praxis Caption", 9.5),
                           ("Praxis Reference", 10.5), ("Praxis List", 11)):
            style = self.doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
            style.font.size = Pt(size)
            style.paragraph_format.line_spacing = 1.04 if name != "Praxis List" else 1.08
            style.paragraph_format.space_after = Pt(4 if name != "Praxis Reference" else 8)
        table = self.doc.styles["Praxis Table"]
        table.font.name = "Arial"
        table.paragraph_format.space_before = Pt(2)
        table.paragraph_format.keep_together = True
        reference = self.doc.styles["Praxis Reference"]
        reference.paragraph_format.left_indent = Inches(.25)
        reference.paragraph_format.first_line_indent = Inches(-.25)
        reference.paragraph_format.keep_together = True
        self.doc.styles["Praxis Code"].font.name = "Consolas"
        self.doc.styles["Praxis Caption"].font.italic = True
        for header in (section.header,):
            p = header.paragraphs[0]
            p.text = "EMPIRICAL PRAXIS  |  Lateral-movement detection"
            for run in p.runs:
                run.font.name, run.font.size = "Arial", Pt(8)
                run.font.color.rgb = RGBColor.from_string(GRAY)
        for footer in (section.footer, section.first_page_footer):
            p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run("False alarms and lateral detection  |  ")
            run.font.name, run.font.size = "Arial", Pt(8)
            run.font.color.rgb = RGBColor.from_string(GRAY)
            add_field(p, "PAGE")
        self.doc.core_properties.title = self.title
        self.doc.core_properties.subject = "Empirical praxis: false alarms, lateral detection, and auditable evaluation"
        self.doc.core_properties.author = "Praxis research project"

    def resolve_link(self, target: str):
        if target.startswith("#"):
            return None, self.bookmarks.get(unquote(target[1:])), None
        split = urlsplit(target)
        if split.scheme in {"https", "http", "mailto"}:
            return target, None, None
        if split.scheme:
            raise ValueError("Unsupported hyperlink scheme: " + split.scheme)
        local = (self.source.parent / unquote(split.path)).resolve()
        try:
            relative = local.relative_to(REPO_ROOT)
        except ValueError:
            return None, None, Path(split.path).name
        if self.repo_url:
            url = self.repo_url + "/" + quote(relative.as_posix(), safe="/")
            if split.fragment:
                url += "#" + split.fragment
            return url, None, None
        return None, None, relative.as_posix()

    def inline(self, paragraph, text: str, *, bold=False, italic=False):
        for part in TOKEN.split(text):
            if not part:
                continue
            if part.startswith("`"):
                run = paragraph.add_run(safe_text(part[1:-1]))
                run.font.name, run.font.size = "Consolas", Pt(9)
                run.bold, run.italic = bold, italic
            elif part.startswith("**"):
                self.inline(paragraph, part[2:-2], bold=True, italic=italic)
            elif part.startswith("*"):
                self.inline(paragraph, part[1:-1], bold=bold, italic=True)
            elif part.startswith("[") and LINK.fullmatch(part):
                match = LINK.fullmatch(part)
                url, anchor, fallback = self.resolve_link(match[2])
                if url or anchor:
                    link_run(paragraph, match[1], url, anchor, bold=bold, italic=italic)
                    self.links += 1
                else:
                    text = match[1] + (" (linked file: " + fallback + ")" if fallback else "")
                    self.inline(paragraph, text, bold=bold, italic=italic)
                    self.local_only_links.append(match[2])
            elif part.startswith(("http://", "https://")):
                url = part.rstrip(".,;")
                link_run(paragraph, url, url, bold=bold, italic=italic)
                if len(url) < len(part):
                    paragraph.add_run(part[len(url):])
                self.links += 1
            else:
                run = paragraph.add_run(safe_text(part))
                run.bold, run.italic = bold, italic

    def heading(self, level: int, text: str):
        first = not self.title_seen and level == 1
        subtitle = self.title_seen and self.headings == 1 and level == 2 and text.lower().startswith(("an empirical", "a praxis", "a study"))
        if first:
            style = "Title"; self.title_seen = True
        elif subtitle:
            style = "Subtitle"
        elif level == 1 or text.lower() in {"abstract", "executive explanation"}:
            style = "Heading 1"
        else:
            style = "Heading " + str(min(3, level))
        p = self.doc.add_paragraph(style=style)
        if not first and (level == 1 or text.lower() == "abstract"):
            p.paragraph_format.page_break_before = True
        self.inline(p, text)
        key = slug(text); count = self.heading_occurrences.get(key, 0)
        self.heading_occurrences[key] = count + 1
        key += "-" + str(count) if count else ""
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self.bookmark_counter)); start.set(qn("w:name"), self.bookmarks[key])
        end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(self.bookmark_counter))
        p._p.insert(1 if p._p.pPr is not None else 0, start)
        p._p.append(end); self.bookmark_counter += 1
        self.headings += 1
        self.in_references = text.lower() in {"references", "bibliography"}
        self.current_numbering = None

    def numbering(self, kind: str, start: int = 1):
        numbering = self.doc.part.numbering_part.element
        ids = [int(n.get(qn("w:abstractNumId"))) for n in numbering.findall(qn("w:abstractNum"))]
        abstract_id = max(ids + [0]) + 1
        abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(abstract_id))
        level = OxmlElement("w:lvl"); level.set(qn("w:ilvl"), "0")
        for tag, value in (("start", str(start)), ("numFmt", kind),
                           ("lvlText", "%1." if kind == "decimal" else "•"), ("lvlJc", "left")):
            node = OxmlElement("w:" + tag); node.set(qn("w:val"), value); level.append(node)
        props = OxmlElement("w:pPr"); indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "330"); indent.set(qn("w:hanging"), "230")
        props.append(indent); level.append(props); abstract.append(level); numbering.append(abstract)
        return numbering.add_num(abstract_id).numId

    def list_item(self, text: str, number: int | None):
        kind = "decimal" if number is not None else "bullet"
        if self.current_numbering is None or self.current_numbering[0] != kind or number == 1:
            self.current_numbering = (kind, self.numbering(kind, number or 1))
        p = self.doc.add_paragraph(style="Praxis List")
        num = p._p.get_or_add_pPr().get_or_add_numPr()
        num.get_or_add_ilvl().val = 0; num.get_or_add_numId().val = self.current_numbering[1]
        self.inline(p, text)

    @staticmethod
    def table_rows(lines: list[str]):
        rows = []
        for line in lines:
            cells, cell, code, escape = [], [], False, False
            for char in line.strip().strip("|"):
                if escape:
                    cell.append(char); escape = False
                elif char == "\\":
                    escape = True
                elif char == "`":
                    code = not code; cell.append(char)
                elif char == "|" and not code:
                    cells.append("".join(cell).strip()); cell = []
                else:
                    cell.append(char)
            cells.append("".join(cell).strip())
            if not all(re.fullmatch(r":?\s*-{3,}\s*:?", c) for c in cells):
                rows.append(cells)
        if not rows or len(rows[0]) < 2 or len(rows[0]) > 6:
            raise ValueError("Tables must contain two through six columns.")
        if any(len(row) != len(rows[0]) for row in rows):
            raise ValueError("Markdown table has inconsistent column counts.")
        return rows

    @staticmethod
    def table_widths(rows):
        count = len(rows[0]); weights = []
        for col in range(count):
            body = [plain(row[col]) for row in rows[1:]] or [plain(rows[0][col])]
            numeric = sum(bool(re.fullmatch(r"[\d.,%+<>=\s-]+", x)) for x in body) >= max(1, math.ceil(len(body) * .75))
            mean = sum(min(110, len(x)) for x in body) / len(body)
            weight = 1.05 if numeric else max(1.25, min(3.2, math.sqrt(mean + 8) / 3.1))
            if col == 0 and not numeric:
                weight = max(weight, 1.7)
            weights.append(weight)
        minimum = .76 if count >= 5 else 1.15 if count == 4 else 1.45
        extra = TEXT_WIDTH - count * minimum
        total = sum(weights)
        return [minimum + extra * w / total for w in weights]

    def table(self, lines: list[str]):
        rows = self.table_rows(lines); widths = self.table_widths(rows)
        keep_whole = len(rows) <= 9 and sum(len(plain(c)) for row in rows for c in row) < 1800
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"; table.autofit = False
        props = table._tbl.tblPr
        width = props.find(qn("w:tblW")); width.set(qn("w:type"), "dxa"); width.set(qn("w:w"), str(round(TEXT_WIDTH * 1440)))
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for i, row in enumerate(table.rows):
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trpr.append(OxmlElement("w:tblHeader"))
            for j, cell in enumerate(row.cells):
                cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margins = OxmlElement("w:tcMar")
                for side, value in (("top", 70), ("bottom", 70), ("left", 85), ("right", 85)):
                    node = OxmlElement("w:" + side); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa"); margins.append(node)
                cell._tc.get_or_add_tcPr().append(margins)
                if i == 0 or i % 2 == 0:
                    shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), PALE if i == 0 else "F8FAFC")
                    cell._tc.get_or_add_tcPr().append(shade)
                p = cell.paragraphs[0]; p.style = "Praxis Table"
                p.paragraph_format.keep_with_next = i == 0 or (keep_whole and i < len(rows) - 1)
                self.inline(p, rows[i][j], bold=i == 0)
                if i and re.fullmatch(r"[\d.,%+<>=\s-]+", plain(rows[i][j])):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p = self.doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = .25
        self.tables += 1; self.current_numbering = None

    def image(self, alt: str, target: str):
        if urlsplit(target).scheme and not Path(target).is_absolute():
            raise ValueError("Images must be local files; download and verify them before document building.")
        path = (self.source.parent / unquote(target)).resolve()
        if not path.is_file():
            raise FileNotFoundError("Markdown image does not exist: " + str(path))
        from PIL import Image
        with Image.open(path) as im:
            ratio = im.height / im.width
        width = min(TEXT_WIDTH, 6.8 / ratio)
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = bool(alt)
        picture = p.add_run().add_picture(str(path), width=Inches(width))
        picture._inline.docPr.set("descr", plain(alt))
        if alt:
            p = self.doc.add_paragraph(style="Praxis Caption"); self.inline(p, alt)
        self.images += 1

    def render(self):
        i = 0
        while i < len(self.lines):
            line = self.lines[i].strip()
            if not line:
                i += 1; continue
            heading = HEADING.match(line)
            image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            ordered = re.match(r"(\d+)\.\s+(.+)", line)
            bullet = re.match(r"[-+*]\s+(.+)", line)
            if line == "<!-- PAGE BREAK -->":
                self.doc.add_page_break(); self.current_numbering = None
            elif line.startswith("<!--"):
                if "-->" not in line:
                    while i < len(self.lines) and "-->" not in self.lines[i]:
                        i += 1
            elif line.startswith("```"):
                i += 1
                while i < len(self.lines) and not self.lines[i].strip().startswith("```"):
                    p = self.doc.add_paragraph(style="Praxis Code")
                    p.add_run(safe_text(self.lines[i])); i += 1
            elif heading:
                self.heading(len(heading[1]), heading[2])
            elif image:
                self.image(image[1], image[2]); self.current_numbering = None
            elif line.startswith("|"):
                end = i
                while end < len(self.lines) and self.lines[end].strip().startswith("|"):
                    end += 1
                self.table(self.lines[i:end]); i = end; continue
            elif ordered:
                self.list_item(ordered[2], int(ordered[1]))
            elif bullet:
                self.list_item(bullet[1], None)
            elif re.fullmatch(r"[-*_]{3,}", line):
                self.current_numbering = None
            else:
                parts = [line]
                while i + 1 < len(self.lines):
                    next_line = self.lines[i + 1].strip()
                    if not next_line or re.match(r"^(?:#{1,6}\s|\||```|<!--|!\[|[-+*]\s|\d+\.\s)", next_line):
                        break
                    parts.append(next_line); i += 1
                p = self.doc.add_paragraph(style="Praxis Reference" if self.in_references else "Normal")
                self.inline(p, " ".join(parts)); self.current_numbering = None
            i += 1
        return self.doc


def build(source: Path, output: Path, *, repo_url: str = DEFAULT_REPO_URL, title: str | None = None):
    if source.resolve() == output.resolve() or output.suffix.lower() != ".docx":
        raise ValueError("Output must be a separate .docx path.")
    renderer = Renderer(source, repo_url, title)
    document = renderer.render()
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return {"source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "docx_sha256": hashlib.sha256(output.read_bytes()).hexdigest(),
            "title": renderer.title, "headings": renderer.headings, "tables": renderer.tables,
            "images": renderer.images, "hyperlinks": renderer.links,
            "local_only_links": renderer.local_only_links, "rendered_visual_qa": "REQUIRED_SEPARATELY"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--repo-url", default=DEFAULT_REPO_URL,
                        help="Public repository blob URL, preferably pinned to the final source commit.")
    parser.add_argument("--title", help="Optional document metadata title; manuscript text remains authoritative.")
    args = parser.parse_args()
    print(json.dumps(build(args.source, args.output, repo_url=args.repo_url, title=args.title), ensure_ascii=True))
