"""Build a complete Praxis manuscript as DOCX/PDF with editable mathematics.

Requirements: python-docx, matplotlib, pillow, pymupdf, latex2mathml, lxml;
LibreOffice for PDF; Microsoft's MML2OMML.XSL for editable Word equations.
No inference, experiment execution or cloud access.
"""
import argparse
import faulthandler
import datetime as dt
import hashlib
import json
import os
from pathlib import Path
import re
import subprocess

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
import pymupdf
from lxml import etree
from latex2mathml.converter import convert as latex_to_mathml

HERE=Path(__file__).resolve().parent
ACCENT='193E50'
MATH_TRANSFORM=None
DASHES=str.maketrans({chr(n):'-' for n in (0x2010,0x2011,0x2012,0x2013,0x2014,0x2212)})
def clean(s):return s.translate(DASHES)
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

def inline(p,s):
    pattern=r'(\$(?=[^$\n]*(?:\\|[_^=]))[^$\n]+\$|\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|(?<!\*)\*[^*]+\*(?!\*))'
    last=0
    for m in re.finditer(pattern,s):
        if m.start()>last:p.add_run(clean(s[last:m.start()]))
        token=m.group()
        if token.startswith('$'):
            mathml=etree.fromstring(latex_to_mathml(token[1:-1]).encode('utf-8'))
            p._p.append(MATH_TRANSFORM(mathml).getroot())
        elif token.startswith('['):
            label,target=re.match(r'\[([^\]]+)\]\(([^)]+)\)',token).groups()
            node=OxmlElement('w:hyperlink');node.set(qn('r:id'),p.part.relate_to(target,RT.HYPERLINK,is_external=True))
            run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),ACCENT);props.append(color)
            run.append(props);text=OxmlElement('w:t');text.text=clean(label);run.append(text);node.append(run);p._p.append(node)
        else:
            content=token[2:-2] if token.startswith('**') else token[1:-1]
            run=p.add_run(clean(content));run.bold=token.startswith('**');run.italic=token.startswith('*') and not token.startswith('**')
            if token.startswith('`'):run.font.name='Consolas';run.font.size=Pt(9)
        last=m.end()
    if last<len(s):p.add_run(clean(s[last:]))

def configure(doc,title,label):
    section=doc.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    for attr in ('top_margin','bottom_margin','left_margin','right_margin'):setattr(section,attr,Inches(1))
    section.header_distance=section.footer_distance=Inches(.45)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12)
    normal.paragraph_format.line_spacing=1.25;normal.paragraph_format.space_after=Pt(7)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',25),('Heading 1',16),('Heading 2',13)]:
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor.from_string(ACCENT)
        s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(13);s.paragraph_format.space_after=Pt(8);s.paragraph_format.line_spacing=1.05
    s=doc.styles['Caption'];s.font.name='Times New Roman';s.font.size=Pt(10);s.font.italic=True
    s.paragraph_format.line_spacing=1.1;s.paragraph_format.space_after=Pt(9)
    doc.core_properties.author='Gary Pagan';doc.core_properties.title=title
    doc.core_properties.subject=label+': complete empirical research manuscript'

def cover(doc,title,label):
    p=doc.add_paragraph(label);p.alignment=1;p.paragraph_format.space_before=Pt(55);p.paragraph_format.space_after=Pt(30)
    p.runs[0].bold=True;p.runs[0].font.color.rgb=RGBColor.from_string(ACCENT)
    p=doc.add_paragraph(title,'Title');p.alignment=1;p.paragraph_format.space_after=Pt(35)
    for text in ['Gary Pagan','September 2026','Completed empirical research manuscript']:
        p=doc.add_paragraph(text);p.alignment=1
    p=doc.add_paragraph('Methods, completed results, and an accompanying evidence package');p.alignment=1
    for r in p.runs:r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string('666666')
    sec=doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec.header.is_linked_to_previous=False;sec.footer.is_linked_to_previous=False
    number=OxmlElement('w:pgNumType');number.set(qn('w:start'),'1');sec._sectPr.append(number)
    p=sec.header.paragraphs[0];p.alignment=2;r=p.add_run(label);r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('666666')
    p=sec.footer.paragraphs[0];p.alignment=1
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)

def table(doc,lines):
    rows=[[v.strip() for v in line.strip().strip('|').split('|')] for line in lines]
    rows=[row for row in rows if not all(re.fullmatch(r':?-+:?',v.replace(' ','')) for v in row)]
    cols=len(rows[0]);tab=doc.add_table(rows=1,cols=cols);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
    if cols==6:widths=[.60,.65,.80,.80,1.80,1.85]
    elif cols==5:widths=[1.95,1.12,1.02,1.35,1.06]
    elif cols==4:widths=[1.7,1.3,1.3,2.2]
    elif cols==3:widths=[3.1,1.7,1.7]
    elif cols==2:widths=[2.0,4.5]
    else:widths=[6.5/cols]*cols
    for c,w in zip(tab.columns,widths):c.width=Inches(w)
    for i,values in enumerate(rows):
        row=tab.rows[0] if i==0 else tab.add_row();props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
        if i==0:props.append(OxmlElement('w:tblHeader'))
        for j,value in enumerate(values):
            cell=row.cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'DFEAF0' if i==0 else 'F4F7F8' if i%2==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(fill)
            p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
            if len(rows)<=10:p.paragraph_format.keep_with_next=i<len(rows)-1
            inline(p,value)
            for run in p.runs:run.font.size=Pt(9.5 if cols>=4 else 10);run.bold=i==0 or run.bold
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def picture(doc,path,caption=None,max_width=6.4):
    with Image.open(path) as img:w,h=img.size
    width=min(max_width,6.8*w/h)
    p=doc.add_paragraph();p.alignment=1;p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=bool(caption)
    p.add_run().add_picture(str(path),width=Inches(width))
    if caption:
        p=doc.add_paragraph(clean(caption),'Caption');p.alignment=0

def diagram(path):
    path.parent.mkdir(parents=True,exist_ok=True)
    fig,ax=plt.subplots(figsize=(9,4.2));ax.set_xlim(0,10);ax.set_ylim(0,5);ax.axis('off')
    nodes={'code':(0.3,3.65,2.1,.85,'Fixed code pair\nand specification'),
           'supplier':(2.95,3.65,2.1,.85,'Supplier W\nup to 16 tests'),
           'display':(5.6,3.65,2.1,.85,'Disclosed records\nup to 2 tests'),
           'select':(2.95,1.85,2.1,.9,'Independent A\ncommit up to 8 IDs'),
           'execute':(5.6,1.85,2.1,.9,'Trusted execution\nof selected A tests'),
           'review':(8.1,2.7,1.65,.95,'Reviewer\naccept / keep'),
           'outcome':(2.95,.15,2.1,.85,'Reserved H\noutcome labels'),
           'analysis':(8.1,.15,1.65,.85,'Paired\nanalysis')}
    for key,(x,y,w,h,label) in nodes.items():
        box=FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.05',linewidth=1,edgecolor='#305466',facecolor='#E5EFF3' if key!='outcome' else '#F0EADB');ax.add_patch(box)
        ax.text(x+w/2,y+h/2,label,ha='center',va='center',fontsize=10,color='#132C36')
    def arrow(a,b):ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=11,linewidth=1,color='#43616E',connectionstyle='arc3'))
    for a,b in [((2.45,4.07),(2.85,4.07)),((5.1,4.07),(5.5,4.07)),((7.75,4.05),(8.55,3.68)),((5.1,2.3),(5.5,2.3)),((7.75,2.3),(8.55,2.62)),((8.95,2.65),(8.95,1.08)),((5.12,.575),(8,.575)),((1.35,3.55),(3.25,2.82)),((6.25,3.55),(4.4,2.82))]:arrow(a,b)
    ax.text(5.15,3.13,'input features only',fontsize=8,ha='center',color='#52656C',bbox={'facecolor':'white','edgecolor':'none','pad':2})
    fig.tight_layout(pad=.3);fig.savefig(path,dpi=220,facecolor='white');plt.close(fig)

def equation(doc,formula,transform):
    # Native Word mathematics remain editable and avoid rasterized math fonts.
    for part in formula.split(r'\qquad'):
        mathml=etree.fromstring(latex_to_mathml(part.rstrip(',')).encode('utf-8'))
        math=transform(mathml).getroot()
        p=doc.add_paragraph();p.alignment=1;p.paragraph_format.line_spacing=1.15
        for run in math.iter(qn('m:r')):
            props=OxmlElement('w:rPr');font=OxmlElement('w:rFonts')
            font.set(qn('w:ascii'),'Cambria Math');font.set(qn('w:hAnsi'),'Cambria Math');props.append(font)
            size=OxmlElement('w:sz');size.set(qn('w:val'),'22');props.append(size);run.insert(0,props)
        p._p.append(math)

def build(source,out,work,transform,label):
    content=source.read_text(encoding='utf-8');lines=content.splitlines();title=lines[0][2:]
    doc=Document();configure(doc,title,label);cover(doc,title,label)
    i=next(i for i,line in enumerate(lines) if line.lower().startswith('## executive summary'));tables=0;figures=[];equations=[]
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('#'):
            m=re.match(r'^(#+)\s+(.*)$',line);level=min(2,max(1,len(m[1])-1));p=doc.add_paragraph(clean(m[2]),f'Heading {level}')
            if level==1 and m[2].lower()!='executive summary':p.paragraph_format.page_break_before=True
            i+=1;continue
        if line=='$$':
            parts=[];i+=1
            while i<len(lines) and lines[i].strip()!='$$':parts.append(lines[i].strip());i+=1
            formula=' '.join(parts);equation(doc,formula,transform);equations.append(formula);i+=1;continue
        if line.startswith('$$') and line.endswith('$$'):
            formula=line[2:-2];equation(doc,formula,transform);equations.append(formula);i+=1;continue
        m=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if m:
            path=(source.parent/m[2]).resolve();picture(doc,path,m[1]);figures.append({'path':str(path),'sha256':sha(path)});i+=1;continue
        if line.startswith('|'):
            block=[]
            while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
            table(doc,block);tables+=1;continue
        if line.startswith('```'):
            i+=1
            while i<len(lines) and not lines[i].startswith('```'):
                p=doc.add_paragraph();p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(2)
                r=p.add_run(clean(lines[i]));r.font.name='Consolas';r.font.size=Pt(8);i+=1
            i+=1;continue
        paragraph=[line];i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','$$','![','```')):paragraph.append(lines[i].strip());i+=1
        p=doc.add_paragraph();inline(p,' '.join(paragraph))
        next_content=next((v.strip() for v in lines[i:] if v.strip()),'')
        if next_content.startswith('```'):p.paragraph_format.keep_with_next=True
        if paragraph[0].startswith('**Table '):p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(5)
        if re.match(r'^\[\d+\]',paragraph[0]):p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25);p.paragraph_format.line_spacing=1.1
    doc.save(out)
    return {'source_sha256':sha(source),'docx_sha256':sha(out),'tables':tables,'figures':figures,'equations':equations,'words_markdown_approx':len(content.split())}

def main():
    global MATH_TRANSFORM
    faulthandler.dump_traceback_later(45, repeat=True)
    p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--label',required=True);p.add_argument('--output-dir',type=Path,required=True);p.add_argument('--work-dir',type=Path,required=True)
    p.add_argument('--soffice',type=Path,default=Path('C:/Program Files/LibreOffice/program/soffice.exe'))
    p.add_argument('--math-xsl',type=Path,default=Path('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'));a=p.parse_args()
    a.output_dir.mkdir(parents=True,exist_ok=True);a.work_dir.mkdir(parents=True,exist_ok=True)
    source=a.source.resolve();docx=a.output_dir/'PAPER.docx';pdf=docx.with_suffix('.pdf')
    transform=etree.XSLT(etree.parse(str(a.math_xsl)))
    MATH_TRANSFORM=transform
    result=build(source,docx,a.work_dir,transform,a.label)
    profile=a.work_dir/'lo_profile';profile.mkdir(exist_ok=True)
    run=subprocess.run([str(a.soffice),'-env:UserInstallation='+profile.resolve().as_uri(),'--headless','--convert-to','pdf','--outdir',str(a.output_dir),str(docx)],capture_output=True,text=True,timeout=180,creationflags=subprocess.CREATE_NO_WINDOW if os.name=='nt' else 0)
    if run.returncode or not pdf.exists() or pdf.stat().st_mtime<docx.stat().st_mtime-1:raise RuntimeError('PDF render failed or stale: '+run.stdout+' '+run.stderr)
    pages=a.work_dir/'pages';pages.mkdir(exist_ok=True);document=pymupdf.open(pdf);geometry=[]
    for i,page in enumerate(document):
        page.get_pixmap(matrix=pymupdf.Matrix(1.5,1.5),alpha=False).save(pages/f'page-{i+1:02}.png')
        outside=[list(b[:4]) for b in page.get_text('blocks') if b[0]<28 or b[2]>page.rect.width-28 or b[1]<14 or b[3]>page.rect.height-14]
        geometry.append({'page':i+1,'text_chars':len(page.get_text()),'outside_safe_bounds':outside})
    result.update(created_utc=dt.datetime.now(dt.timezone.utc).isoformat(),builder_sha256=sha(Path(__file__)),pdf_sha256=sha(pdf),page_count=len(document),page_geometry=geometry,visual_review='PENDING',no_experiment_execution=True)
    document.close();(a.output_dir/'BUILD_MANIFEST.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:result[k] for k in ('page_count','words_markdown_approx','tables','source_sha256','docx_sha256','pdf_sha256')},indent=2))
    faulthandler.cancel_dump_traceback_later()

if __name__=='__main__':main()
