"""Build portable, styled Word reports from the reviewed Markdown sources.

Preset: standard_business_brief. Overrides: ResearchTitle (24 pt),
ResearchTable (9.5 pt), Code (8.5 pt), running furniture (9 pt).
No manuscript or human-approval claim is created by formatting.
"""
import json
import re
from pathlib import Path
from urllib.parse import quote
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPO = ROOT.parents[1]
WEB = "https://github.com/garypagangit/praxis/blob/Final-Praxis-Three-Closure-20260917/"
TOKENS = {
    "preset": "standard_business_brief", "header_pattern": "memo_masthead",
    "page_inches": [8.5, 11], "margins_inches": [1, 1, 1, 1],
    "header_footer_inches": 0.492, "font": "Calibri", "body_pt": 11,
    "body_spacing": {"before_pt": 0, "after_pt": 6, "line": 1.10},
    "headings": {"1": [16, "2E74B5", 16, 8], "2": [13, "2E74B5", 12, 6], "3": [12, "1F4D78", 8, 4]},
    "list": {"indent_dxa": 720, "hanging_dxa": 360, "after_pt": 8, "line_240": 280},
    "table": {"width_dxa": 9360, "indent_dxa": 120, "cell_margins_dxa": [80, 80, 120, 120], "header_fill": "F2F4F7"},
    "table_citation_text": {"before_pt": 4, "after_pt": 4},
    "named_overrides": {"ResearchTitle": "24pt ink blue; before0 after8", "ResearchTable": "9.5pt; line1.05; after3", "Code": "Consolas8.5pt; line1.05; after3", "Furniture": "Calibri9pt gray; line1; after0"}
}

def el(tag, **attrs):
    e = OxmlElement("w:" + tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), str(v))
    return e

def set_style(style, size, color="202A33", before=0, after=6, line=1.1, bold=False):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    f = style.paragraph_format
    f.space_before, f.space_after, f.line_spacing = Pt(before), Pt(after), line
    f.widow_control = True

def link_target(target, source):
    if re.match(r"https?://", target):
        return target
    rel = (source.parent / target.split("#")[0]).resolve().relative_to(REPO).as_posix()
    return WEB + quote(rel, safe="/")

def inline(p, s, source, size=None):
    # Links, bold, and code have real Word semantics; ordinary punctuation is retained.
    pattern = r"(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|(?<!\w)\*(?!\s)[^*]+(?<!\s)\*(?!\w)|`[^`]+`)"
    for item in re.split(pattern, s):
        if not item:
            continue
        match = re.fullmatch(r"\[([^\]]+)\]\(([^\)]+)\)", item)
        if match:
            label, target = match.groups()
            h = OxmlElement("w:hyperlink")
            h.set(qn("r:id"), p.part.relate_to(link_target(target, source), "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True))
            r = el("r")
            rp = el("rPr")
            rp.append(el("color", val="245D8F"))
            rp.append(el("u", val="single"))
            if size: rp.append(el("sz", val=round(size * 2)))
            r.append(rp)
            t = el("t"); t.text = label; r.append(t); h.append(r); p._p.append(h)
        else:
            bold = item.startswith("**") and item.endswith("**")
            italic = not bold and item.startswith("*") and item.endswith("*")
            code = item.startswith("`") and item.endswith("`")
            text = item[2:-2] if bold else item[1:-1] if code or italic else item
            r = p.add_run(text)
            if bold: r.bold = True
            if italic: r.italic = True
            if code:
                r.font.name = "Consolas"; r.font.size = Pt(size or 9)
            if size: r.font.size = Pt(size)

def numbering(doc):
    n = doc.part.numbering_part.element
    for idx, fmt, marker in ((90, "bullet", "•"), (91, "decimal", "%1.")):
        a = el("abstractNum", abstractNumId=idx); a.append(el("multiLevelType", val="singleLevel"))
        lev = el("lvl", ilvl=0)
        for tag, vals in [("start", {"val": 1}), ("numFmt", {"val": fmt}), ("lvlText", {"val": marker}), ("lvlJc", {"val": "left"})]:
            lev.append(el(tag, **vals))
        pp = el("pPr"); tabs = el("tabs"); tabs.append(el("tab", val="num", pos=720)); pp.append(tabs)
        pp.append(el("ind", left=720, hanging=360)); pp.append(el("spacing", before=0, after=160, line=280, lineRule="auto")); lev.append(pp)
        a.append(lev); n.append(a)
        num = el("num", numId=idx); num.append(el("abstractNumId", val=idx)); n.append(num)

def table(doc, rows, source):
    cols = len(rows[0]); t = doc.add_table(rows=0, cols=cols); t.autofit = False
    if cols == 6: widths = [1500, 1200, 1740, 1250, 1900, 1770]
    elif cols == 5: widths = [2600, 1500, 1300, 2300, 1660]
    elif cols == 4: widths = [1880, 2320, 2580, 2580]
    elif cols == 3: widths = [4400, 2480, 2480]
    else: widths = [9360 // cols] * cols; widths[-1] += 9360 - sum(widths)
    pr = t._tbl.tblPr
    pr.find(qn("w:tblW")).set(qn("w:type"), "dxa"); pr.find(qn("w:tblW")).set(qn("w:w"), "9360")
    pr.append(el("tblInd", w=120, type="dxa")); pr.append(el("tblLayout", type="fixed"))
    margins = el("tblCellMar")
    for side, val in [("top",80),("bottom",80),("start",120),("end",120)]: margins.append(el(side,w=val,type="dxa"))
    pr.append(margins)
    borders = el("tblBorders")
    for side in ("top","left","bottom","right","insideH","insideV"): borders.append(el(side,val="single",sz=4,color="CBD3DA"))
    pr.append(borders)
    grid = t._tbl.tblGrid
    for c in list(grid): grid.remove(c)
    for w in widths: grid.append(el("gridCol",w=w))
    for ridx, values in enumerate(rows):
        row = t.add_row(); row._tr.get_or_add_trPr().append(el("cantSplit"))
        if ridx == 0: row._tr.get_or_add_trPr().append(el("tblHeader"))
        for j, value in enumerate(values):
            c = row.cells[j]; c.width = Inches(widths[j]/1440)
            tcpr = c._tc.get_or_add_tcPr(); tcpr.find(qn("w:tcW")).set(qn("w:w"), str(widths[j])); tcpr.append(el("vAlign",val="center"))
            if ridx == 0: tcpr.append(el("shd",fill="F2F4F7"))
            p = c.paragraphs[0]; p.style = "ResearchTable"; inline(p,value,source,9.5)
            if len(rows)<=8 and ridx<len(rows)-1: p.paragraph_format.keep_with_next=True
            if ridx == 0:
                for run in p.runs: run.bold=True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def build(source, out, short):
    doc = Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(.492)
    set_style(doc.styles["Normal"],11)
    set_style(doc.styles["Title"],24,"0B2545",0,8,1.05,True)
    set_style(doc.styles["Subtitle"],13,"54616E",0,12,1.1)
    for i, vals in TOKENS["headings"].items():
        size,col,bef,aft = vals; st=doc.styles["Heading "+i]; set_style(st,size,col,bef,aft,1.1,True); st.paragraph_format.keep_with_next=True
    for name,size,col,bef,aft,line in [("ResearchTable",9.5,"202A33",0,3,1.05),("Code",8.5,"243746",0,3,1.05),("Furniture",9,"67737E",0,0,1),("TableCitation",10,"465766",4,4,1.1)]:
        st = doc.styles.add_style(name,1); set_style(st,size,col,bef,aft,line)
    doc.styles["Code"].font.name="Consolas"
    numbering(doc)
    header=sec.header.paragraphs[0]; header.style="Furniture"; header.add_run("PRAXIS RESEARCH REPORT  |  "+short.upper())
    footer=sec.footer.paragraphs[0]; footer.style="Furniture"; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("17 September 2026  •  ")
    fld=el("fldSimple",instr="PAGE"); footer._p.append(fld)
    lines=source.read_text(encoding="utf-8").splitlines(); i=0; title_seen=False; abstract_seen=False; in_references=False; next_num_id=100; current_num_id=91; keep_section=False
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line: continue
        if line.startswith("```"):
            while i<len(lines) and not lines[i].startswith("```"):
                p=doc.add_paragraph(style="Code"); p.add_run(lines[i]); i+=1
            i+=1; continue
        if line.startswith("|"):
            block=[line]
            while i<len(lines) and lines[i].strip().startswith("|"): block.append(lines[i].strip()); i+=1
            rows=[]
            for row in block:
                vals=[v.strip() for v in row.strip("|").split("|")]
                if all(re.fullmatch(r":?-+:?",v) for v in vals): continue
                rows.append(vals)
            table(doc,rows,source); continue
        if line.startswith("# "):
            p=doc.add_paragraph(style="Title"); inline(p,line[2:],source); title_seen=True; continue
        if line.startswith("#"):
            level=len(line)-len(line.lstrip("#")); title=line[level:].strip()
            if title=="Abstract": abstract_seen=True
            in_references = title=="References"
            keep_section = "Conclusion" in title
            sty="Subtitle" if title_seen and not abstract_seen and level==2 else "Heading "+str(min(3,level-1))
            p=doc.add_paragraph(style=sty); inline(p,title,source)
            if in_references: p.paragraph_format.page_break_before=True
            if short=="010" and title.startswith("4.2"): p.paragraph_format.page_break_before=True
            continue
        match=re.match(r"^(?:([-*]) |(\d+)\. )(.*)",line)
        if match:
            if match[2]=="1":
                current_num_id=next_num_id; next_num_id+=1
                n=el("num",numId=current_num_id); n.append(el("abstractNumId",val=91)); override=el("lvlOverride",ilvl=0); override.append(el("startOverride",val=1)); n.append(override); doc.part.numbering_part.element.append(n)
            p=doc.add_paragraph(); pp=p._p.get_or_add_pPr(); num=el("numPr"); num.append(el("ilvl",val=0)); num.append(el("numId",val=90 if match[1] else current_num_id)); pp.append(num)
            p.paragraph_format.space_after=Pt(8); p.paragraph_format.line_spacing=280/240
            if in_references: p.paragraph_format.keep_together=True
            inline(p,match[3],source); continue
        paragraph=[line]
        while i<len(lines) and lines[i].strip() and not re.match(r"^(#|\||```|[-*] |\d+\. )",lines[i].strip()): paragraph.append(lines[i].strip()); i+=1
        p=doc.add_paragraph(); inline(p," ".join(paragraph),source)
        if keep_section: p.paragraph_format.keep_together=True
        if re.match(r"^\*\*Table \d",line):
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.keep_together=True
    doc.core_properties.title=next(x[2:] for x in lines if x.startswith("# "))
    doc.core_properties.author="AI-assisted research compilation for Gary Pagan; human review pending"
    doc.core_properties.subject="Completed experiment report with preserved evidence and human-review requirements"
    doc.core_properties.comments="Sources and claims preserved in the accompanying Markdown and evidence map. Not an institutional approval record."
    out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)
    print(out)

if __name__=="__main__":
    out=ROOT/"deliverables"
    for dirname,name,label in [("01_cti","CTI_Final_Experiment_Report.docx","CTI"),("02_008","008_Final_Experiment_Report.docx","008"),("03_010","010_Final_Experiment_Report.docx","010")]:
        build(ROOT/dirname/"PAPER.md",out/name,label)
    (ROOT/"code"/"DOCUMENT_STYLE_TOKENS.json").write_text(json.dumps(TOKENS,indent=2)+"\n",encoding="utf-8")
